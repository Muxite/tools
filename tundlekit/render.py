"""render: PDF pages, Office slides and contact sheets as PNGs (MANIFEST.md §8, after tundle's render.py).

Optional packages (pymupdf, Pillow, python-pptx, pywin32) are imported inside the functions that need them.
render_office never opens the source: it copies it into out_dir and works on the copy only.
"""
from __future__ import annotations

import contextlib
import importlib.util
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from tundlekit.cli_support import CliResult, common_flags
from tundlekit.registry import ToolError, tool

SHEET_PAD = 14          # px between and around thumbnails
SHEET_LABEL = 26        # px of label above each thumbnail
SMALL_PNG = 10_240      # renders under this many bytes are suspicious (blank slides)
DOC_DPI = 110
SLIDE_SIZE = (1600, 900)
OFFICE_PROCESSES = ("POWERPNT.EXE", "WINWORD.EXE")
OFFICE_EXTRA = ("EXCEL.EXE",)
OFFICE_POLL = 2.0       # seconds between office_check polls
BACKENDS = ("auto", "powerpoint", "libreoffice", "text")


# ---------------------------------------------------------------------------------------------- optional packages

def _has_module(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False


def load_pymupdf():
    """pymupdf (or its old name fitz) with its messages routed to stderr, or None when not installed."""
    try:
        import pymupdf as mod
    except ImportError:
        try:
            import fitz as mod
        except ImportError:
            return None
    try:
        mod.set_messages(fd=2)
    except Exception:  # older releases lack set_messages or its fd argument
        pass
    return mod


def _import_pymupdf():
    mod = load_pymupdf()
    if mod is None:
        raise ToolError("pymupdf is not installed (pip install pymupdf)")
    return mod


@contextlib.contextmanager
def native_stdout_to_stderr():
    """While native code runs, send anything written to file descriptor 1 to stderr (stdout carries JSON)."""
    try:
        sys.stdout.flush()
        saved = os.dup(1)
    except (OSError, ValueError, AttributeError):
        yield
        return
    try:
        os.dup2(2, 1)
    except OSError:
        os.close(saved)
        yield
        return
    try:
        yield
    finally:
        try:
            sys.stdout.flush()
        except (OSError, ValueError):
            pass
        os.dup2(saved, 1)
        os.close(saved)


def _import_pil():
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        raise ToolError("Pillow is not installed (pip install Pillow)") from None
    return Image, ImageDraw, ImageFont


def _import_pptx():
    try:
        from pptx import Presentation
    except ImportError:
        raise ToolError("python-pptx is not installed (pip install python-pptx)") from None
    return Presentation


# ---------------------------------------------------------------------------------------------- backend discovery

class ProcessListError(RuntimeError):
    """The process list (tasklist / ps) could not be read, so the Office state is unknown."""


def _scan_office(all_apps: bool) -> list[str]:
    """Running Office process names; ProcessListError when tasklist (Windows) or ps (elsewhere) fails."""
    if sys.platform != "win32":
        return _soffice_scan() if all_apps else []
    try:
        proc = subprocess.run(["tasklist", "/FO", "CSV", "/NH"], capture_output=True, text=True,
                              errors="replace", timeout=60)
    except (OSError, subprocess.SubprocessError, ValueError) as exc:
        raise ProcessListError(f"could not run tasklist ({exc})") from None
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or "").strip()[:200]
        raise ProcessListError(f"tasklist failed with exit code {proc.returncode}" + (f": {detail}" if detail else ""))
    images = set()
    for line in (proc.stdout or "").splitlines():
        first = line.split(",", 1)[0].strip().strip('"').upper()
        if first:
            images.add(first)
    names = OFFICE_PROCESSES + (OFFICE_EXTRA if all_apps else ())
    return [name for name in names if name in images]


def office_running(all_apps: bool = False) -> list[str]:
    """Names of running Office processes (POWERPNT.EXE, WINWORD.EXE) that COM automation would attach to.

    With all_apps (office_check, MANIFEST §16.7), EXCEL.EXE is reported too, and off Windows the running
    LibreOffice processes (`soffice`, from `ps`) are reported. Without it, the result is always [] off Windows.
    On Windows, if tasklist cannot be run, the state is unknown, so this raises ToolError rather than report
    "nothing running".
    """
    try:
        return _scan_office(all_apps)
    except ProcessListError as exc:
        if sys.platform != "win32":
            return []
        raise ToolError(f"refusing to use Office: {exc} to check for open Office") from None


_OFFICE_RUNNING = office_running


def _soffice_scan() -> list[str]:
    """Distinct LibreOffice process names from `ps` (`soffice`, `soffice.bin`); ProcessListError on failure."""
    try:
        proc = subprocess.run(["ps", "-A", "-o", "comm="], capture_output=True, text=True, errors="replace",
                              timeout=60)
    except (OSError, subprocess.SubprocessError, ValueError) as exc:
        raise ProcessListError(f"could not run ps ({exc})") from None
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or "").strip()[:200]
        raise ProcessListError(f"ps failed with exit code {proc.returncode}" + (f": {detail}" if detail else ""))
    found = []
    for line in (proc.stdout or "").splitlines():
        name = os.path.basename(line.strip())
        if name.lower().startswith("soffice") and name not in found:
            found.append(name)
    return found


def _com_registered(progid: str) -> bool:
    if sys.platform != "win32" or not _has_module("win32com"):
        return False
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid + "\\CLSID"):
            return True
    except OSError:
        return False


def _powerpoint_available(ext: str = ".pptx") -> bool:
    if ext == ".docx":
        return _com_registered("Word.Application") and _has_pymupdf()
    return _com_registered("PowerPoint.Application")


def _has_pymupdf() -> bool:
    return _has_module("pymupdf") or _has_module("fitz")


def _soffice_candidates() -> list[str]:
    if sys.platform == "win32":
        bases = [os.environ.get(k) for k in ("ProgramFiles", "ProgramFiles(x86)", "ProgramW6432")]
        bases += [r"C:\Program Files", r"C:\Program Files (x86)"]
        return [str(Path(b) / "LibreOffice" / "program" / "soffice.exe") for b in bases if b]
    if sys.platform == "darwin":
        return ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    found = ["/usr/bin/soffice", "/usr/local/bin/soffice", "/usr/lib/libreoffice/program/soffice",
             "/snap/bin/libreoffice"]
    found += sorted(str(p) for p in Path("/opt").glob("libreoffice*/program/soffice")) if Path("/opt").is_dir() else []
    return found


def find_soffice() -> str | None:
    """The LibreOffice executable: TUNDLEKIT_SOFFICE if set, else soffice/libreoffice on PATH, else install dirs."""
    env = os.environ.get("TUNDLEKIT_SOFFICE")
    if env:
        if Path(env).is_file():
            return str(Path(env).resolve())
        return shutil.which(env)
    for name in ("soffice", "libreoffice"):
        hit = shutil.which(name)
        if hit:
            return hit
    for cand in _soffice_candidates():
        if Path(cand).is_file():
            return cand
    return None


_INKSCAPE_PATHS = ("C:/Program Files/Inkscape/bin/inkscape.exe", "C:/Program Files (x86)/Inkscape/bin/inkscape.exe",
                   "/Applications/Inkscape.app/Contents/MacOS/inkscape")


def _svg_to_png() -> str | None:
    """The SVG to PNG converter diagram_render would use, in its order (MANIFEST §14.6):
    cairosvg, rsvg-convert, inkscape (paths for executables), then "pymupdf"."""
    if _has_module("cairosvg"):
        return "cairosvg"
    hit = shutil.which("rsvg-convert")
    if hit:
        return hit
    hit = shutil.which("inkscape") or next((p for p in _INKSCAPE_PATHS if os.path.isfile(p)), None)
    if hit:
        return hit
    if _has_pymupdf():
        return "pymupdf"
    return None


@tool(
    "render_backends",
    "Report which render backends and optional packages are available: PowerPoint automation (Windows, "
    "pywin32, Office), the LibreOffice executable path, pymupdf, Pillow, python-pptx, and an SVG to PNG "
    "converter. Use it before render_office or render_pdf to know what will work on this machine.",
    {"type": "object", "properties": {}, "additionalProperties": False},
    readOnlyHint=True,
)
def render_backends() -> dict:
    return {
        "powerpoint": _powerpoint_available(),
        "libreoffice": find_soffice(),
        "pymupdf": _has_pymupdf(),
        "pillow": _has_module("PIL"),
        "pptx": _has_module("pptx"),
        "svg_to_png": _svg_to_png(),
    }


@tool(
    "office_check",
    "Check whether Word, PowerPoint or Excel is running (Windows tasklist; elsewhere LibreOffice soffice "
    "via ps). Run it before any build or render that automates Office: COM attaches to the open instance. "
    "With wait (seconds), polls every 2 s until nothing is running or the time is up. Returns the running "
    "process names and ok (true when nothing is running).",
    {"type": "object",
     "properties": {
         "wait": {"type": "number", "minimum": 0,
                  "description": "seconds to keep polling while something is running (default 0)"},
     },
     "additionalProperties": False},
    readOnlyHint=True,
)
def office_check(wait: float | None = None) -> dict:
    import math
    import time

    if wait is None:
        wait = 0
    if isinstance(wait, bool) or not isinstance(wait, (int, float)) or not math.isfinite(wait) or wait < 0:
        raise ToolError(f"wait must be a finite number of seconds >= 0, got {wait!r}")
    deadline = time.monotonic() + wait

    def poll() -> list[str]:
        names = _scan_office(all_apps=True)          # ProcessListError: the state is unknown (§17.6)
        if office_running is not _OFFICE_RUNNING:     # a replaced hook (embedding code, tests) supplies the names
            names = office_running(all_apps=True)
        return list(names)

    try:
        running = poll()
        while running:
            left = deadline - time.monotonic()
            if left <= 0:
                break
            time.sleep(min(OFFICE_POLL, left))
            running = poll()
    except (ProcessListError, ToolError) as exc:
        return {"running": [], "ok": False, "error": f"cannot tell whether Office is running: {exc}"}
    return {"running": running, "ok": not running}


# ---------------------------------------------------------------------------------------------- PDF pages

def _check_pdf_file(pdf: str) -> Path:
    path = Path(pdf)
    if not path.is_file():
        raise ToolError(f"no such PDF file: {pdf}")
    try:
        with open(path, "rb") as f:
            head = f.read(1024)
    except OSError as exc:
        raise ToolError(f"cannot read {pdf}: {exc}") from None
    if b"%PDF-" not in head:
        raise ToolError(f"not a PDF file: {pdf}")
    return path


def _render_pages(pdf: Path, out_dir: Path, name: str, dpi: float | None = None, width: int | None = None,
                  first: int | None = None, last: int | None = None) -> tuple[int, list[str]]:
    """Render pages first..last of `pdf` to out_dir/<name % page>. Returns (page count, files).

    Either dpi, or a target pixel width per page (dpi derived from the page width).
    """
    pymupdf = _import_pymupdf()
    with native_stdout_to_stderr():
        return _render_with(pymupdf, pdf, out_dir, name, dpi, width, first, last)


def _render_with(pymupdf, pdf, out_dir, name, dpi, width, first, last) -> tuple[int, list[str]]:
    try:
        doc = pymupdf.open(str(pdf))
    except Exception as exc:  # pymupdf raises its own error types for damaged files
        raise ToolError(f"cannot open PDF {pdf}: {exc}") from None
    files = []
    try:
        if not doc.is_pdf:
            raise ToolError(f"not a PDF file: {pdf}")
        count = doc.page_count
        lo = 1 if first is None else first
        hi = count if last is None else min(last, count)
        if lo < 1:
            raise ToolError("first must be at least 1")
        if last is not None and last < lo:
            raise ToolError(f"last ({last}) is before first ({lo})")
        if count and lo > count:
            raise ToolError(f"first ({lo}) is past the last page ({count})")
        out_dir.mkdir(parents=True, exist_ok=True)
        for n in range(lo, hi + 1):
            page = doc[n - 1]
            page_dpi = dpi
            if width is not None:
                page_dpi = width * 72.0 / max(page.rect.width, 1)
            target = out_dir / (name % n)
            page.get_pixmap(dpi=max(1, round(page_dpi))).save(str(target))
            files.append(str(target))
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError(f"cannot render {pdf}: {exc}") from None
    finally:
        doc.close()
    return count, files


@tool(
    "render_pdf",
    "Render PDF pages to PNG files page-001.png, page-002.png, ... (numbered by page) in out_dir, which is "
    "created if needed. first/last are 1-based and inclusive; dpi defaults to 110. Needs pymupdf. "
    "Returns the PDF's page count and the files written.",
    {"type": "object",
     "properties": {
         "pdf": {"type": "string", "description": "path to the PDF"},
         "out_dir": {"type": "string", "description": "output directory"},
         "dpi": {"type": "number", "description": "resolution (default 110)", "exclusiveMinimum": 0},
         "first": {"type": "integer", "description": "first page, 1-based (default 1)", "minimum": 1},
         "last": {"type": "integer", "description": "last page, inclusive (default the last page)", "minimum": 1},
     },
     "required": ["pdf", "out_dir"],
     "additionalProperties": False},
    readOnlyHint=False,
)
def render_pdf(pdf: str, out_dir: str, dpi: float = DOC_DPI, first: int | None = None,
               last: int | None = None) -> dict:
    if dpi is None or dpi <= 0:
        raise ToolError("dpi must be positive")
    check_path_string(out_dir)
    path = _check_pdf_file(pdf)
    out = Path(out_dir)
    if out.exists() and not out.is_dir():
        raise ToolError(f"out_dir is not a directory: {out_dir}")
    count, files = _render_pages(path, out, "page-%03d.png", dpi=dpi, first=first, last=last)
    return {"pages": count, "files": files}


# ---------------------------------------------------------------------------------------------- contact sheets

def _label(path: Path) -> str:
    m = re.search(r"(\d+)$", path.stem)
    if not m:
        return path.stem
    return m.group(1).lstrip("0") or "0"


def _font(ImageFont):
    for name in ("arial.ttf", "DejaVuSans.ttf", "Arial.ttf"):
        try:
            return ImageFont.truetype(name, 18)
        except OSError:
            continue
    return ImageFont.load_default()


def _contact_sheets(images: list[Path], out_dir: Path, cols: int, rows: int, thumb_width: int) -> list[str]:
    Image, ImageDraw, ImageFont = _import_pil()
    font = _font(ImageFont)
    per_sheet = cols * rows
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    for sheet_no, start in enumerate(range(0, len(images), per_sheet), 1):
        chunk = images[start:start + per_sheet]
        thumbs = []
        for p in chunk:
            try:
                with Image.open(p) as src:
                    im = src.convert("RGB")
            except Exception as exc:  # Pillow raises several types for unreadable images
                raise ToolError(f"cannot read image {p}: {exc}") from None
            if im.width > thumb_width:
                height = max(1, round(im.height * thumb_width / im.width))
                im = im.resize((thumb_width, height), Image.LANCZOS)
            thumbs.append(im)
        cell_h = max(t.height for t in thumbs) + SHEET_LABEL + SHEET_PAD
        sheet = Image.new("RGB", (cols * (thumb_width + SHEET_PAD) + SHEET_PAD, rows * cell_h + SHEET_PAD), "white")
        draw = ImageDraw.Draw(sheet)
        for k, (p, im) in enumerate(zip(chunk, thumbs)):
            x = SHEET_PAD + (k % cols) * (thumb_width + SHEET_PAD)
            y = SHEET_PAD + (k // cols) * cell_h
            draw.text((x, y), _label(p), fill="black", font=font)
            sheet.paste(im, (x, y + SHEET_LABEL))
        target = out_dir / f"contact-{sheet_no:02d}.png"
        sheet.save(str(target))
        made.append(str(target))
    return made


@tool(
    "render_contact_sheet",
    "Tile images into contact sheets contact-01.png, contact-02.png, ... in out_dir: cols x rows thumbnails "
    "per sheet (default 4 x 3), each thumb_width px wide (default 480, never upscaled), aspect ratio kept, "
    "labelled with the number at the end of the file name. Use it to eyeball a whole deck or document at "
    "once. Needs Pillow.",
    {"type": "object",
     "properties": {
         "images": {"type": "array", "items": {"type": "string"}, "description": "image paths, in order"},
         "out_dir": {"type": "string", "description": "output directory"},
         "cols": {"type": "integer", "minimum": 1, "description": "thumbnails per row (default 4)"},
         "rows": {"type": "integer", "minimum": 1, "description": "rows per sheet (default 3)"},
         "thumb_width": {"type": "integer", "minimum": 1, "description": "thumbnail width in px (default 480)"},
     },
     "required": ["images", "out_dir"],
     "additionalProperties": False},
    readOnlyHint=False,
)
def render_contact_sheet(images: list[str], out_dir: str, cols: int = 4, rows: int = 3,
                         thumb_width: int = 480) -> dict:
    if not images:
        raise ToolError("no images given")
    for name, value in (("cols", cols), ("rows", rows), ("thumb_width", thumb_width)):
        if not isinstance(value, int) or value < 1:
            raise ToolError(f"{name} must be a positive integer")
    check_path_string(out_dir)
    paths = [Path(p) for p in images]
    missing = [str(p) for p in paths if not p.is_file()]
    if missing:
        raise ToolError("no such image: " + ", ".join(missing))
    out = Path(out_dir)
    if out.exists() and not out.is_dir():
        raise ToolError(f"out_dir is not a directory: {out_dir}")
    return {"sheets": _contact_sheets(paths, out, cols, rows, thumb_width)}


# ---------------------------------------------------------------------------------------------- office: safety

RENDER_MARKER = ".tundlekit-render"
_WINDOWS = sys.platform == "win32"
_WIN_BAD_CHARS = frozenset('<>"|?*')
_WIN_DEVICE = re.compile(r"(?i)(con|prn|aux|nul|conin\$|conout\$|clock\$|com[0-9¹²³]|"
                         r"lpt[0-9¹²³])(\..*)?")
_ADMIN_SHARE = re.compile(r"\\\\([^\\]+)\\([A-Za-z])\$(\\.*)?")


def check_path_string(value, what: str = "out_dir") -> None:
    """ToolError unless `value` is a usable path string (no NUL; on Windows no reserved characters or devices)."""
    if not isinstance(value, str) or not value or "\0" in value:
        raise ToolError(f"{what} is not a valid path: {value!r}")
    if not _WINDOWS:
        return
    body = value.replace("/", "\\")
    for prefix in ("\\\\?\\", "\\\\.\\"):
        if body.startswith(prefix):
            body = body[len(prefix):]
            break
    rest = os.path.splitdrive(body)[1]
    if ":" in rest or any(c in _WIN_BAD_CHARS or ord(c) < 32 for c in rest):
        raise ToolError(f"{what} is not a valid path: {value!r}")
    for part in rest.split("\\"):
        if _WIN_DEVICE.fullmatch(part.rstrip(" .")):
            raise ToolError(f"{what} is not a valid path (it names a device): {value!r}")


def _local_hosts() -> set[str]:
    import socket

    hosts = {"localhost", "127.0.0.1", "::1", "[::1]", "."}
    try:
        hosts.add(socket.gethostname().lower())
    except OSError:
        pass
    return hosts


def _plain_windows_spelling(p: str) -> str:
    """Drop a \\\\?\\ or \\\\.\\ prefix and turn \\\\localhost\\C$\\x into C:\\x."""
    p = p.replace("/", "\\")
    if p.startswith("\\\\?\\UNC\\"):
        p = "\\\\" + p[8:]
    elif p.startswith(("\\\\?\\", "\\\\.\\")):
        p = p[4:]
    m = _ADMIN_SHARE.fullmatch(p)
    if m and m.group(1).lower() in _local_hosts():
        p = m.group(2) + ":" + (m.group(3) or "\\")
    return p


def _norm(path: str) -> str:
    """A canonical string for comparing locations: real path, plain spelling, case-folded."""
    p = os.fspath(path)
    if _WINDOWS:
        p = _plain_windows_spelling(p)
    p = os.path.realpath(os.path.abspath(p))
    if _WINDOWS:
        p = _plain_windows_spelling(p)            # realpath may add the prefix back
    return os.path.normcase(os.path.normpath(p))


def _ident(path: str) -> tuple[int, int] | None:
    try:
        st = os.stat(path)
    except (OSError, ValueError):
        return None
    return (st.st_dev, st.st_ino) if st.st_ino else None


def _chain(path: str) -> list[str]:
    """`path` (normalised) and every ancestor of it, up to the root."""
    p = _norm(path)
    chain = [p]
    while os.path.dirname(p) != p:
        p = os.path.dirname(p)
        chain.append(p)
    return chain


def _is_link(path: str) -> bool:
    """A symlink, a junction, or any other reparse point: never descend into it."""
    try:
        st = os.lstat(path)
    except (OSError, ValueError):
        return False
    if stat.S_ISLNK(st.st_mode):
        return True
    return bool(getattr(st, "st_file_attributes", 0) & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0))


def _contains_git(top: str) -> bool:
    stack = [top]
    while stack:
        folder = stack.pop()
        try:
            entries = list(os.scandir(folder))
        except OSError:
            continue
        for entry in entries:
            if entry.name.lower() == ".git":
                return True
            try:
                if entry.is_dir(follow_symlinks=False) and not _is_link(entry.path):
                    stack.append(entry.path)
            except OSError:
                continue
    return False


def _protected_bases(src: str) -> list[tuple[str, str]]:
    homes = {os.path.expanduser("~")}
    try:
        homes.add(str(Path.home()))
    except (RuntimeError, OSError):
        pass
    for var in ("HOME", "USERPROFILE"):
        if os.environ.get(var):
            homes.add(os.environ[var])
    bases = [("the home directory or an ancestor of it", h) for h in sorted(homes)]
    bases.append(("the current directory or an ancestor of it", os.getcwd()))
    bases.append(("the source's directory or an ancestor of it", os.path.dirname(os.path.abspath(src))))
    return bases


def _refusal(out_dir: str, src: str) -> str | None:
    """Why out_dir must not be used (a message containing 'refusing'), or None. Touches nothing."""
    try:
        out_norm = _norm(out_dir)
    except (OSError, ValueError) as exc:
        raise ToolError(f"out_dir is not a valid path: {out_dir!r} ({exc})") from None
    out_id = _ident(out_dir)
    if os.path.dirname(out_norm) == out_norm or os.path.ismount(out_norm):
        return f"refusing to use a filesystem root as out_dir: {out_dir}"
    for label, base in _protected_bases(src):
        try:
            chain = _chain(base)
        except (OSError, ValueError):
            continue
        for protected in chain:
            if out_norm == protected or (out_id is not None and out_id == _ident(protected)):
                return f"refusing to use {label} as out_dir: {out_dir}"
    if not os.path.lexists(out_dir):
        return None
    if not os.path.isdir(out_dir):
        return f"refusing to use out_dir {out_dir}: it exists and is not a directory"
    try:
        entries = os.listdir(out_dir)
    except OSError as exc:
        return f"refusing to use out_dir {out_dir}: cannot list it ({exc})"
    if any(e.lower() == ".git" for e in entries) or _contains_git(out_dir):
        return f"refusing to use a directory containing .git as out_dir: {out_dir}"
    if entries and not os.path.isfile(os.path.join(out_dir, RENDER_MARKER)):
        return (f"refusing to empty {out_dir}: it is not empty and has no {RENDER_MARKER} marker "
                "(use a new or empty directory)")
    return None


def _remove_entry(path: str) -> None:
    if _is_link(path):
        try:
            os.rmdir(path)                        # directory symlink or junction: removes the link only
        except OSError:
            os.unlink(path)
    elif os.path.isdir(path):
        shutil.rmtree(path)                       # rmtree removes nested links without following them
    else:
        os.unlink(path)


def _empty_dir(out: Path) -> None:
    try:
        out.mkdir(parents=True, exist_ok=True)
        entries = os.listdir(out)
    except (OSError, ValueError) as exc:
        raise ToolError(f"cannot create out_dir {out}: {exc}") from None
    for name in entries:
        try:
            _remove_entry(os.path.join(out, name))
        except OSError as exc:
            raise ToolError(f"cannot clear {name} from out_dir: {exc}") from None
    try:
        with open(out / RENDER_MARKER, "wb"):
            pass
    except OSError as exc:
        raise ToolError(f"cannot write {RENDER_MARKER} in out_dir: {exc}") from None


# ---------------------------------------------------------------------------------------------- office: backends

def _write(path: Path, text: str) -> str:
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    return str(path)


def _notes_text(slide) -> str:
    """A slide's notes text; a missing notes slide or notes placeholder counts as empty notes (§16.1)."""
    try:
        if not slide.has_notes_slide:
            return ""
        tf = slide.notes_slide.notes_text_frame
        return tf.text if tf is not None else ""
    except Exception:  # damaged or unusual notes parts read as empty notes, never a crash
        return ""


def _write_notes(copy: Path, out: Path, prs=None) -> list[str]:
    if prs is None:
        prs = _open_pptx(_import_pptx(), copy)
    files = []
    for n, slide in enumerate(prs.slides, 1):
        files.append(_write(out / f"notes-{n:02d}.txt", _notes_text(slide)))
    return files


def _check_package(path: Path, what: str) -> None:
    """ToolError unless every entry of the zip package decompresses cleanly (§16.1 corrupt packages)."""
    import zlib

    try:
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                with z.open(info) as f:
                    while f.read(1 << 20):
                        pass
    except (zipfile.BadZipFile, zipfile.LargeZipFile, zlib.error, EOFError, OSError, NotImplementedError,
            RuntimeError, ValueError) as exc:
        raise ToolError(f"cannot read {path.name} as a {what}: corrupt package ({exc})") from None


def _open_pptx(Presentation, path: Path):
    try:
        return Presentation(str(path))
    except Exception as exc:  # python-pptx raises assorted errors for damaged packages
        raise ToolError(f"cannot read {path.name} as a presentation: {exc}") from None


class _RawShape:
    """A shape element python-pptx cannot wrap (p:contentPart, mc:AlternateContent, unknown tags)."""
    has_text_frame = False

    def __init__(self, element):
        self._element = element


_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_TREE_PROPS = {_P + "nvGrpSpPr", _P + "grpSpPr", _P + "extLst"}


def _shape_items(shapes) -> list:
    """Every child shape of a shape tree, wrapped by python-pptx where it can, else as a _RawShape."""
    tree = getattr(shapes, "_spTree", None)
    factory = getattr(shapes, "_shape_factory", None)
    if tree is None or factory is None:
        try:
            return list(shapes)
        except Exception:  # an unreadable shape tree gives no shapes rather than a crash
            return []
    items = []
    for el in tree:
        if not isinstance(el.tag, str) or el.tag in _TREE_PROPS:
            continue
        try:
            items.append(factory(el))
        except Exception:
            items.append(_RawShape(el))
    return items


def _iter_shapes(shapes):
    """Shapes in order, descending into group shapes (as deck_inspect reads them)."""
    for sh in _shape_items(shapes):
        inner = None
        if sh.__class__.__name__ == "GroupShape":
            try:
                inner = sh.shapes
            except Exception:
                inner = None
        if inner is not None:
            yield from _iter_shapes(inner)
        else:
            yield sh


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _xml_text(element) -> str:
    """Paragraph text of any shape element (a:p / a:t), ignoring mc:Fallback copies."""
    paras = []

    def walk(el, inside):
        tag = getattr(el, "tag", None)
        if not isinstance(tag, str) or tag == _MC_FALLBACK:
            return
        if tag == _A + "p" and not inside:
            parts = []
            for sub in el.iter():
                if sub.tag == _A + "t":
                    parts.append(sub.text or "")
                elif sub.tag == _A + "br":
                    parts.append("\v")
            paras.append("".join(parts))
            return
        for child in el:
            walk(child, inside)

    try:
        walk(element, False)
    except Exception:
        return ""
    return "\n".join(paras)


def _shape_text(sh) -> tuple[str | None, list]:
    """(text, [(run text, size)]) for a shape with text, or (None, []) when it holds no text frame.

    Recognised shapes are read through python-pptx; anything else (unrecognised types such as p:contentPart,
    or a shape python-pptx cannot read) is read from its XML, so its text still counts (§16.1)."""
    try:
        if getattr(sh, "has_text_frame", False):
            tf = sh.text_frame
            runs = []
            for p in tf.paragraphs:
                for r in p.runs:
                    try:
                        size = r.font.size or p.font.size or 0
                    except Exception:
                        size = 0
                    runs.append((r.text, size))
            return tf.text, runs
        if type(sh).__name__ not in ("_RawShape", "BaseShape"):
            return None, []           # pictures, tables, connectors: no text frame, as before
    except Exception:
        pass
    element = getattr(sh, "_element", None)
    if element is None:
        return None, []
    try:
        has_body = any(el.tag == _A + "p" for el in element.iter())
    except Exception:
        return None, []
    if not has_body:
        return None, []
    text = _xml_text(element)
    return text, [(text, 0)]


def _slide_frames(slide) -> tuple[list[str], int | None]:
    """(text of every text frame in shape order, index of the title frame or None).

    The title is the same as deck_inspect's (MANIFEST §14.9): the frame holding the non-blank run with the
    largest font (first such frame on ties), else the first non-blank frame.
    """
    frames, best_size, title = [], -1, None
    for sh in _iter_shapes(slide.shapes):
        text, runs = _shape_text(sh)
        if text is None:
            continue
        frames.append(text)
        for run_text, size in runs:
            if run_text.strip() and size > best_size:
                best_size, title = size, len(frames) - 1
    if title is None:
        title = next((k for k, t in enumerate(frames) if t.strip()), None)
    return frames, title


def _text_deck(copy: Path, out: Path) -> tuple[int, list[str]]:
    Presentation = _import_pptx()
    prs = _open_pptx(Presentation, copy)
    files = []
    for n, slide in enumerate(prs.slides, 1):
        frames, title = _slide_frames(slide)
        if title is None:
            parts = ["TITLE: "] + frames[1:] if frames else ["TITLE: "]
        else:
            parts = ["TITLE: " + frames[title]] + frames[:title] + frames[title + 1:]
        files.append(_write(out / f"slide-{n:02d}.txt", "\n\n".join(parts) + "\n"))
    files += _write_notes(copy, out, prs)
    return len(prs.slides), files


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _text_docx(copy: Path, out: Path) -> tuple[int, list[str]]:
    try:
        with zipfile.ZipFile(copy) as z:
            xml = z.read("word/document.xml")
        root = ElementTree.fromstring(xml)
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, OSError, EOFError, ValueError,
            NotImplementedError, RuntimeError) as exc:
        raise ToolError(f"cannot read {copy.name} as a Word document: {exc}") from None
    except Exception as exc:  # zlib.error and other decompression failures
        raise ToolError(f"cannot read {copy.name} as a Word document: corrupt package ({exc})") from None
    lines = []
    for p in root.iter(_W + "p"):
        style = p.find(f"{_W}pPr/{_W}pStyle")
        style_id = style.get(_W + "val") if style is not None else None
        parts = []
        for el in p.iter():
            if el.tag == _W + "t":
                parts.append(el.text or "")
            elif el.tag == _W + "tab":
                parts.append("\t")
            elif el.tag in (_W + "br", _W + "cr"):
                parts.append(" ")
        lines.append(f"[{style_id or 'Normal'}] {''.join(parts)}")
    return len(lines), [_write(out / "paragraphs.txt", "\n".join(lines) + ("\n" if lines else ""))]


def _powerpoint_deck(copy: Path, out: Path) -> tuple[int, list[str]]:
    import win32com.client

    app = win32com.client.Dispatch("PowerPoint.Application")   # never set Visible=False: PowerPoint refuses
    pres = None
    files = []
    try:
        pres = app.Presentations.Open(str(copy), ReadOnly=True, Untitled=False, WithWindow=False)
        count = pres.Slides.Count
        for i in range(1, count + 1):
            target = out / f"slide-{i:02d}.png"
            pres.Slides(i).Export(str(target), "PNG", *SLIDE_SIZE)
            files.append(str(target))
    finally:
        if pres is not None:
            pres.Close()
        if app.Presentations.Count == 0:
            app.Quit()
    return count, files


def _word_pdf(copy: Path, out: Path) -> Path:
    import win32com.client

    pdf = out / "src.pdf"
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0
    doc = None
    try:
        doc = app.Documents.Open(str(copy), ReadOnly=True, AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(pdf), 17)                  # 17 = wdExportFormatPDF
    finally:
        if doc is not None:
            doc.Close(0)
        if app.Documents.Count == 0:
            app.Quit()
    return pdf


def _libreoffice_pdf(copy: Path, out: Path, soffice: str) -> Path:
    with tempfile.TemporaryDirectory(prefix="tundlekit-lo-") as profile:
        cmd = [soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--norestore",
               "--convert-to", "pdf", "--outdir", str(out), str(copy)]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, errors="replace", timeout=600)
        except (OSError, subprocess.SubprocessError) as exc:
            raise ToolError(f"LibreOffice failed to start: {exc}") from None
    pdf = out / (copy.stem + ".pdf")
    if not pdf.is_file():
        detail = (proc.stderr or proc.stdout).strip()
        raise ToolError(f"LibreOffice did not produce a PDF (exit {proc.returncode}): {detail}")
    return pdf


def _visual(backend: str, copy: Path, out: Path, deck: bool) -> tuple[int, list[str]]:
    if backend == "powerpoint":
        if deck:
            return _powerpoint_deck(copy, out)
        pdf = _word_pdf(copy, out)
        return _render_pages(pdf, out, "page-%03d.png", dpi=DOC_DPI)
    pdf = _libreoffice_pdf(copy, out, find_soffice())
    if deck:
        return _render_pages(pdf, out, "slide-%02d.png", width=SLIDE_SIZE[0])
    return _render_pages(pdf, out, "page-%03d.png", dpi=DOC_DPI)


def _available(backend: str, ext: str) -> bool:
    if backend == "powerpoint":
        return _powerpoint_available(ext)
    if backend == "libreoffice":
        return find_soffice() is not None and _has_pymupdf()
    return True


@tool(
    "render_office",
    "Render a .pptx or .docx for review. Copies the source to out_dir/src.<ext> (the source is never opened), "
    "empties out_dir first, then renders the copy: backend powerpoint (Windows COM; refuses while PowerPoint "
    "or Word is running) or libreoffice gives slide-NN.png / page-NNN.png plus contact sheets; backend text "
    "gives slide-NN.txt and notes-NN.txt, or paragraphs.txt. auto picks powerpoint, then libreoffice, then "
    "text. Refuses an out_dir that is a filesystem root, the home directory, the source's directory or an "
    "ancestor of it, or a directory containing .git.",
    {"type": "object",
     "properties": {
         "src": {"type": "string", "description": "path to a .pptx or .docx"},
         "out_dir": {"type": "string", "description": "output directory (emptied first)"},
         "backend": {"type": "string", "enum": list(BACKENDS), "description": "default auto"},
     },
     "required": ["src", "out_dir"],
     "additionalProperties": False},
    readOnlyHint=False,
    destructiveHint=True,
)
def render_office(src: str, out_dir: str, backend: str = "auto") -> dict:
    if backend not in BACKENDS:
        raise ToolError(f"backend must be one of {', '.join(BACKENDS)}, got {backend!r}")
    source = Path(src)
    if not source.is_file():
        raise ToolError(f"no such file: {src}")
    ext = source.suffix.lower()
    if ext not in (".pptx", ".docx"):
        raise ToolError(f"unsupported file type {source.suffix or '(none)'}: expected .pptx or .docx")
    deck = ext == ".pptx"

    # Every check happens before anything on disk is touched.
    check_path_string(out_dir)
    reason = _refusal(out_dir, src)
    if reason:
        raise ToolError(reason)

    if backend == "auto":
        chosen = next(b for b in ("powerpoint", "libreoffice", "text") if _available(b, ext))
    else:
        if not _available(backend, ext):
            raise ToolError(f"backend {backend} is not available on this machine (see render_backends)")
        chosen = backend
    if chosen == "powerpoint":
        busy = office_running()
        if busy:
            raise ToolError(f"refusing to render: {', '.join(busy)} is running. Office COM is single-instance, "
                            "so this would attach to the open window and quitting would close it. Close Office first.")
    if chosen == "text" and deck:
        _import_pptx()

    out_path = Path(out_dir)
    _empty_dir(out_path)
    copy = out_path / ("src" + ext)
    try:
        shutil.copy2(source, copy)
    except OSError as exc:
        raise ToolError(f"cannot copy {src} to {copy}: {exc}") from None
    _check_package(copy, "presentation" if deck else "Word document")

    warnings: list[str] = []
    sheets: list[str] = []
    if chosen == "text":
        count, files = _text_deck(copy, out_path) if deck else _text_docx(copy, out_path)
        kind = "text"
    else:
        try:
            count, files = _visual(chosen, copy, out_path, deck)
        except ToolError:
            raise
        except Exception as exc:  # COM and conversion failures surface as assorted exception types
            raise ToolError(f"{chosen} render failed: {type(exc).__name__}: {exc}") from None
        kind = "slides" if deck else "pages"
        pngs = [Path(f) for f in files]
        small = [p for p in pngs if p.stat().st_size < SMALL_PNG]
        if small:
            warnings.append(f"{len(small)} render(s) under 10 KB: " + ", ".join(p.name for p in small[:6]))
        if deck and _has_module("pptx"):
            files += _write_notes(copy, out_path)
        if pngs and _has_module("PIL"):
            sheets = _contact_sheets(pngs, out_path, 4, 3, 480)
    return {"backend": chosen, "kind": kind, "count": count, "files": files, "sheets": sheets,
            "warnings": warnings}


# ---------------------------------------------------------------------------------------------- CLI

def add_cli(groups) -> None:
    p = groups.add_parser("render", help="render PDFs and Office files to PNG, make contact sheets")
    cmds = p.add_subparsers(dest="command", required=True)

    c = cmds.add_parser("pdf", help="render PDF pages to page-NNN.png")
    c.add_argument("pdf")
    c.add_argument("-o", "--out-dir", dest="out_dir", required=True)
    c.add_argument("--dpi", type=float, default=DOC_DPI)
    c.add_argument("--first", type=int)
    c.add_argument("--last", type=int)
    common_flags(c)
    c.set_defaults(handler=_cli_pdf)

    c = cmds.add_parser("sheet", help="tile images into contact-NN.png sheets")
    c.add_argument("images", nargs="+")
    c.add_argument("-o", "--out-dir", dest="out_dir", required=True)
    c.add_argument("--cols", type=int, default=4)
    c.add_argument("--rows", type=int, default=3)
    c.add_argument("--thumb-width", dest="thumb_width", type=int, default=480)
    common_flags(c)
    c.set_defaults(handler=_cli_sheet)

    c = cmds.add_parser("office", help="render a .pptx or .docx (copy only) for review")
    c.add_argument("src")
    c.add_argument("-o", "--out-dir", dest="out_dir", required=True)
    c.add_argument("--backend", choices=BACKENDS, default="auto")
    common_flags(c)
    c.set_defaults(handler=_cli_office)

    c = cmds.add_parser("backends", help="show which render backends are available")
    common_flags(c)
    c.set_defaults(handler=_cli_backends)

    p = groups.add_parser("office", help="check for running Office applications")
    cmds = p.add_subparsers(dest="command", required=True)
    c = cmds.add_parser("check", help="list running Word/PowerPoint/Excel; exit 1 when any is running")
    c.add_argument("--wait", type=float, default=0, metavar="SECONDS",
                   help="poll every 2 s for up to SECONDS until nothing is running")
    common_flags(c)
    c.set_defaults(handler=_cli_office_check)


def _cli_office_check(args) -> CliResult:
    r = office_check(wait=args.wait)
    if r.get("error"):
        text = r["error"]
    else:
        text = "ok: no Office application is running" if r["ok"] else "running: " + ", ".join(r["running"])
    return CliResult(r, text, exit_code=0 if r["ok"] else 1)


def _cli_pdf(args) -> CliResult:
    r = render_pdf(args.pdf, args.out_dir, dpi=args.dpi, first=args.first, last=args.last)
    return CliResult(r, f"{len(r['files'])} of {r['pages']} page(s) rendered to {args.out_dir}")


def _cli_sheet(args) -> CliResult:
    r = render_contact_sheet(args.images, args.out_dir, cols=args.cols, rows=args.rows,
                             thumb_width=args.thumb_width)
    return CliResult(r, "\n".join(r["sheets"]))


def _cli_office(args) -> CliResult:
    r = render_office(args.src, args.out_dir, backend=args.backend)
    lines = [f"{r['backend']}: {r['count']} {r['kind']}, {len(r['files'])} file(s), "
             f"{len(r['sheets'])} contact sheet(s) in {args.out_dir}"]
    lines += [f"warning: {w}" for w in r["warnings"]]
    return CliResult(r, "\n".join(lines))


def _cli_backends(args) -> CliResult:
    r = render_backends()
    return CliResult(r, "\n".join(f"{k}: {v if v is not None else 'no'}" for k, v in r.items()))
