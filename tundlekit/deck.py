"""deck: presentations from a JSON spec (MANIFEST.md §3), after tundle's deckkit.py and checks/.

1 slide = stage-coloured eyebrow + takeaway title + a body that carries the point (bullets, lines, table,
figure, excerpt or chart) + source footer + slide number. Research slides add a DEMONSTRATED BY strip, design
slides a THUS strip, a phase's first slide an IN/OUT strip. Notes: TIME, SAY, IF ASKED.

Insert slides are optional slides numbered after the core slide they follow ("9a"); their notes open with an
INSERT line and their time is counted apart from the core time.

python-pptx and Pillow are imported inside the functions that need them; deck_lint on a spec needs neither.
"""
from __future__ import annotations

import json
import math
import pathlib
import re
import unicodedata

from tundlekit.palette import OUTCOME, PPT, STAGE, tint
from tundlekit.registry import ToolError, tool

SLIDE_TYPES = ("title", "divider", "content")
BODY_KINDS = ("bullets", "lines", "table", "figure", "excerpt", "chart", "point")
CHART_TYPES = ("bar", "line")
BLOCK_BOX = ("x", "y", "w", "h")             # optional block position and size, inches (§14.5)
INSERT_MODES = ("shown", "hidden", "off")
GEOM_MAX, GEOM_MIN = 1000, {"w": 0.01, "h": 0.01}   # §16.1: 0 <= x, y <= 1000 in; 0.01 <= w, h <= 1000 in
NO_HIGHLIGHT = ("pending",)                        # §16.7: white, refused as a highlight colour
DEFAULT_TARGET, DEFAULT_MAX, DEFAULT_WPS = "40:00", "45:00", 2.3
INSERT_LINE = "INSERT: optional slide; delete or hide it and the talk flows unchanged"

MAX_INSERTS = 26                    # letters a..z after 1 core slide

TIME_RE = re.compile(r"([0-9]{1,3}):([0-5][0-9])")
NUMBER_RE = re.compile(r"[0-9]+[a-z]?")
NOTES_TIME_RE = re.compile(r"^TIME[ \t]+([0-9]{1,3}):([0-9][0-9])(?![0-9])", re.M)
NOTES_SAY_RE = re.compile(r"^SAY:(.*?)(?=^IF ASKED:|^MUST HIT:|\Z)", re.M | re.S)


# ------------------------------------------------------------------ small helpers
def mmss(seconds: int) -> str:
    return f"{seconds // 60}:{seconds % 60:02d}"


def parse_time(text) -> int | None:
    """M:SS -> seconds, or None when the text is not in that format."""
    m = TIME_RE.fullmatch(text) if isinstance(text, str) else None
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def budget_status(seconds: int, target: int, maximum: int) -> str:
    if seconds > maximum:
        return "OVER BUDGET"
    if seconds > target:
        return "over target"
    return "ok"


def _is_num(v) -> bool:
    """A JSON number (not a bool) that converts to a finite float; 10**400 does not."""
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return False
    try:
        return math.isfinite(float(v))
    except OverflowError:
        return False


def _short(value) -> str:
    """A bounded repr for problem messages (never fails, even for huge ints)."""
    try:
        text = repr(value)
    except ValueError:
        text = f"<{type(value).__name__}>"
    return text if len(text) <= 40 else text[:37] + "..."


def bad_chars(text: str) -> bool:
    """Control characters other than newline and tab (and lone surrogates) can't be written to PPTX or SVG."""
    return any((unicodedata.category(ch) in ("Cc", "Cs") and ch not in "\n\t") or ch in "￾￿"
               for ch in text)


def _check_text(value, where: str, problems: list, required: bool = False) -> None:
    if value is None and not required:
        return
    if not isinstance(value, str):
        problems.append(f"{where}: must be a string")
    elif bad_chars(value):
        problems.append(f"{where}: contains a control character")


def _is_file(path: pathlib.Path) -> bool:
    try:
        return path.is_file()
    except (OSError, ValueError):
        return False


def _words(text: str) -> int:
    return len(text.split())


def _display(path) -> str:
    return str(path).replace("\\", "/")


def _require_office(pillow: bool = True):
    try:
        import pptx  # noqa: F401
    except ImportError:
        raise ToolError("python-pptx is not installed (pip install python-pptx, or tundlekit[office])") from None
    if pillow:
        try:
            import PIL  # noqa: F401
        except ImportError:
            raise ToolError("Pillow is not installed (pip install Pillow, or tundlekit[office])") from None


# ------------------------------------------------------------------ spec loading and validation (§3.1)
def read_json(path: pathlib.Path, what: str):
    """Parse a JSON file (a UTF-8 BOM is allowed); any failure is a ToolError."""
    try:
        text = path.read_text(encoding="utf-8-sig")
    except FileNotFoundError:
        raise ToolError(f"{what} not found: {path}") from None
    except (OSError, ValueError) as e:
        raise ToolError(f"cannot read {what} {path}: {e}") from None
    try:
        return json.loads(text)
    except (ValueError, RecursionError) as e:
        raise ToolError(f"{what} {path} is not valid JSON: {e}") from None


def load_spec(spec, spec_path) -> tuple[dict, pathlib.Path, str]:
    """Return (spec, base directory for relative paths, display path for findings)."""
    if (spec is None) == (spec_path is None):
        raise ToolError("give exactly 1 of spec or spec_path")
    if spec_path is None:
        base, shown = pathlib.Path.cwd(), "<spec>"
    else:
        p = pathlib.Path(spec_path)
        spec = read_json(p, "spec file")
        base, shown = p.resolve().parent, _display(spec_path)
    if not isinstance(spec, dict):
        raise ToolError("spec: must be a JSON object")
    return spec, base, shown


def _resolve(base: pathlib.Path, path: str) -> pathlib.Path:
    p = pathlib.Path(path)
    return p if p.is_absolute() else base / p


def _is_image(path: pathlib.Path) -> bool:
    try:
        with open(path, "rb") as f:
            head = f.read(8)
    except OSError:
        return False
    return head.startswith(b"\x89PNG\r\n\x1a\n") or head.startswith(b"\xff\xd8\xff")


class Plan:
    """A validated spec with every slide's number, timing and text worked out."""

    def __init__(self, spec: dict, base: pathlib.Path, shown: str):
        self.spec, self.base, self.shown = spec, base, shown
        problems: list[str] = []
        self.meta = self._meta(spec.get("meta", {}), problems)
        slides = spec.get("slides")
        if not isinstance(slides, list) or not slides:
            problems.append("slides: must be a non-empty list")
            slides = []
        self.slides = []
        seen_core, run = False, 0
        for i, s in enumerate(slides):
            where = f"slides[{i}]"
            if not isinstance(s, dict):
                problems.append(f"{where}: must be an object")
                continue
            sl = self._slide(s, where, seen_core, problems)
            self.slides.append(sl)
            if sl["insert"]:
                run += 1
                if run == MAX_INSERTS + 1:
                    problems.append(f"{where}.insert: more than {MAX_INSERTS} insert slides after 1 core slide")
            else:
                seen_core, run = True, 0
        if problems:
            raise ToolError("invalid deck spec:\n" + "\n".join(f"- {p}" for p in problems))
        self._number()

    # ------------------------------------------------------------ meta
    def _meta(self, meta, problems) -> dict:
        if not isinstance(meta, dict):
            problems.append("meta: must be an object")
            meta = {}
        out = {"id": meta.get("id"), "title": meta.get("title"), "logo": None,
               "inserts": meta.get("inserts", "shown"), "words_per_second": meta.get("words_per_second", DEFAULT_WPS)}
        if out["id"] is not None and (not isinstance(out["id"], str) or not out["id"] or bad_chars(out["id"])):
            problems.append("meta.id: must be a non-empty string without control characters")
        _check_text(out["title"], "meta.title", problems)
        if not isinstance(out["inserts"], str) or out["inserts"] not in INSERT_MODES:
            problems.append(f"meta.inserts: must be one of shown, hidden, off, not {_short(out['inserts'])}")
        if not _is_num(out["words_per_second"]) or out["words_per_second"] <= 0:
            problems.append("meta.words_per_second: must be a positive finite number")
        logo = meta.get("logo")
        if logo is not None:
            if not isinstance(logo, str) or not logo:
                problems.append("meta.logo: must be a path string")
            else:
                p = _resolve(self.base, logo)
                out["logo"] = p if _is_file(p) else None      # a missing logo is ignored (§3.1a)
        budget = meta.get("budget", {})
        if not isinstance(budget, dict):
            problems.append("meta.budget: must be an object")
            budget = {}
        for key, default in (("target", DEFAULT_TARGET), ("max", DEFAULT_MAX)):
            value = budget.get(key, default)
            if parse_time(value) is None:
                problems.append(f"meta.budget.{key}: must be M:SS, not {_short(value)}")
                value = default
            out[key] = value
        return out

    # ------------------------------------------------------------ slides
    def _slide(self, s: dict, where: str, seen_core: bool, problems: list) -> dict:
        kind = s.get("type", "content")
        if not isinstance(kind, str) or kind not in SLIDE_TYPES:
            problems.append(f"{where}.type: unknown slide type {_short(kind)} (one of title, divider, content)")
            kind = "content"
        title = s.get("title")
        if not isinstance(title, str) or not title.strip():
            problems.append(f"{where}.title: required non-empty string")
            title = ""
        else:
            _check_text(title, f"{where}.title", problems)
        for key in ("subtitle", "byline", "eyebrow", "source", "thus"):
            _check_text(s.get(key), f"{where}.{key}", problems)
        stage = s.get("stage")
        if stage is not None and (not isinstance(stage, str) or stage not in STAGE):
            problems.append(f"{where}.stage: unknown stage {_short(stage)} (one of {', '.join(STAGE)})")
        insert = s.get("insert", False)
        if not isinstance(insert, bool):
            problems.append(f"{where}.insert: must be true or false")
            insert = False
        if insert and kind != "content":
            problems.append(f"{where}.insert: only content slides can be insert slides")
            insert = False
        elif insert and not seen_core:
            problems.append(f"{where}.insert: an insert slide must follow at least 1 core slide")
        demo = s.get("demonstrated_by")
        if demo is not None:
            if not isinstance(demo, dict):
                problems.append(f"{where}.demonstrated_by: must be an object with string paper and setup")
            else:
                _check_text(demo.get("paper"), f"{where}.demonstrated_by.paper", problems, required=True)
                _check_text(demo.get("setup"), f"{where}.demonstrated_by.setup", problems)
        phase = s.get("phase")
        if phase is not None:
            if not isinstance(phase, dict):
                problems.append(f"{where}.phase: must be an object with string in and out")
            else:
                _check_text(phase.get("in"), f"{where}.phase.in", problems, required=True)
                _check_text(phase.get("out"), f"{where}.phase.out", problems, required=True)
        body = s.get("body")
        if isinstance(body, list):                          # §14.5: a list of blocks
            for i, block in enumerate(body):
                self._body(block, f"{where}.body[{i}]", problems)
        elif body is not None:
            self._body(body, f"{where}.body", problems)
        seconds, say, asked = self._notes(s.get("notes"), f"{where}.notes", problems)
        if not asked and s.get("asked") is not None:      # tolerated: asked beside notes instead of inside
            asked = self._asked(s["asked"], f"{where}.asked", problems)
        return {"spec": s, "type": kind, "title": title,
                "insert": insert, "seconds": seconds, "say": say, "asked": asked}

    def _body(self, body, where: str, problems: list) -> None:
        if not isinstance(body, dict):
            problems.append(f"{where}: must be an object or a list of objects")
            return
        kind = body.get("kind")
        if not isinstance(kind, str) or kind not in BODY_KINDS:
            problems.append(f"{where}.kind: unknown body kind {_short(kind)} (one of {', '.join(BODY_KINDS)})")
            return
        for key in BLOCK_BOX:
            v = body.get(key)
            if v is None:
                continue
            if not _is_num(v):
                problems.append(f"{where}.{key}: must be a finite number (inches)")
            elif not GEOM_MIN.get(key, 0) <= v <= GEOM_MAX:      # §16.1
                problems.append(f"{where}.{key}: must be from {GEOM_MIN.get(key, 0):g} to {GEOM_MAX:g} inches, "
                                f"not {_short(v)}")
        size = body.get("size")
        if size is not None and (not _is_num(size) or not 1 <= size <= 400):
            problems.append(f"{where}.size: must be a number of points from 1 to 400")
        if kind == "point":
            _check_text(body.get("text"), f"{where}.text", problems, required=True)
        elif kind in ("bullets", "lines"):
            self._texts(body.get("items"), f"{where}.items", problems)
            if kind == "lines" and not isinstance(body.get("mono", False), bool):
                problems.append(f"{where}.mono: must be true or false")
        elif kind == "table":
            rows = body.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(r, list) and r for r in rows):
                problems.append(f"{where}.rows: must be a non-empty list of non-empty rows")
                return
            if len({len(r) for r in rows}) != 1:
                problems.append(f"{where}.rows: rows have different lengths "
                                f"({', '.join(str(len(r)) for r in rows)})")
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    cw = f"{where}.rows[{i}][{j}]"
                    if isinstance(cell, str):
                        _check_text(cell, cw, problems)
                    elif cell is not None and not isinstance(cell, bool) and not _is_num(cell):
                        problems.append(f"{cw}: must be a string, a finite number, a boolean or null")
            widths = body.get("widths")
            if widths is not None:
                if not isinstance(widths, list) or not all(_is_num(w) and w > 0 for w in widths):
                    problems.append(f"{where}.widths: must be a list of positive numbers")
                elif len(widths) != len(rows[0]):
                    problems.append(f"{where}.widths: {len(widths)} widths for {len(rows[0])} columns")
            mono = body.get("mono_cols")
            if mono is not None:
                ncols = len(rows[0])
                if not isinstance(mono, list) or not all(isinstance(c, int) and not isinstance(c, bool)
                                                         for c in mono):
                    problems.append(f"{where}.mono_cols: must be a list of 0-based column indices")
                else:
                    bad = [c for c in mono if not 0 <= c < ncols]
                    if bad:
                        problems.append(f"{where}.mono_cols: {_short(bad)} out of range for {ncols} columns "
                                        f"(0-based, 0..{ncols - 1})")
        elif kind == "figure":
            path = body.get("path")
            if not isinstance(path, str) or not path:
                problems.append(f"{where}.path: required path to a PNG or JPEG file")
                return
            p = _resolve(self.base, path)
            if not _is_file(p):
                problems.append(f"{where}.path: figure file not found: {path}")
            elif not _is_image(p):
                problems.append(f"{where}.path: not a PNG or JPEG file: {path}")
        elif kind == "excerpt":
            _check_text(body.get("text"), f"{where}.text", problems, required=True)
            _check_text(body.get("caption"), f"{where}.caption", problems)
        elif kind == "chart":
            cats = body.get("categories")
            ok = self._texts(cats, f"{where}.categories", problems, non_empty=True)
            self._chart_values(body, where, cats if ok else None, problems)
            hl = body.get("highlight")
            if hl is not None and ok:
                if isinstance(hl, bool) or not isinstance(hl, (int, str)):
                    problems.append(f"{where}.highlight: must be a category index or name")
                elif isinstance(hl, int) and not 0 <= hl < len(cats):
                    problems.append(f"{where}.highlight: index {_short(hl)} is out of range")
                elif isinstance(hl, str) and hl not in cats:
                    problems.append(f"{where}.highlight: {_short(hl)} is not one of the categories")
            for key in ("unit", "takeaway"):
                _check_text(body.get(key), f"{where}.{key}", problems)
            ctype = body.get("chart_type", "bar")
            if not isinstance(ctype, str) or ctype not in CHART_TYPES:
                problems.append(f"{where}.chart_type: must be bar or line, not {_short(ctype)}")
            hc = body.get("highlight_color")
            if hc is not None and (not isinstance(hc, str) or (hc not in STAGE and hc not in OUTCOME)):
                problems.append(f"{where}.highlight_color: unknown colour key {_short(hc)} "
                                f"(a stage: {', '.join(STAGE)}; or an outcome: "
                                f"{', '.join(k for k in OUTCOME if k not in NO_HIGHLIGHT)})")
            elif hc in NO_HIGHLIGHT:
                problems.append(f"{where}.highlight_color: {hc!r} is white and can't be a highlight colour")

    @classmethod
    def _chart_values(cls, body: dict, where: str, cats, problems: list) -> None:
        """Exactly 1 of values (1 series) or series (a list of {name, values}), 1 finite number per category."""
        vals, series = body.get("values"), body.get("series")
        if (vals is None) == (series is None):
            problems.append(f"{where}: give exactly 1 of values or series")
            return
        if vals is not None:
            if not isinstance(vals, list) or not all(_is_num(v) for v in vals):
                problems.append(f"{where}.values: must be a list of finite numbers")
            elif cats is not None and len(cats) != len(vals):
                problems.append(f"{where}.values: {len(vals)} values for {len(cats)} categories")
            return
        if not isinstance(series, list) or not series:
            problems.append(f"{where}.series: must be a non-empty list of {{name, values}} objects")
            return
        for k, sr in enumerate(series):
            sw = f"{where}.series[{k}]"
            if not isinstance(sr, dict):
                problems.append(f"{sw}: must be an object with a name and values")
                continue
            name = sr.get("name")
            if not isinstance(name, str) or not name.strip():
                problems.append(f"{sw}.name: required non-empty string")
            else:
                _check_text(name, f"{sw}.name", problems)
            sv = sr.get("values")
            if not isinstance(sv, list) or not all(_is_num(v) for v in sv):
                problems.append(f"{sw}.values: must be a list of finite numbers")
            elif cats is not None and len(sv) != len(cats):
                problems.append(f"{sw}.values: {len(sv)} values for {len(cats)} categories")

    @staticmethod
    def _texts(items, where: str, problems: list, non_empty: bool = False) -> bool:
        """A list of strings without control characters; returns whether it is usable."""
        if not isinstance(items, list) or not all(isinstance(t, str) for t in items) or (non_empty and not items):
            problems.append(f"{where}: must be a {'non-empty ' if non_empty else ''}list of strings")
            return False
        for i, t in enumerate(items):
            _check_text(t, f"{where}[{i}]", problems)
        return True

    def _notes(self, notes, where: str, problems: list):
        if not isinstance(notes, dict):
            problems.append(f"{where}: required object with a time")
            return 0, "", []
        time = notes.get("time")
        seconds = parse_time(time)
        if time is None:
            problems.append(f"{where}.time: required (M:SS)")
        elif seconds is None:
            problems.append(f"{where}.time: must be M:SS (1-3 digit minutes, seconds 00-59), not {_short(time)}")
        say = notes.get("say") or ""
        if not isinstance(say, str):
            problems.append(f"{where}.say: must be a string")
            say = ""
        else:
            _check_text(say, f"{where}.say", problems)
        return seconds or 0, say, self._asked(notes.get("asked"), f"{where}.asked", problems)

    @classmethod
    def _asked(cls, asked, where: str, problems: list) -> list:
        asked = asked or []
        return asked if cls._texts(asked, where, problems) else []

    # ------------------------------------------------------------ derived values
    def _number(self) -> None:
        core, letter = 0, 0
        for pos, sl in enumerate(self.slides, 1):
            sl["position"] = pos
            if sl["insert"]:
                letter += 1
                sl["number"] = f"{core}{chr(96 + letter)}"
            else:
                core += 1
                letter = 0
                sl["number"] = str(core) if sl["type"] == "content" else None
            sl["label"] = sl["number"] or str(pos)
            sl["notes_text"] = notes_text(sl)
            sl["face"] = face_text(sl["spec"], sl["type"])
            sl["face_body"] = face_text(sl["spec"], sl["type"], with_title=False)

    @property
    def core_seconds(self) -> int:
        return sum(s["seconds"] for s in self.slides if not s["insert"])

    @property
    def insert_seconds(self) -> int:
        return sum(s["seconds"] for s in self.slides if s["insert"])


def notes_text(sl: dict) -> str:
    """The speaker notes of a slide in the §3.2 format."""
    parts = [INSERT_LINE] if sl["insert"] else []
    parts.append(f"TIME {sl['spec']['notes']['time']}")
    if sl["say"]:
        parts.append(f"SAY: {sl['say']}")
    if sl["asked"]:
        parts.append("IF ASKED:")
        parts.extend(f"- {a}" for a in sl["asked"])
    return "\n".join(parts)


def _strip_texts(s: dict) -> list[str]:
    out = []
    demo = s.get("demonstrated_by")
    if isinstance(demo, dict):
        out.append(f"DEMONSTRATED BY   {demo.get('paper', '')}  ·  {demo.get('setup', '')}")
    if isinstance(s.get("phase"), dict):
        out.append(f"IN   {s['phase'].get('in', '')}        OUT   {s['phase'].get('out', '')}")
    if s.get("thus"):
        out.append(f"THUS   {s['thus']}")
    return out


def body_blocks(body) -> list[dict]:
    """A slide body as a list of blocks: a single block (§3.1) or a list of blocks (§14.5)."""
    if isinstance(body, dict):
        return [body]
    if isinstance(body, list):
        return [b for b in body if isinstance(b, dict)]
    return []


def _point_text(text: str) -> str:
    """A point block is 1 line: line breaks become spaces."""
    return " ".join(text.splitlines())


def face_text(s: dict, kind: str, with_title: bool = True) -> str:
    """All text a spec slide shows (§3.3): title, eyebrow, subtitle, byline, body, strips, cells, takeaway."""
    out = [s.get("title") or ""] if with_title else []
    for key in ("eyebrow", "subtitle", "byline"):
        if s.get(key):
            out.append(s[key].upper() if key == "eyebrow" else s[key])
    if kind == "content":
        if s.get("source"):
            out.append(s["source"])
        out.extend(_strip_texts(s))
        for body in body_blocks(s.get("body")):
            k = body.get("kind")
            if k in ("bullets", "lines"):
                out.extend(body.get("items", []))
            elif k == "table":
                out.extend(str(c) for r in body.get("rows", []) for c in r)
            elif k == "excerpt":
                out.append(body.get("text", ""))
                if body.get("caption"):
                    out.append(body["caption"])
            elif k == "point":
                out.append(_point_text(body.get("text", "")))
            elif k == "chart":
                out.extend(body.get("categories", []))
                out.extend(sr["name"] for sr in body.get("series") or [] if isinstance(sr, dict) and sr.get("name"))
                if body.get("takeaway"):
                    out.append(body["takeaway"])
    return "\n".join(t for t in out if t)


# ------------------------------------------------------------------ building (§3.2)
W, H = 13.333, 7.5
INK, INK2, MUTED, HAIRC, PALEC, CODEBG, GREYBAR = "000000", "404040", "767676", "BFBFBF", "EDEDED", "F4F4F4", "9A9A9A"
FONT, MONO = "Calibri", "Consolas"
LOGO_W = 0.42
BODY_X, BODY_W = 0.75, 11.8
FIG_W, FIG_H = 11.8, 5.3
BODY_BOTTOM, BODY_BOTTOM_THUS = 6.9, 6.25    # the body area ends at 6.9 in, or above the THUS strip (6.38 in)
GAP = 0.15                                   # between stacked body blocks
POINT_SIZE, EXCERPT_SIZE = 20, 15
MIN_BLOCK_H, MIN_CHART_H = 0.3, 1.0
CAPTION_H = 0.45                             # an excerpt caption line under the box
SERIES_GREYS = ("9A9A9A", "4D4D4D", "C8C8C8", "737373", "262626")
OVERLAP_EPS = 0.01                           # in; boxes that only touch do not overlap


class _Builder:
    """The deckkit slide kit, driven by a Plan instead of by build scripts."""

    def __init__(self, plan: Plan, inserts: str):
        from pptx import Presentation
        from pptx.util import Inches

        self.plan, self.inserts = plan, inserts
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []
        self.label = ""
        self.strips: list[tuple[str, tuple]] = []     # (name, (x, y, w, h)) of the current slide's strips

    @staticmethod
    def overlaps(a, b) -> bool:
        """Whether 2 (x, y, w, h) boxes overlap by more than OVERLAP_EPS in both directions."""
        return (min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]) > OVERLAP_EPS
                and min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]) > OVERLAP_EPS)

    # ------------------------------------------------------------ primitives
    @staticmethod
    def rgb(h):
        from pptx.dml.color import RGBColor

        return RGBColor.from_string(h)

    def warn(self, text: str) -> None:
        self.warnings.append(f"slide {self.label}: {text}")

    @staticmethod
    def text_need(w, paras, spacing) -> float:
        """deckkit's estimate of the height (in) that paras of (text, size, mono) need in a box w in wide."""
        need = 0.0
        for text, size, mono in paras:
            cpl = max(1, int((w - 0.1) * 72 / (size * (0.6 if mono else 0.5))))
            lines = sum(max(1, math.ceil(len(seg) / cpl)) for seg in text.split("\n"))
            need += lines * size * 1.2 / 72 + size * spacing / 72
        return need

    def check_text(self, y, w, h, paras, spacing, body=True) -> None:
        """deckkit's overflow heuristic: warn when the estimated text height exceeds its box or the slide."""
        need = self.text_need(w, paras, spacing)
        if need > h + 0.08 or (body and y + need > H - 0.45):
            first = paras[0][0] if paras else ""
            self.warn(f"text needs {need:.2f} in, box {h:.2f} in: {first[:48]!r}")

    def tb(self, s, x, y, w, h, paras, align=None, anchor=None, spacing=0.35, body=True):
        """A text box; paras are (text, size, bold, colour[, font])."""
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt

        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        f = t.text_frame
        f.word_wrap = True
        f.vertical_anchor = anchor if anchor is not None else MSO_ANCHOR.TOP
        f.margin_left = f.margin_right = Inches(0.04)
        f.margin_top = f.margin_bottom = Inches(0.02)
        for i, p in enumerate(paras):
            text, size, bold, color = p[:4]
            font = p[4] if len(p) > 4 else FONT
            para = f.paragraphs[0] if i == 0 else f.add_paragraph()
            para.alignment = align if align is not None else PP_ALIGN.LEFT
            para.space_after = Pt(size * spacing)
            r = para.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = self.rgb(color)
            r.font.name = font
        self.check_text(y, w, h, [(p[0], p[1], len(p) > 4 and p[4] == MONO) for p in paras], spacing, body)
        return t

    def rect(self, s, x, y, w, h, fill, line=None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid()
        b.fill.fore_color.rgb = self.rgb(fill)
        if line:
            b.line.color.rgb = self.rgb(line)
            b.line.width = Pt(0.75)
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        return b

    def logo(self, s, x=None, y=0.24, w=LOGO_W):
        from pptx.util import Inches

        path = self.plan.meta["logo"]
        if path is None:
            return None
        x = W - 0.6 - w if x is None else x
        try:
            return s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        except Exception as e:  # an unreadable image is a user problem, not a crash
            raise ToolError(f"cannot use logo {path}: {e}") from None

    def strip(self, s, y, runs, colour, fill, h, text_len, size):
        """A full-width strip with a coloured left edge; runs are (text, bold, colour, size)."""
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        b = self.rect(s, 0.75, y, W - 1.5, h, fill)
        self.rect(s, 0.75, y, 0.09, h, colour)
        name = runs[0][0].strip()
        self.strips.append(("IN/OUT" if name == "IN" else name, (0.75, y, W - 1.5, h)))
        f = b.text_frame
        f.word_wrap = True
        f.vertical_anchor = MSO_ANCHOR.MIDDLE
        f.margin_left = Inches(0.25)
        p = f.paragraphs[0]
        for text, bold, col, sz in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.color.rgb = self.rgb(col)
            r.font.name = FONT
        cpl = int((W - 1.9) * 72 / (size * 0.5))
        if text_len > cpl * max(1, round(h / 0.45)):
            self.warn(f"{runs[0][0].strip()} strip too long for 1 line ({text_len} chars)")
        return b

    # ------------------------------------------------------------ slides
    def build(self):
        for sl in self.plan.slides:
            if sl["insert"] and self.inserts == "off":
                continue
            self.label = sl["label"]
            s = self.prs.slides.add_slide(self.blank)
            if sl["type"] == "title":
                self.title_slide(s, sl["spec"])
            elif sl["type"] == "divider":
                self.divider(s, sl["spec"])
            else:
                self.content(s, sl)
            if sl["insert"] and self.inserts == "hidden":
                s._element.set("show", "0")
            s.notes_slide.notes_text_frame.text = sl["notes_text"]
        return self.prs

    def title_slide(self, s, spec):
        self.rect(s, 0.95, 2.4, 1.4, 0.08, INK)
        self.tb(s, 0.9, 2.65, 11.5, 1.0, [(spec["title"], 44, True, INK)])
        if spec.get("subtitle"):
            self.tb(s, 0.9, 3.75, 11.5, 1.5, [(spec["subtitle"], 24, False, INK2)])
        if spec.get("byline"):
            self.tb(s, 0.9, 5.9, 9, 0.5, [(spec["byline"], 16, False, MUTED)])
        self.logo(s, x=0.9, y=0.9, w=0.9)

    def divider(self, s, spec):
        self.rect(s, 0, 0, W, H, "141414")
        self.rect(s, 0.95, 2.55, 1.2, 0.06, "FFFFFF")
        self.tb(s, 0.9, 2.8, 11.5, 1.0, [(spec["title"], 40, True, "FFFFFF")])
        if spec.get("subtitle"):
            self.tb(s, 0.9, 3.8, 11.5, 0.9, [(spec["subtitle"], 24, False, "BDBDBD")])

    def content(self, s, sl):
        from pptx.enum.text import PP_ALIGN

        spec = sl["spec"]
        stage = spec.get("stage")
        col = PPT.get(stage, MUTED)
        if stage:
            self.rect(s, 0.6, 0.36, 0.22, 0.22, col)
        eyebrow = (spec.get("eyebrow") or "").upper()
        if eyebrow:
            self.tb(s, 0.9 if stage else 0.6, 0.3, 10.5, 0.35, [(eyebrow, 11, True, col)])
        self.tb(s, 0.58, 0.58, W - 1.2, 0.75, [(spec["title"], 30, True, INK)])
        self.rect(s, 0.6, 1.38, W - 1.2, 0.02 if stage else 0.012, col if stage else HAIRC)
        if spec.get("source"):
            self.tb(s, 0.6, 7.02, 11.2, 0.3, [(spec["source"], 10, False, MUTED)], body=False)
        self.logo(s)
        self.tb(s, W - 1.2, 7.02, 0.6, 0.3, [(sl["number"], 10, False, MUTED)], align=PP_ALIGN.RIGHT, body=False)

        top, bottom = 1.6, BODY_BOTTOM
        self.strips = []
        demo = spec.get("demonstrated_by")
        if demo:
            text = f"{demo['paper']}  ·  {demo.get('setup', '')}"
            self.strip(s, top - 0.08, [("DEMONSTRATED BY   ", True, "555555", 12), (text, False, INK, 15)],
                       "555555", "F2F2F2", 0.5, len("DEMONSTRATED BY") + 3 + len(text), 15)
            top += 0.58
        phase = spec.get("phase")
        if phase:
            edge = PPT[stage] if stage else PALEC
            runs = [("IN   ", True, MUTED, 12), (phase["in"], False, INK, 15),
                    ("        OUT   ", True, MUTED, 12), (phase["out"], False, INK, 15)]
            self.strip(s, top - 0.08, runs, edge, PALEC, 0.46, 19 + len(phase["in"]) + len(phase["out"]), 15)
            top += 0.54
        if spec.get("thus"):
            thus_col = PPT.get(stage, INK)
            fill = tint(thus_col, 0.12).lstrip("#").upper()
            self.strip(s, 6.38, [("THUS   ", True, thus_col, 12), (spec["thus"], False, INK, 17)],
                       thus_col, fill, 0.56, 7 + len(spec["thus"]), 17)
            bottom = BODY_BOTTOM_THUS
        blocks = body_blocks(spec.get("body"))
        if blocks:
            self.layout(s, blocks, top, bottom, stage)

    # ------------------------------------------------------------ bodies (§3.2, §14.5)
    def natural_h(self, b, w) -> float | None:
        """Estimated height (in) of a block with no h; None for charts and figures, which take what is left."""
        kind = b["kind"]
        if kind in ("bullets", "lines", "point"):
            _, paras, spacing = self.block_paras(b)
            need = self.text_need(w, [(p[0], p[1], len(p) > 4 and p[4] == MONO) for p in paras], spacing)
            return need + 0.1
        if kind == "table":
            size, head, row_h = self.table_sizes(b)
            return self.table_est(b["rows"], b.get("widths"), w, size, head, row_h, set(b.get("mono_cols") or ()))
        if kind == "excerpt":
            size = b.get("size") or EXCERPT_SIZE
            lines = b["text"].rstrip("\n").split("\n")
            need = self.text_need(w - 0.36, [(line, size, True) for line in lines], 0.0)
            return need + 0.34 + (CAPTION_H if b.get("caption") else 0)
        return None

    @staticmethod
    def block_paras(b):
        """(size, paragraphs for tb, paragraph spacing) of a text block."""
        kind = b["kind"]
        if kind == "bullets":
            size = b.get("size") or 24
            return size, [(f"•  {t}", size, False, INK) for t in b["items"]], 0.6
        if kind == "lines":
            size = b.get("size") or 22
            font = MONO if b.get("mono") else FONT
            return size, [(t, size, False, INK, font) for t in b["items"]], 0.5
        size = b.get("size") or POINT_SIZE
        return size, [(_point_text(b["text"]), size, True, INK)], 0.0

    @staticmethod
    def table_sizes(b):
        size = b.get("size")
        if size is None:
            return 17, 15, 0.62
        return size, round(size * 15 / 17, 1), max(0.3, 0.62 * size / 17)

    def layout(self, s, blocks, top, bottom, stage):
        """Place the blocks: positioned ones where they say, the others stacked top to bottom (§14.5).

        Text blocks and tables take their estimated height; charts and figures share what is left above
        `bottom` after the stacked blocks below them. The last stacked text block runs to `bottom`, as a single
        body always did. A figure without h shrinks to fit, so it never runs past the body area."""
        start = top + 0.2
        cursor = start
        stacked = [i for i, b in enumerate(blocks) if b.get("y") is None]
        placed, flowed = [], []                 # (index, kind, (x, y, w, h)) of positioned and stacked blocks
        for i, b in enumerate(blocks):
            kind = b["kind"]
            x = b["x"] if b.get("x") is not None else BODY_X
            if b.get("w") is not None:
                w = b["w"]
            elif b.get("x") is not None:
                w = max(0.5, min(BODY_X + BODY_W, W - 0.6) - x)
            else:
                w = BODY_W
            positioned = b.get("y") is not None
            later = [blocks[j] for j in stacked if j > i] if not positioned else []
            if positioned:
                y = b["y"]
            elif kind == "figure" and cursor == start:
                y = top                        # a figure at the head of the body starts right under the strips
            else:
                y = cursor
            reserve, flex = 0.0, 1
            for r in later:
                rw = r["w"] if r.get("w") is not None else BODY_W
                nat = r["h"] if r.get("h") is not None else self.natural_h(r, rw)
                if nat is None:
                    flex += 1
                    reserve += GAP
                else:
                    reserve += nat + GAP
            room = bottom - y - reserve
            h = b.get("h")
            if kind in ("bullets", "lines", "point"):
                if h is None:
                    nat = self.natural_h(b, w)
                    h = room if not later and not positioned else min(nat, room)
                    h = max(h, MIN_BLOCK_H)
                size, paras, spacing = self.block_paras(b)
                self.tb(s, x, y, w, h, paras, spacing=spacing)
                used = h
            elif kind == "table":
                size, head, row_h = self.table_sizes(b)
                if h is not None:
                    row_h = h / len(b["rows"])
                used = self.table(s, b["rows"], y, b.get("widths"), x=x, w=w, size=size, head=head, row_h=row_h,
                                  box_h=h, mono_cols=set(b.get("mono_cols") or ()))
            elif kind == "excerpt":
                caption = b.get("caption")
                if h is None:
                    nat = self.natural_h(b, w)
                    h = room if not later and not positioned else min(nat, room)
                    h = max(h, MIN_BLOCK_H + (CAPTION_H if caption else 0))
                if caption and h - CAPTION_H < MIN_BLOCK_H:        # §16.1: never a negative or tiny box
                    self.warn(f"excerpt box {h:.2f} in tall has no room for its caption; caption omitted")
                    caption = None
                self.excerpt(s, b["text"], y, h - (CAPTION_H if caption else 0), caption, x=x, w=w,
                             size=b.get("size") or EXCERPT_SIZE)
                used = h
            elif kind == "chart":
                if h is None:
                    h = max(room / flex, MIN_CHART_H)
                self.chart(s, b, x, y, w, h, stage)
                used = h
            else:                                   # figure
                limit = h if h is not None else max(room / flex, MIN_BLOCK_H)
                used = self.figure(s, _resolve(self.plan.base, b["path"]), y, bottom, x=x, max_w=w,
                                   max_h=limit, fixed_h=h is not None)
            cursor = max(cursor, y + used + GAP)
            (placed if positioned else flowed).append((i, kind, (x, y, w, used)))
        self.check_overlaps(placed, flowed)

    def check_overlaps(self, placed, flowed) -> None:
        """§16.7: warn when a positioned block overlaps a strip, or a stacked block overlaps a positioned one."""
        for i, kind, box in placed:
            for name, strip in self.strips:
                if self.overlaps(box, strip):
                    self.warn(f"{kind} block (body[{i}]) at y={box[1]:.2f} in overlaps the {name} strip")
        for i, kind, box in flowed:
            for k, other_kind, other in placed:
                if self.overlaps(box, other):
                    self.warn(f"stacked {kind} block (body[{i}]) at y={box[1]:.2f} in overlaps the positioned "
                              f"{other_kind} block (body[{k}]) at y={other[1]:.2f} in")

    def table_est(self, rows, widths, w, size, head, row_h, mono_cols=frozenset()) -> float:
        nc = len(rows[0])
        tot = sum(widths) if widths else nc
        colw = [w * (widths[j] if widths else 1) / tot for j in range(nc)]
        est = 0.0
        for i, row in enumerate(rows):
            row_need = row_h
            for j, cell in enumerate(row):
                sz = head if i == 0 else size
                cpl = max(1, int((colw[j] - 0.16) * 72 / (sz * (0.6 if j in mono_cols else 0.5))))
                row_need = max(row_need, math.ceil(max(1, len(str(cell))) / cpl) * sz * 1.2 / 72 + 0.12)
            est += row_need
        return est

    def table(self, s, rows, y, widths, x=BODY_X, w=BODY_W, size=17, head=15, row_h=0.62, box_h=None,
              mono_cols=frozenset()):
        """A native table; returns its estimated height (in). Columns in mono_cols (the header row too) use the
        monospace font (§16.7)."""
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        nr, nc = len(rows), len(rows[0])
        shp = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
        tot = sum(widths) if widths else nc
        colw = [w * (widths[j] if widths else 1) / tot for j in range(nc)]
        for j in range(nc):
            shp.columns[j].width = Inches(colw[j])
        shp.first_row = True
        for i, row in enumerate(rows):
            shp.rows[i].height = Inches(row_h)
            for j, cell in enumerate(row):
                c = shp.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = self.rgb(PALEC if i == 0 else "FFFFFF")
                c.margin_left = c.margin_right = Inches(0.08)
                c.margin_top = c.margin_bottom = Inches(0.03)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame
                tf.word_wrap = True
                text = str(cell)
                r = tf.paragraphs[0].add_run()
                r.text = text
                sz = head if i == 0 else size
                r.font.size = Pt(sz)
                r.font.name = MONO if j in mono_cols else FONT
                r.font.bold = i == 0 or (j == 0 and nc > 1)
                r.font.color.rgb = self.rgb(INK)
        est = self.table_est(rows, widths, w, size, head, row_h, mono_cols)
        if y + est > H - 0.4:
            self.warn(f"table needs {est:.2f} in from y={y:.2f} and runs past the slide bottom")
        elif box_h is not None and est > box_h + 0.08:
            self.warn(f"table needs {est:.2f} in, box {box_h:.2f} in")
        return est

    def figure(self, s, path, y, bottom, x=BODY_X, max_w=FIG_W, max_h=FIG_H, fixed_h=False):
        """A picture fitted into max_w x max_h keeping its aspect ratio, centred across max_w; returns its height.

        Without a fixed h, max_h is the room left in the body area, so the picture shrinks to fit (§14.5)."""
        from PIL import Image
        from pptx.util import Inches

        try:
            with Image.open(path) as im:
                iw, ih = im.size
        except Exception as e:
            raise ToolError(f"cannot read figure {path}: {e}") from None
        if not fixed_h:
            max_h = min(max_h, FIG_H)               # at most the §3.2 size, less when the body area is shorter
        fw, fh = max_w, max_w * ih / iw
        if fh > max_h:
            fh, fw = max_h, max_h * iw / ih
        if y + fh > bottom + 0.05:
            self.warn(f"figure needs {fh:.2f} in from y={y:.2f} and runs past the body area")
        s.shapes.add_picture(str(path), Inches(x + (max_w - fw) / 2), Inches(y), width=Inches(fw), height=Inches(fh))
        return fh

    def excerpt(self, s, text, y, h, caption, x=BODY_X, w=BODY_W, size=EXCERPT_SIZE):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt

        box = self.rect(s, x, y, w, h, CODEBG, line=HAIRC)
        f = box.text_frame
        f.word_wrap = True
        f.vertical_anchor = MSO_ANCHOR.TOP
        f.margin_left = f.margin_right = Inches(0.18)
        f.margin_top = f.margin_bottom = Inches(0.12)
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            p = f.paragraphs[0] if i == 0 else f.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(0)
            p.font.name = MONO
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.name = MONO
            r.font.color.rgb = self.rgb(INK)
        self.check_text(y, w - 0.36, h - 0.2, [(line, size, True) for line in lines], 0.0)
        if caption:
            self.tb(s, x, y + h + 0.06, w, 0.35, [(caption, 12, False, MUTED)])
        return box

    def chart(self, s, body, x, y, w, h, stage):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
        from pptx.util import Inches, Pt

        cats = body["categories"]
        multi = body.get("series") is not None
        series_list = ([(sr["name"], sr["values"]) for sr in body["series"]] if multi
                       else [("values", body["values"])])
        takeaway = body.get("takeaway")
        size = body.get("size") or 14
        take_size = size + 4
        take_h = max(0.5, take_size * 1.2 / 72 + 0.2)
        chart_h = max(h - (take_h + 0.1 if takeaway else 0), MIN_BLOCK_H)
        line = body.get("chart_type", "bar") == "line"
        data = CategoryChartData()
        data.categories = cats
        for name, values in series_list:
            data.add_series(name, values)
        kind = XL_CHART_TYPE.LINE_MARKERS if line else XL_CHART_TYPE.BAR_CLUSTERED
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(chart_h), data)
        chart = gf.chart
        chart.has_legend = multi                     # §16.7: a legend names the series
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        chart.font.size = Pt(size)
        chart.font.name = FONT
        if not line:
            chart.category_axis.reverse_order = True     # first category on top, as in the spec
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.visible = False
        plot = chart.plots[0]
        if not line:
            plot.gap_width = 60
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_value = True
        decimals = 0 if all(float(v).is_integer() for _, values in series_list for v in values) else 1
        unit = (body.get("unit") or "").replace('"', "")
        labels.number_format = ("0" if decimals == 0 else "0.0") + (f'"{unit}"' if unit else "")
        labels.number_format_is_linked = False
        hl = body.get("highlight")
        hi = cats.index(hl) if isinstance(hl, str) else hl
        key = body.get("highlight_color") or stage or "gates"      # §14.5
        colour = (STAGE.get(key) or OUTCOME[key]).lstrip("#").upper()
        if line:
            labels.position = XL_LABEL_POSITION.ABOVE
        if multi:
            # 1 colour per series (the highlight colour first, then greys) so the legend tells them apart; on a
            # line chart the highlighted category gets bigger markers
            for k, series in enumerate(plot.series):
                rgb = self.rgb(colour if k == 0 else SERIES_GREYS[(k - 1) % len(SERIES_GREYS)])
                if not line:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = rgb
                    continue
                series.smooth = False
                series.format.line.color.rgb = rgb
                series.format.line.width = Pt(2.25)
                for i in range(len(cats)):
                    marker = series.points[i].marker
                    marker.style = XL_MARKER_STYLE.CIRCLE
                    marker.size = 11 if i == hi else 8
                    marker.format.fill.solid()
                    marker.format.fill.fore_color.rgb = rgb
                    marker.format.line.color.rgb = rgb
        else:
            series = plot.series[0]
            if line:
                series.smooth = False
                series.format.line.color.rgb = self.rgb(GREYBAR)
                series.format.line.width = Pt(2.25)
            for i in range(len(cats)):
                point = series.points[i]
                rgb = self.rgb(colour if i == hi else GREYBAR)
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb
                if line:
                    point.marker.style = XL_MARKER_STYLE.CIRCLE
                    point.marker.size = 11 if i == hi else 8
                    point.marker.format.fill.solid()
                    point.marker.format.fill.fore_color.rgb = rgb
                    point.marker.format.line.color.rgb = rgb
        if takeaway:
            self.tb(s, x, y + chart_h + 0.1, w, take_h, [(takeaway, take_size, True, INK)])
        return gf


def _read_times(times_file: str) -> tuple[pathlib.Path, dict]:
    """Check a times ledger before anything is written: its directory must exist, and the file, when present,
    must be a JSON object of non-negative numbers."""
    try:
        path = pathlib.Path(times_file).resolve()
        parent_ok, exists = path.parent.is_dir(), path.exists()
    except (OSError, ValueError) as e:
        raise ToolError(f"invalid times file path {times_file!r}: {e}") from None
    if not parent_ok:
        raise ToolError(f"times file directory does not exist: {path.parent}")
    if not exists:
        return path, {}
    data = read_json(path, "times file")
    if not isinstance(data, dict) or not all(_is_num(v) and v >= 0 for v in data.values()):
        raise ToolError(f"times file {path} must be a JSON object of non-negative numbers")
    return path, data


@tool("deck_build",
      "Build a PowerPoint deck (.pptx, 13.333 x 7.5 in) from a JSON deck spec (spec object or spec_path). Slides "
      "are title, divider or content; content slides take an eyebrow, a takeaway title, a stage colour, a source "
      "footer, DEMONSTRATED BY / THUS / IN-OUT strips and a body (bullets, lines, table, figure, excerpt, chart). "
      "Every slide needs notes.time (M:SS); notes become TIME / SAY / IF ASKED. Insert slides (insert: true) are "
      "numbered after the core slide they follow (3a, 3b) and can be shown, hidden or left out (inserts). "
      "Returns slide numbers, core and insert times against the budget, overflow warnings, and, with "
      "times_file, a shared time ledger across decks (needs meta.id). Needs python-pptx and Pillow.",
      {"type": "object",
       "properties": {
           "spec": {"type": "object", "description": "the deck spec"},
           "spec_path": {"type": "string", "description": "path to a JSON deck spec"},
           "out": {"type": "string", "description": "output .pptx path (overwritten)"},
           "inserts": {"type": "string", "enum": list(INSERT_MODES),
                       "description": "overrides meta.inserts: shown, hidden or off"},
           "times_file": {"type": "string", "description": "JSON time ledger shared by several decks"}},
       "required": ["out"],
       "additionalProperties": False},
      readOnlyHint=False)
def deck_build(out: str, spec: dict | None = None, spec_path: str | None = None, inserts: str | None = None,
               times_file: str | None = None) -> dict:
    plan = Plan(*load_spec(spec, spec_path))
    mode = inserts or plan.meta["inserts"]
    if not isinstance(mode, str) or mode not in INSERT_MODES:
        raise ToolError(f"inserts must be one of shown, hidden, off, not {_short(mode)}")
    if times_file is not None and not plan.meta["id"]:
        raise ToolError("meta.id is required when times_file is given")
    try:
        out_path = pathlib.Path(out).resolve()
        parent_ok = out_path.parent.is_dir()
    except (OSError, ValueError) as e:
        raise ToolError(f"invalid output path {out!r}: {e}") from None
    if not parent_ok:
        raise ToolError(f"output directory does not exist: {out_path.parent}")
    tpath, times = _read_times(times_file) if times_file is not None else (None, None)
    core, ins = plan.core_seconds, plan.insert_seconds
    combined = None
    if times is not None:                  # §16.1: the ledger sum is checked before anything is written
        deck_id = plan.meta["id"]
        times[deck_id] = core
        times[deck_id + "_inserts"] = 0 if mode == "off" else ins
        combined = _ledger_sum(times, {}, tpath)
    _require_office()

    builder = _Builder(plan, mode)
    prs = builder.build()
    try:
        prs.save(str(out_path))
    except (OSError, ValueError) as e:
        raise ToolError(f"cannot write {out}: {e}") from None

    target, maximum = parse_time(plan.meta["target"]), parse_time(plan.meta["max"])
    warnings = builder.warnings
    if core > maximum:
        warnings.append(f"slide all: core time {mmss(core)} is over the {plan.meta['max']} maximum")

    ledger = None
    if times is not None:
        try:
            tpath.write_text(json.dumps(times, indent=1) + "\n", encoding="utf-8")
        except (OSError, ValueError) as e:
            raise ToolError(f"cannot write times file {times_file}: {e}") from None
        decks = {k: v for k, v in times.items() if not k.endswith("_inserts")}
        ledger = {"path": str(tpath), "decks": decks, "combined_seconds": combined,
                  "combined_time": mmss(int(combined)), "status": budget_status(combined, target, maximum)}
        if combined > maximum:
            warnings.append(f"slide all: decks in {tpath.name} together take {mmss(int(combined))}, "
                            f"over the {plan.meta['max']} maximum")

    return {
        "out": str(out_path),
        "core_slides": sum(1 for s in plan.slides if not s["insert"]),
        "insert_slides": sum(1 for s in plan.slides if s["insert"]),
        "inserts": mode,
        "core_seconds": core,
        "insert_seconds": ins,
        "core_time": mmss(core),
        "total_time": mmss(core if mode == "off" else core + ins),
        "budget": {"target": plan.meta["target"], "max": plan.meta["max"],
                   "status": budget_status(core, target, maximum)},
        "ledger": ledger,
        "warnings": warnings,
        "slides": [{"number": s["number"], "type": s["type"], "title": s["title"], "insert": s["insert"],
                    "seconds": s["seconds"], "say_words": _words(s["say"])} for s in plan.slides],
    }


# ------------------------------------------------------------------ reading a .pptx (§3.3, §3.4)
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_KNOWN_SHAPES = {_P + t for t in ("sp", "cxnSp", "pic", "graphicFrame")}
_OTHER_SHAPES = {_P + "contentPart", _MC + "AlternateContent"}


def _raw_text(elm) -> str:
    """The text of any shape element: its a:p paragraphs (a:t runs, a:br as a line break), newline-joined.

    For mc:AlternateContent only the first mc:Choice is read, so the fallback copy is not counted twice."""
    if elm.tag == _MC + "AlternateContent":
        choice = next((c for c in elm if c.tag == _MC + "Choice"), None)
        if choice is None:
            choice = next((c for c in elm if c.tag == _MC + "Fallback"), None)
        return _raw_text(choice) if choice is not None else ""
    paras = []
    for p in elm.iter(_A + "p"):
        parts = []
        for node in p.iter(_A + "t", _A + "br"):
            parts.append("\n" if node.tag == _A + "br" else (node.text or ""))
        paras.append("".join(parts))
    return "\n".join(paras)


class _RawShape:
    """A shape python-pptx does not recognise (§16.1): no geometry, no text frame, only its raw text."""

    has_text_frame = has_table = has_chart = False
    top = left = width = height = shape_type = None

    def __init__(self, elm):
        self._element = elm
        try:
            self.raw_text = _raw_text(elm)
        except Exception:  # noqa: BLE001 - odd XML is ordinary content, never a crash
            self.raw_text = ""


def _iter_shapes(shapes):
    """Every leaf shape of a slide's shape tree in document order, groups flattened (§16.1: never a crash).

    Shapes python-pptx can't proxy (p:contentPart, mc:AlternateContent, anything that fails) come back as
    _RawShape, which has no geometry and no text frame but keeps its text."""
    tree = getattr(shapes, "_spTree", None)
    if tree is None:                       # not a python-pptx shape collection: iterate it as given
        yield from shapes
        return
    yield from _iter_tree(tree, shapes)


def _iter_tree(tree, shapes):
    from pptx.shapes.base import BaseShape

    for elm in tree.iterchildren():
        tag = elm.tag
        if tag == _P + "grpSp":
            yield from _iter_tree(elm, shapes)
        elif tag in _KNOWN_SHAPES:
            try:
                sh = shapes._shape_factory(elm)
            except Exception:  # noqa: BLE001
                sh = None
            yield sh if sh is not None and type(sh) is not BaseShape else _RawShape(elm)
        elif tag in _OTHER_SHAPES:
            yield _RawShape(elm)


def _shape_top(sh):
    try:
        return sh.top
    except Exception:  # noqa: BLE001 - geometry of an odd shape is skipped
        return None


def _open_pptx(pptx_path: str):
    _require_office(pillow=False)
    from pptx import Presentation

    p = pathlib.Path(pptx_path)
    try:
        is_file = p.is_file()
    except (OSError, ValueError):
        is_file = False
    if not is_file:
        raise ToolError(f"file not found: {pptx_path}")
    try:
        return Presentation(str(p))
    except Exception as e:  # python-pptx raises assorted errors for files that are not decks (zlib, zip, XML)
        raise ToolError(f"cannot open {pptx_path} as a .pptx: {type(e).__name__}: {e}") from None


def _slide_notes(slide) -> str:
    """Notes text; a notes slide without a notes placeholder (or any odd notes part) is empty notes (§16.1)."""
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        return frame.text if frame is not None else ""
    except Exception:  # noqa: BLE001
        return ""


def read_slides(prs, path) -> list[dict]:
    """read_slide for every slide; anything unexpected in a damaged deck is a ToolError, never a crash."""
    out = []
    try:
        for i, slide in enumerate(prs.slides, 1):
            out.append(read_slide(slide, i))
    except ToolError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"cannot read {path} as a .pptx: {type(e).__name__}: {e}") from None
    return out


def parse_notes(notes: str) -> tuple[int | None, str | None]:
    """(seconds or None, SAY text or None) from notes in the §3.2 format."""
    m = NOTES_TIME_RE.search(notes)
    seconds = int(m.group(1)) * 60 + int(m.group(2)) if m else None
    say = NOTES_SAY_RE.search(notes)
    return seconds, (say.group(1).strip() if say else None)


FOOTER_TOP_EMU = round(6.9 * 914400)     # §14.4: a footer frame's top is at or below 6.9 in


def _frame_read(sh):
    """(text, [(run text, size)]) of a shape's text frame, or None when it has none."""
    if isinstance(sh, _RawShape):
        return (sh.raw_text, []) if sh.raw_text.strip() else None
    try:
        if not (getattr(sh, "has_text_frame", False) and sh.has_text_frame):
            return None
        tf = sh.text_frame
        runs = []
        for p in tf.paragraphs:
            psize = p.font.size
            for r in p.runs:
                runs.append((r.text, r.font.size or psize or 0))
        return tf.text, runs
    except Exception:  # noqa: BLE001 - an odd text body: keep whatever text the XML holds
        text = _raw_text(sh._element)
        return (text, []) if text.strip() else None


def _table_rows(sh) -> list[list[str]] | None:
    try:
        if not (getattr(sh, "has_table", False) and sh.has_table):
            return None
        return [[c.text for c in row.cells] for row in sh.table.rows]
    except Exception:  # noqa: BLE001
        return None


def read_slide(slide, index: int) -> dict:
    frames, tables, low = [], [], []
    best_size, title = -1, None
    for sh in _iter_shapes(slide.shapes):
        got = _frame_read(sh)
        if got is not None:
            text, runs = got
            frames.append(text)
            top = None if isinstance(sh, _RawShape) else _shape_top(sh)
            if top is not None and top >= FOOTER_TOP_EMU:
                low.append(text.strip())
            for rtext, size in runs:
                if rtext.strip() and size > best_size:
                    best_size, title = size, text
        rows = _table_rows(sh)
        if rows is not None:
            tables.append(rows)
    if title is None:
        title = next((t for t in frames if t.strip()), "")
    number = next((t.strip() for t in frames if NUMBER_RE.fullmatch(t.strip())), None)
    notes = _slide_notes(slide)
    notes = notes.replace("\r\n", "\n").replace("\r", "\n").replace("\x0b", "\n")
    seconds, say = parse_notes(notes)
    footer = " ".join(t for t in low if t and not NUMBER_RE.fullmatch(t))
    cells = [c for rows in tables for row in rows for c in row]
    try:
        hidden = slide._element.get("show") in ("0", "false")
    except Exception:  # noqa: BLE001
        hidden = False
    return {"index": index, "number": number, "title": title, "texts": frames + cells, "notes": notes,
            "seconds": seconds, "say_words": _words(say or ""), "insert": notes.startswith("INSERT:"),
            "hidden": hidden, "say": say or "", "footer": footer, "frames": frames, "tables": tables}


@tool("deck_inspect",
      "Read any .pptx and list its slides: 1-based index, slide number text (a text box that is just a number "
      "such as 3 or 3a), title (the text in the largest font), every text in shape order (table cells after), "
      "speaker notes, TIME seconds and SAY word count parsed from the notes, whether it is an insert slide "
      "(notes start with INSERT:) and whether it is hidden. Needs python-pptx.",
      {"type": "object",
       "properties": {"pptx_path": {"type": "string", "description": "path to the .pptx file"}},
       "required": ["pptx_path"],
       "additionalProperties": False},
      readOnlyHint=True)
def deck_inspect(pptx_path: str) -> dict:
    return {"path": str(pptx_path), "slides": [_public(info) for info in _inspect(pptx_path)]}


_PRIVATE_KEYS = ("say", "footer", "frames", "tables")


def _inspect(pptx_path: str) -> list[dict]:
    """Every slide's read_slide record (with the private keys)."""
    return read_slides(_open_pptx(pptx_path), pptx_path)


def _public(info: dict) -> dict:
    return {k: v for k, v in info.items() if k not in _PRIVATE_KEYS}


# ------------------------------------------------------------------ linting (§3.3, §14.4)
DEFENSIVE = ("i'm not going to", "i am not going to", "i'm not claiming", "i'll show", "i will show",
             "this talk will", "that's fine for",
             "i'll go through", "i'm going to show", "i will go through", "instead i'll")
BACKREF_RE = re.compile(r"(?i)\b(as we saw|as i said|as mentioned|next slide|coming up)\b")
INSERT_REF_RE = re.compile(r"(?i)\bslide\s+\d+[a-z]\b")
EQ_RE = re.compile(r"\bEq\.?\s*\d")
# D014 typography: an em dash; a double period that is not part of `...`; exactly 2 spaces between a
# lower-case letter or .,;:) and a capital (3 or more spaces are a deliberate gutter, as in the phase strip)
TYPO_RES = (("an em dash", re.compile("—")),
            ("a double period", re.compile(r"(?<!\.)\.\.(?!\.)")),
            ("a double space mid-sentence", re.compile(r"(?<=[a-z.,;:)])  (?=[A-Z])")))


def _excerpt(text: str, start: int = 0, end: int | None = None) -> str:
    end = start if end is None else end
    a, b = max(0, start - 30), min(len(text), end + 30)
    return " ".join(text[a:b].split())[:80]


def _lint_records(records: list[dict], shown: str, wps: float) -> list[dict]:
    """Per-slide findings (every rule but D008 and D013) for 1 deck."""
    findings = []

    def add(rule, severity, rec, message, excerpt=""):
        findings.append({"rule": rule, "severity": severity, "path": shown,
                         "line": rec["position"], "message": message, "excerpt": excerpt})

    for r in records:
        say, notes, face, title = r["say"], r["notes"], r["face"], r["title"]
        words = _words(say)
        if r["seconds"] is None:
            add("D012", "error", r, "slide has no TIME in its notes")
        elif words > r["seconds"] * wps + 1e-9:
            add("D001", "error", r, f"SAY is {words} words for {mmss(r['seconds'])} "
                                    f"(cap {int(r['seconds'] * wps)} at {wps:g} words/s)", _excerpt(say))
        if r["content"] and not r["insert"] and words < 20:
            add("D002", "warning", r, f"SAY is only {words} words on a content slide", _excerpt(say))
        if "—" in title:
            add("D003", "error", r, "title contains an em dash", _excerpt(title, title.index("—")))
        for phrase in ("MUST HIT", "The point of this slide"):
            if phrase in notes:
                i = notes.index(phrase)
                add("D004", "error", r, f"notes contain {phrase!r}", _excerpt(notes, i, i + len(phrase)))
        low = say.replace("’", "'").lower()
        for phrase in DEFENSIVE:
            if phrase in low:
                i = low.index(phrase)
                add("D005", "warning", r, f"SAY contains the meta or defensive line {phrase!r}",
                    _excerpt(say, i, i + len(phrase)))
                break
        if r["insert"]:
            for where, text in (("face text", face), ("SAY", say)):
                m = BACKREF_RE.search(text)
                if m:
                    add("D006", "error", r, f"insert slide refers to the flow in its {where}: {m.group(0)!r}",
                        _excerpt(text, m.start(), m.end()))
                    break
        else:
            for where, text in (("face text", face), ("SAY", say)):
                m = INSERT_REF_RE.search(text)
                if m:
                    add("D007", "error", r, f"core slide mentions an insert slide in its {where}: {m.group(0)!r}",
                        _excerpt(text, m.start(), m.end()))
                    break
        if r["content"] and not r["source"]:
            add("D009", "warning", r, "content slide has no source footer")
        m = re.search(r"et al\.", face) or EQ_RE.search(face)
        if m:
            add("D010", "warning", r, f"face text contains {m.group(0)!r}", _excerpt(face, m.start(), m.end()))
        if len(title) > 90:
            add("D011", "warning", r, f"title is {len(title)} characters (over 90)", title[:80])
        found, first = [], None
        for where, text in (("face text", r["face_body"]), ("SAY", say)):
            for what, rx in TYPO_RES:
                m = rx.search(text)
                if m:
                    found.append(f"{what} in {where}")
                    if first is None:
                        first = _excerpt(text, m.start(), m.end())
        if found:
            add("D014", "warning", r, "typography: " + "; ".join(found), first)
    return findings


def _d008(core: int, target: int, maximum: int, shown: str, what: str) -> dict | None:
    if core > maximum:
        severity, message = "error", f"{what} {mmss(core)} is over the {mmss(maximum)} maximum"
    elif core > target:
        severity, message = "warning", f"{what} {mmss(core)} is over the {mmss(target)} target"
    else:
        return None
    return {"rule": "D008", "severity": severity, "path": shown, "line": None, "message": message, "excerpt": ""}


def _time_arg(name: str, value, default: str) -> int:
    value = default if value is None else value
    seconds = parse_time(value)
    if seconds is None:
        raise ToolError(f"{name} must be M:SS, not {_short(value)}")
    return seconds


def _pptx_records(path: str) -> list[dict]:
    records = []
    for info in _inspect(path):
        i = info["index"]
        texts = list(info["texts"])
        if info["title"] in texts:
            texts.remove(info["title"])          # D014 leaves titles to D003
        records.append({"position": i, "title": info["title"], "face": "\n".join(info["texts"]),
                        "face_body": "\n".join(texts), "say": info["say"], "notes": info["notes"],
                        "seconds": info["seconds"], "insert": info["insert"],
                        "content": info["number"] is not None, "source": info["footer"]})
    return records


def _read_ledger(times_file) -> tuple[str, dict]:
    """The times ledger for linting: it must exist and be a JSON object of non-negative numbers."""
    if not isinstance(times_file, str) or not times_file:
        raise ToolError("times_file must be a non-empty path string")
    path = pathlib.Path(times_file)
    data = read_json(path, "times file")
    if not isinstance(data, dict) or not all(_is_num(v) and v >= 0 for v in data.values()):
        raise ToolError(f"times file {times_file} must be a JSON object of non-negative numbers")
    return _display(times_file), data


def _ledger_sum(ledger: dict, own: dict, shown) -> float:
    """Combined core seconds of every non-`_inserts` id, with `own` values replacing the ledger's (§14.4).

    A sum that is not finite (or too large to use as a time) is a ToolError (§16.1), never an OverflowError."""
    try:
        total = math.fsum([float(v) for k, v in ledger.items() if not k.endswith("_inserts") and k not in own]
                          + [float(v) for v in own.values()])
    except (OverflowError, ValueError):
        total = math.inf
    if not math.isfinite(total) or total > 1e15:
        raise ToolError(f"times file {shown}: the combined core time is not a finite number of seconds")
    return int(total) if total.is_integer() else total


@tool("deck_lint",
      "Check a deck spec (spec or spec_path) or built .pptx decks (pptx_path: 1 path or a list) against the "
      "speaker-note and slide rules: D001 SAY longer than TIME x words_per_second; D002 core content slide with "
      "under 20 SAY words; D003 em dash in a title; D004 MUST HIT or 'The point of this slide' in notes; D005 "
      "meta or defensive lines in SAY; D006 insert slide referring to the flow; D007 core slide mentioning an "
      "insert slide; D008 core time over target (warning) or max (error); D009 content slide without a source "
      "footer; D010 'et al.' or 'Eq. N' on the face; D011 title over 90 chars; D012 slide without TIME; D013 "
      "core time differs from the times_file entry (deck id = file stem, or ids); D014 em dash, '..' or a "
      "mid-sentence double space in face text or SAY. With times_file, D008 checks the combined time of all "
      "decks. A spec needs no optional packages; a .pptx needs python-pptx.",
      {"type": "object",
       "properties": {
           "spec": {"type": "object", "description": "the deck spec"},
           "spec_path": {"type": "string", "description": "path to a JSON deck spec"},
           "pptx_path": {"type": ["string", "array"], "items": {"type": "string"},
                         "description": "path to a built .pptx, or a list of paths"},
           "words_per_second": {"type": "number", "description": "speaking rate (default meta, else 2.3)"},
           "target": {"type": "string", "description": "target core time M:SS (default meta budget, else 40:00)"},
           "max": {"type": "string", "description": "maximum core time M:SS (default meta budget, else 45:00)"},
           "times_file": {"type": "string",
                          "description": "JSON time ledger {deck_id: seconds} shared by several decks (read only)"},
           "ids": {"type": "array", "items": {"type": "string"},
                   "description": "deck ids for the times_file, 1 per deck (default: the file stems)"}},
       "additionalProperties": False},
      readOnlyHint=True)
def deck_lint(spec: dict | None = None, spec_path: str | None = None, pptx_path=None,
              words_per_second: float | None = None, target: str | None = None, max: str | None = None,
              times_file: str | None = None, ids: list[str] | None = None) -> dict:
    if sum(x is not None for x in (spec, spec_path, pptx_path)) != 1:
        raise ToolError("give exactly 1 of spec, spec_path or pptx_path")
    if words_per_second is not None and (not _is_num(words_per_second) or words_per_second <= 0):
        raise ToolError("words_per_second must be a positive number")
    if ids is not None and (not isinstance(ids, list) or not all(isinstance(i, str) and i for i in ids)):
        raise ToolError("ids must be a list of non-empty strings")

    decks = []                          # (shown path, records, deck id or None)
    if pptx_path is not None:
        paths = [pptx_path] if isinstance(pptx_path, str) else pptx_path
        if not isinstance(paths, list) or not paths or not all(isinstance(p, str) and p for p in paths):
            raise ToolError("pptx_path must be a path or a non-empty list of paths")
        if ids is not None and len(ids) != len(paths):
            raise ToolError(f"ids has {len(ids)} entries for {len(paths)} deck(s)")
        wps = words_per_second or DEFAULT_WPS
        target_s = _time_arg("target", target, DEFAULT_TARGET)
        max_s = _time_arg("max", max, DEFAULT_MAX)
        for k, p in enumerate(paths):
            deck_id = ids[k] if ids is not None else pathlib.Path(p).stem
            decks.append((_display(p), _pptx_records(p), deck_id))
    else:
        plan = Plan(*load_spec(spec, spec_path))
        if ids is not None and len(ids) != 1:
            raise ToolError(f"ids has {len(ids)} entries for 1 deck")
        wps = words_per_second or plan.meta["words_per_second"]
        target_s = _time_arg("target", target, plan.meta["target"])
        max_s = _time_arg("max", max, plan.meta["max"])
        records = [{"position": s["position"], "title": s["title"], "face": s["face"], "face_body": s["face_body"],
                    "say": s["say"], "notes": s["notes_text"], "seconds": s["seconds"], "insert": s["insert"],
                    "content": s["type"] == "content", "source": s["spec"].get("source")}
                   for s in plan.slides]
        deck_id = (ids[0] if ids is not None else plan.meta["id"]
                   or (pathlib.Path(spec_path).stem if spec_path is not None else None))
        decks.append((plan.shown, records, deck_id))

    ledger_path, ledger = _read_ledger(times_file) if times_file is not None else (None, None)
    if ledger is not None and any(d[2] is None for d in decks):
        raise ToolError("an inline spec needs meta.id (or ids) when times_file is given")

    findings, summary, measured = [], [], {}
    for shown, records, deck_id in decks:
        core = sum(r["seconds"] or 0 for r in records if not r["insert"])
        ins = sum(r["seconds"] or 0 for r in records if r["insert"])
        summary.append({"path": shown, "id": deck_id, "core_seconds": core, "insert_seconds": ins,
                        "core_time": mmss(core), "slides": len(records)})
        findings.extend(_lint_records(records, shown, wps))
        if ledger is None:
            f = _d008(core, target_s, max_s, shown, "core time")
            if f:
                findings.append(f)
            continue
        measured[deck_id] = measured.get(deck_id, 0) + core
        entry = ledger.get(deck_id)
        if entry is None:
            keys = sorted(ledger)
            present = ", ".join(repr(k) for k in keys[:30]) + (", ..." if len(keys) > 30 else "") or "(none)"
            findings.append({"rule": "D013", "severity": "error", "path": shown, "line": None,
                             "message": f"{ledger_path} has no entry {deck_id!r} (measured core time {mmss(core)}); "
                                        f"keys present: {present}; use --ids (ids) to give each deck's "
                                        f"ledger id", "excerpt": ""})
        elif abs(float(entry) - core) > 1e-9:
            findings.append({"rule": "D013", "severity": "error", "path": shown, "line": None,
                             "message": f"core time {mmss(core)} ({core} s) differs from {ledger_path} entry "
                                        f"{deck_id!r} ({entry:g} s)", "excerpt": ""})

    combined = None
    if ledger is not None:
        combined = _ledger_sum(ledger, measured, ledger_path)
        f = _d008(int(math.ceil(combined)), target_s, max_s, decks[0][0], f"combined core time in {ledger_path}")
        if f:
            findings.append(f)

    findings.sort(key=lambda f: (f["path"], f["line"] or 0, f["rule"]))
    counts = {sev: sum(1 for f in findings if f["severity"] == sev) for sev in ("error", "warning", "info")}
    core = sum(d["core_seconds"] for d in summary)
    result = {"ok": counts["error"] == 0, "findings": findings, "counts": counts,
              "core_seconds": core, "insert_seconds": sum(d["insert_seconds"] for d in summary),
              "core_time": mmss(core), "slides": sum(d["slides"] for d in summary), "decks": summary}
    if ledger is not None:
        result["ledger"] = {"path": ledger_path, "combined_seconds": combined,
                            "combined_time": mmss(int(math.ceil(combined))),
                            "status": budget_status(combined, target_s, max_s)}
    return result


# ------------------------------------------------------------------ diffing decks (§15.4)
HINT_SUFFIXES = (".py", ".json", ".md")
PAIR_SIMILARITY, RELATED_SIMILARITY = 0.8, 0.4      # §16.2 pairing thresholds


_GIT_BASH_RE = re.compile(r"^/([A-Za-z])(?:/(.*))?$")


def _native_path(path: str) -> str:
    """On Windows, a Git Bash drive path `/c/x` means `C:/x` (§16.1) when `/c/x` itself doesn't exist."""
    import os

    if os.name != "nt":
        return path
    m = _GIT_BASH_RE.match(path.replace("\\", "/")) if path.startswith(("/", "\\")) else None
    if not m:
        return path
    try:
        if os.path.exists(path):
            return path
    except (OSError, ValueError):
        pass
    return f"{m.group(1).upper()}:/{m.group(2) or ''}"


def search_files(search: list[str] | None) -> list[tuple[str, str]]:
    """(hint path, text) for every searchable file under the `search` entries (§15.4).

    Directories are walked for .py, .json and .md files; a file given directly is searched whatever its
    suffix. Hint paths join the entry with the walked path, so a relative entry gives relative hints.
    """
    import os

    out, seen = [], set()
    for entry in search or []:
        if not isinstance(entry, str) or not entry:
            raise ToolError("search entries must be non-empty path strings")
        given, entry = entry, _native_path(entry)
        try:
            is_dir, is_file = os.path.isdir(entry), os.path.isfile(entry)
        except (OSError, ValueError):
            is_dir = is_file = False
        if is_file:
            paths = [(given, entry)]
        elif is_dir:
            paths = []
            for dirpath, dirnames, filenames in os.walk(entry):
                dirnames.sort()
                for name in sorted(filenames):
                    if name.lower().endswith(HINT_SUFFIXES):
                        full = os.path.join(dirpath, name)
                        paths.append((given + full[len(entry):], full))
        else:
            raise ToolError(f"search path not found: {given}")
        for shown, full in paths:
            try:
                key = os.path.normcase(os.path.realpath(full))
            except (OSError, ValueError):
                key = full
            if key in seen:                 # §16.1: the same file under 2 spellings is 1 input
                continue
            seen.add(key)
            try:
                with open(full, encoding="utf-8-sig", errors="replace") as f:
                    out.append((shown.replace("\\", "/"), f.read()))
            except OSError:
                continue
    return out


def find_hints(old: str, files: list[tuple[str, str]], limit: int = 3, new: str | None = None) -> list[str]:
    """Up to `limit` `path:line` places where the first hint candidate that occurs anywhere occurs (§17.2):
    `old`, then each changed line of `old` against `new`, then its longest line of 12+ characters."""
    from tundlekit.textlint import hint_needles

    if not old or not files:
        return []
    for needle in hint_needles(old, new):
        hits = []
        for shown, text in files:
            start = text.find(needle)
            while start != -1 and len(hits) < limit:
                hits.append(f"{shown}:{text.count(chr(10), 0, start) + 1}")
                start = text.find(needle, start + len(needle))
            if len(hits) >= limit:
                break
        if hits:
            return hits
    return []


def _similar(a: str, b: str) -> float:
    import difflib

    return difflib.SequenceMatcher(None, _norm_title(a), _norm_title(b), autojunk=False).ratio()


def _norm_title(s: str) -> str:
    return " ".join(s.casefold().split())


def _unified(old: str, new: str) -> str:
    import difflib

    return "\n".join(difflib.unified_diff(old.split("\n"), new.split("\n"), "old", "new", lineterm=""))


def _pair_slides(old: list[dict], new: list[dict]) -> dict[int, int]:
    """new index -> old index, in the §16.2 order.

    1. exact title (after whitespace collapse and case folding), then title similarity >= 0.8;
    2. same number text with title similarity >= 0.4;
    3. the rest by position, with title similarity >= 0.4;
    anything left is added or removed. Among several candidates in a step, the same number text wins, then the
    nearest position, then the higher similarity; ties go to the earlier slides."""
    pairs: dict[int, int] = {}
    used: set[int] = set()
    norm_old = [_norm_title(o["title"]) for o in old]
    norm_new = [_norm_title(n["title"]) for n in new]
    cache: dict[tuple[int, int], float] = {}

    def sim(j: int, i: int) -> float:
        if (j, i) not in cache:
            a, b = norm_new[j], norm_old[i]
            if a == b:
                cache[j, i] = 1.0
            else:
                import difflib

                sm = difflib.SequenceMatcher(None, b, a, autojunk=False)
                # the cheap upper bounds first: most pairs of a long deck are far apart
                cache[j, i] = sm.ratio() if sm.real_quick_ratio() >= 0.4 and sm.quick_ratio() >= 0.4 else 0.0
        return cache[j, i]

    def same_number(j: int, i: int) -> bool:
        return new[j]["number"] is not None and new[j]["number"] == old[i]["number"]

    def assign(candidates) -> None:
        """Greedy assignment of (j, i, score) candidates by preference."""
        ranked = sorted(candidates, key=lambda c: (not same_number(c[0], c[1]), abs(c[0] - c[1]), -c[2], c[0], c[1]))
        for j, i, _ in ranked:
            if j not in pairs and i not in used:
                pairs[j] = i
                used.add(i)

    def free():
        for j in range(len(new)):
            if j in pairs:
                continue
            for i in range(len(old)):
                if i not in used:
                    yield j, i

    # step 1: exact titles (non-empty), then similar titles
    assign((j, i, 1.0) for j, i in free() if norm_new[j] and norm_new[j] == norm_old[i])
    assign((j, i, s) for j, i in free() if norm_new[j] and norm_old[i] and (s := sim(j, i)) >= PAIR_SIMILARITY)
    # step 2: same number text, related titles
    assign((j, i, s) for j, i in free() if same_number(j, i) and (s := sim(j, i)) >= RELATED_SIMILARITY)
    # step 3: the rest by position, related titles
    assign((j, i, s) for j, i in free() if (s := sim(j, i)) >= RELATED_SIMILARITY)
    return pairs


def _kept_in_order(seq: list[int]) -> set[int]:
    """Indexes of seq forming a longest increasing subsequence (the slides that did not move)."""
    import bisect

    tails, tails_idx, prev = [], [], [-1] * len(seq)
    for k, v in enumerate(seq):
        pos = bisect.bisect_left(tails, v)
        if pos == len(tails):
            tails.append(v)
            tails_idx.append(k)
        else:
            tails[pos] = v
            tails_idx[pos] = k
        prev[k] = tails_idx[pos - 1] if pos else -1
    keep, k = set(), tails_idx[-1] if tails_idx else -1
    while k != -1:
        keep.add(k)
        k = prev[k]
    return keep


def _without_title(texts: list[str], title: str) -> list[str]:
    out = list(texts)
    if title in out:
        out.remove(title)
    return out


def _frame_changes(old_texts: list[str], new_texts: list[str]):
    """(old, new) per changed text frame, aligned with difflib."""
    import difflib

    sm = difflib.SequenceMatcher(None, old_texts, new_texts, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        for k in range(max(i2 - i1, j2 - j1)):
            o = old_texts[i1 + k] if i1 + k < i2 else ""
            n = new_texts[j1 + k] if j1 + k < j2 else ""
            if o != n:
                yield o, n


@tool("deck_diff",
      "Compare 2 .pptx decks (old and new, for example a built deck and a hand-edited copy) and list the changes "
      "per slide: title, text (per text frame; 1 per table, rows as cells joined by ' | '), notes, hidden, number "
      "(renumbered), added, removed and moved slides (outside a longest increasing run), with a unified diff for "
      "text changes. Slides pair by exact or similar title (>= 0.8), then same number, then position (titles "
      ">= 0.4 similar). slide is the number text or #position. With "
      "search (files or directories; .py, .json and .md are walked), each change carries up to 3 'path:line' "
      "hints where the old text occurs, to carry hand edits back into build scripts. Needs python-pptx.",
      {"type": "object",
       "properties": {
           "old": {"type": "string", "description": "the original .pptx"},
           "new": {"type": "string", "description": "the edited .pptx"},
           "search": {"type": "array", "items": {"type": "string"},
                      "description": "files or directories to search for the old text"}},
       "required": ["old", "new"],
       "additionalProperties": False},
      readOnlyHint=True)
def deck_diff(old: str, new: str, search: list[str] | None = None) -> dict:
    old_slides = _inspect(old)
    new_slides = _inspect(new)
    files = search_files(search)
    pairs = _pair_slides(old_slides, new_slides)
    changes = []

    def add(slide, field, o, n, diff=None, hint_text=None):
        changes.append({"slide": slide, "field": field, "old": o, "new": n, "diff": diff,
                        "hint": find_hints(hint_text, files, new=n if isinstance(n, str) else None)
                        if hint_text else []})

    paired_new = sorted(pairs)
    kept = _kept_in_order([pairs[j] for j in paired_new])
    moved = {paired_new[k] for k in range(len(paired_new)) if k not in kept}

    for j, n in enumerate(new_slides):
        label = _slide_label(n)
        if j not in pairs:
            add(label, "added", "", n["title"])
            continue
        o = old_slides[pairs[j]]
        if j in moved:
            add(label, "moved", o["index"], n["index"])
        if o["number"] != n["number"]:
            add(label, "number", o["number"] or "", n["number"] or "")
        if o["title"] != n["title"]:
            add(label, "title", o["title"], n["title"], _unified(o["title"], n["title"]), o["title"])
        for ot, nt in _frame_changes(_body_frames(o), _body_frames(n)):
            add(label, "text", ot, nt, _unified(ot, nt), ot)
        for ot, nt in _table_changes(o["tables"], n["tables"]):
            add(label, "text", ot, nt, _unified(ot, nt), ot)
        if o["notes"] != n["notes"]:
            add(label, "notes", o["notes"], n["notes"], _unified(o["notes"], n["notes"]), o["notes"])
        if o["hidden"] != n["hidden"]:
            add(label, "hidden", o["hidden"], n["hidden"])
    matched_old = set(pairs.values())
    for i, o in enumerate(old_slides):
        if i not in matched_old:
            add(_slide_label(o), "removed", o["title"], "", hint_text=o["title"])
    return {"changes": changes}


def _slide_label(info: dict) -> str:
    """§16.2: the number text, or "#<position>" for a slide without one."""
    return info["number"] if info["number"] is not None else f"#{info['index']}"


def _body_frames(info: dict) -> list[str]:
    """A slide's text frames without its title and its slide-number frame (those have their own fields)."""
    out = _without_title(info["frames"], info["title"])
    if info["number"] is not None:
        k = next((k for k, t in enumerate(out) if t.strip() == info["number"]), None)
        if k is not None:
            del out[k]
    return out


def _table_text(rows: list[list[str]]) -> str:
    return "\n".join(" | ".join(row) for row in rows)


def _table_changes(old_tables: list, new_tables: list):
    """(old, new) per changed table (§16.2: at most 1 change per table), rows as lines of cells joined by ' | '.

    Tables pair in order; a table only on 1 side compares with an empty one."""
    for k in range(max(len(old_tables), len(new_tables))):
        o = _table_text(old_tables[k]) if k < len(old_tables) else ""
        n = _table_text(new_tables[k]) if k < len(new_tables) else ""
        if o != n:
            yield o, n


# ------------------------------------------------------------------ CLI
def add_cli(groups) -> None:
    from tundlekit.cli_support import common_flags

    p = groups.add_parser("deck", help="presentations from a JSON spec")
    cmds = p.add_subparsers(dest="command", required=True)

    c = cmds.add_parser("build", help="build a .pptx from a deck spec")
    c.add_argument("spec_path", help="deck spec (.json)")
    c.add_argument("-o", "--out", required=True, help="output .pptx")
    c.add_argument("--inserts", choices=INSERT_MODES, help="shown, hidden or off (overrides meta.inserts)")
    c.add_argument("--times-file", dest="times_file", help="JSON time ledger shared by several decks")
    common_flags(c)
    c.set_defaults(handler=_cli_build)

    c = cmds.add_parser("lint", help="check a deck spec or built .pptx decks")
    c.add_argument("path", nargs="+", help="deck spec (.json), or 1 or more decks (.pptx)")
    c.add_argument("--words-per-second", dest="words_per_second", type=float, help="speaking rate")
    c.add_argument("--target", help="target core time M:SS")
    c.add_argument("--max", help="maximum core time M:SS")
    c.add_argument("--times-file", dest="times_file", help="JSON time ledger shared by several decks (read only)")
    c.add_argument("--ids", nargs="+", metavar="ID", help="deck ids for --times-file, 1 per deck (default: stems)")
    common_flags(c, checker=True)
    c.set_defaults(handler=_cli_lint)

    c = cmds.add_parser("diff", help="list hand edits between 2 .pptx decks")
    c.add_argument("old", help="the original .pptx")
    c.add_argument("new", help="the edited .pptx")
    c.add_argument("--search", nargs="+", action="extend", metavar="PATH",
                   help="files or directories to search for the old text")
    common_flags(c)
    c.set_defaults(handler=_cli_diff)

    c = cmds.add_parser("inspect", help="list the slides, texts and notes of a .pptx")
    c.add_argument("pptx_path", help="deck (.pptx)")
    common_flags(c)
    c.set_defaults(handler=_cli_inspect)


def _cli_build(args):
    from tundlekit.cli_support import CliResult

    r = deck_build(out=args.out, spec_path=args.spec_path, inserts=args.inserts, times_file=args.times_file)
    lines = [f"wrote {r['out']}: {r['core_slides']} core slides, planned {r['core_time']} "
             f"({r['budget']['status']}; target {r['budget']['target']}, max {r['budget']['max']})"]
    if r["insert_slides"]:
        lines.append(f"  {r['insert_slides']} insert slides ({r['inserts']}), together {r['total_time']}")
    if r["ledger"]:
        led = r["ledger"]
        lines.append(f"  ledger {led['path']}: combined {led['combined_time']} ({led['status']})")
    lines += [f"  WARN {w}" for w in r["warnings"]]
    return CliResult(r, text="\n".join(lines))


def _cli_lint(args):
    from tundlekit.cli_support import CliResult, format_findings

    kw = {"words_per_second": args.words_per_second, "target": args.target, "max": args.max,
          "times_file": args.times_file, "ids": args.ids}
    kw = {k: v for k, v in kw.items() if v is not None}
    pptx = [p for p in args.path if p.lower().endswith(".pptx")]
    if len(pptx) == len(args.path):
        r = deck_lint(pptx_path=pptx[0] if len(pptx) == 1 else pptx, **kw)
    elif len(args.path) == 1:
        r = deck_lint(spec_path=args.path[0], **kw)
    else:
        raise ToolError("give either 1 deck spec (.json) or 1 or more decks (.pptx)")
    text = format_findings(r) + f"\n{r['slides']} slides, core time {r['core_time']}"
    if "ledger" in r:
        led = r["ledger"]
        text += f"\ncombined core time in {led['path']}: {led['combined_time']} ({led['status']})"
    return CliResult(r, text=text)


def _cli_inspect(args):
    from tundlekit.cli_support import CliResult

    r = deck_inspect(args.pptx_path)
    lines = []
    for s in r["slides"]:
        flags = " ".join(f for f, on in (("insert", s["insert"]), ("hidden", s["hidden"])) if on)
        time = mmss(s["seconds"]) if s["seconds"] is not None else "no TIME"
        lines.append(f"{s['index']:>3} [{s['number'] or '-'}] {s['title']!r} {time}, "
                     f"{s['say_words']} SAY words {flags}".rstrip())
    return CliResult(r, text="\n".join(lines))


def _cli_diff(args):
    from tundlekit.cli_support import CliResult

    r = deck_diff(args.old, args.new, search=args.search)
    lines = []
    for c in r["changes"]:
        lines.append(f"slide {c['slide']}: {c['field']}: {c['old']!r} -> {c['new']!r}")
        lines += [f"  at {h}" for h in c["hint"]]
    return CliResult(r, text="\n".join(lines) or "no changes")
