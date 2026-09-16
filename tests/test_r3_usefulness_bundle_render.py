"""MANIFEST §14.8 bundle_lint (B005 for git-ignored junk, B014 installer coverage) and §14.9 render_office
(the text backend's TITLE line uses the largest font)."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import pytest

from helpers_core import by_rule, call, check_shape, gitenv, lint_tree, run_cli_json  # noqa: F401

GITIGNORE = "__pycache__/\n*.pyc\n*.tmp\n.pytest_cache/\n"


def lint(root, **kw):
    r = call("bundle_lint", root=root, **kw)
    check_shape(r)
    return r


def sev(r, rule):
    return {f["path"].rstrip("/"): f["severity"] for f in by_rule(r, rule)}


# ====================================================================== B005 and .gitignore
def test_b005_gitignored_junk_is_info(tmp_path, gitenv):  # noqa: F811
    """§14.8: junk that git ignores is info (it still travels with a copied folder); other junk stays a warning."""
    root = lint_tree(tmp_path / "t", files={
        ".gitignore": GITIGNORE,
        "papers/README.md": "# Papers\n",
        "papers/build/__pycache__/figures.cpython-312.pyc": "x",
        "papers/draft.tmp": "x",
        "papers/Thumbs.db": "x",
    })
    r = lint(root)
    assert sev(r, "B005") == {
        "papers/build/__pycache__": "info",
        "papers/draft.tmp": "info",
        "papers/Thumbs.db": "warning",
    }
    for f in by_rule(r, "B005"):
        if f["severity"] == "info":
            assert re.search(r"cop(y|ied|ies)", f["message"], re.I), f["message"]


def test_b005_without_gitignore_stays_warning(tmp_path, gitenv):  # noqa: F811
    """§2.7/§14.8: junk that git does not ignore is a warning."""
    root = lint_tree(tmp_path / "t", files={"notes/README.md": "x", "notes/.DS_Store": "x"})
    assert sev(lint(root), "B005") == {"notes/.DS_Store": "warning"}


def test_b005_info_does_not_fail_strict(tmp_path, gitenv, capsys):  # noqa: F811
    """§14.8 with §0.4: ignored junk alone is info, so --strict still exits 0."""
    root = lint_tree(tmp_path / "t", files={".gitignore": GITIGNORE, "tools/__pycache__/a.pyc": "x"})
    rc, data, _ = run_cli_json(capsys, ["bundle", "lint", "--root", root, "--strict"])
    assert rc == 0
    assert [f["severity"] for f in data["findings"] if f["rule"] == "B005"] == ["info"]


# ====================================================================== B014
SETUP_README = "# Setup\n"


def _readme(*names):
    return "| File | What |\n|---|---|\n" + "".join(f"| `{n}` | installer |\n" for n in names)


def _src(payload, name):
    return f"- File: {name}\n- SHA-256: `{hashlib.sha256(payload).hexdigest()}`\n"


def test_b014_counts_uncovered_installers(tmp_path):
    """§14.8 B014: 1 info finding on `setup`, line null, counting installer files without a SOURCE.md."""
    root = lint_tree(tmp_path / "t", files={
        "setup/README.md": SETUP_README,
        "setup/windows/README.md": _readme("Git-2.45.0-64-bit.exe", "python-3.12.4-amd64.exe", "vscode/"),
        "setup/windows/Git-2.45.0-64-bit.exe": "g",
        "setup/windows/python-3.12.4-amd64.exe": "p",
        "setup/windows/vscode/VSCodeSetup-x64-1.93.1.exe": "v",
    })
    f = by_rule(lint(root), "B014")
    assert len(f) == 1
    assert f[0]["severity"] == "info"
    assert f[0]["line"] is None
    assert f[0]["path"] == "setup"
    assert re.search(r"\b3\b", f[0]["message"]), f[0]["message"]


def test_b014_source_md_covers_directory(tmp_path):
    """§14.8: files in a directory that has a SOURCE.md are covered; the others are counted."""
    exe = b"installer bytes"
    root = lint_tree(tmp_path / "t", files={
        "setup/README.md": SETUP_README,
        "setup/linux/README.md": _readme("code_1.93.1_amd64.deb"),
        "setup/linux/code_1.93.1_amd64.deb": exe,
        "setup/linux/SOURCE.md": _src(exe, "code_1.93.1_amd64.deb"),
        "setup/windows/README.md": _readme("7z2408-x64.exe", "Git-2.45.0-64-bit.exe"),
        "setup/windows/7z2408-x64.exe": "a",
        "setup/windows/Git-2.45.0-64-bit.exe": "b",
    })
    f = by_rule(lint(root), "B014")
    assert len(f) == 1
    assert re.search(r"\b2\b", f[0]["message"]), f[0]["message"]


def test_b014_program_version_folder_with_source_is_covered(tmp_path):
    """§14.8: a `<program>-<version>/` folder that has a SOURCE.md counts as covered."""
    exe = b"python installer"
    root = lint_tree(tmp_path / "t", files={
        "setup/README.md": SETUP_README,
        "setup/windows/README.md": _readme("python-3.12.4/"),
        "setup/windows/python-3.12.4/python-3.12.4-amd64.exe": exe,
        "setup/windows/python-3.12.4/SOURCE.md": _src(exe, "python-3.12.4-amd64.exe"),
    })
    assert by_rule(lint(root), "B014") == []


def test_b014_no_setup_no_finding(tmp_path):
    """§14.8: no finding without a setup/ directory."""
    root = lint_tree(tmp_path / "t", files={"docs/README.md": "x"})
    r = lint(root)
    assert by_rule(r, "B014") == []
    assert r["findings"] == []


# ====================================================================== §14.9 render_office TITLE
def _sized_deck(path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for frames in ([("1 BACKGROUND", 11), ("MCP is a protocol, not a tool", 30), ("3", 10)],
                   [("DSH 0.1.5-rc.2 at c291e7961a", 10), ("Every layer is a plugin", 28),
                    ("Unloading a plugin removes what it registered", 18)]):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, (text, size) in enumerate(frames):
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3 + i), Inches(8), Inches(0.6))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(size)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def test_text_backend_title_is_largest_font(tmp_path):
    """§14.9: `TITLE:` is the frame with the largest font (as deck_inspect), then the other frames in order."""
    src = _sized_deck(tmp_path / "src" / "capsule.pptx")
    out = tmp_path / "render"
    r = call("render_office", src=str(src), out_dir=str(out), backend="text")
    assert r["backend"] == "text"
    s1 = (out / "slide-01.txt").read_text(encoding="utf-8").rstrip("\n")
    s2 = (out / "slide-02.txt").read_text(encoding="utf-8").rstrip("\n")
    assert s1 == "TITLE: MCP is a protocol, not a tool\n\n1 BACKGROUND\n\n3"
    assert s2 == ("TITLE: Every layer is a plugin\n\nDSH 0.1.5-rc.2 at c291e7961a\n\n"
                  "Unloading a plugin removes what it registered")


def test_text_backend_title_matches_deck_inspect(tmp_path):
    """§14.9: the TITLE line equals deck_inspect's title for every slide."""
    src = _sized_deck(tmp_path / "src" / "capsule.pptx")
    out = tmp_path / "render"
    call("render_office", src=str(src), out_dir=str(out), backend="text")
    titles = [s["title"] for s in call("deck_inspect", pptx_path=str(src))["slides"]]
    got = [Path(out / f"slide-{i:02d}.txt").read_text(encoding="utf-8").split("\n", 1)[0]
           for i in range(1, len(titles) + 1)]
    assert got == [f"TITLE: {t}" for t in titles]
