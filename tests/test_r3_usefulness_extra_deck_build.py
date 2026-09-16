"""MANIFEST §14 usefulness fixes: deck_build (§14.5).

Companion to the other tests/test_r3_usefulness_*.py files (helpers in helpers_usefulness.py).
"""
from __future__ import annotations

import hashlib  # noqa: F401
import json  # noqa: F401
import re  # noqa: F401
import sys
from pathlib import Path

import pytest  # noqa: F401

sys.path.insert(0, str(Path(__file__).resolve().parent))
import helpers_usefulness as h  # noqa: E402
from helpers_usefulness import gitenv  # noqa: E402,F401  (fixture)


# ===================================================================== from the deck_build group

TOL = 0.03


@pytest.fixture(autouse=True)
def _office():
    h.need_office()


def pictures(sl):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return [sh for sh in h.iter_shapes(sl.shapes) if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def tables(sl):
    return [sh for sh in h.iter_shapes(sl.shapes) if getattr(sh, "has_table", False) and sh.has_table]


def charts(sl):
    return [sh.chart for sh in h.iter_shapes(sl.shapes) if getattr(sh, "has_chart", False) and sh.has_chart]


def build_err(tmp_path, spec):
    with pytest.raises(h.tool_error()) as exc:
        h.call("deck_build", spec=spec, out=str(tmp_path / "bad.pptx"))
    return str(exc.value)


def test_block_list_with_inserts_keeps_numbering(tmp_path):
    """§14.5 with §3.2: block lists do not change numbering or timing."""
    body = [{"kind": "point", "text": "One"}, {"kind": "point", "text": "Two"}]
    spec = h.deck_spec(h.content(body=body), h.content(body=body, insert=True, time="0:20"), h.content(body=body))
    res, prs = h.build(tmp_path, spec)
    assert [s["number"] for s in res["slides"]] == ["1", "1a", "2"]
    assert res["core_seconds"] == 60 and res["insert_seconds"] == 20


def test_inspect_sees_every_block(tmp_path):
    """§14.5 with §3.4: every block's text is on the slide."""
    body = [{"kind": "point", "text": "Models propose"}, {"kind": "bullets", "items": ["Code checks"]},
            {"kind": "table", "rows": [["Freeze", "hash"]]}]
    h.build(tmp_path, h.deck_spec(h.content(body=body)))
    texts = h.call("deck_inspect", pptx_path=str(tmp_path / "out.pptx"))["slides"][0]["texts"]
    joined = "\n".join(texts)
    for want in ("Models propose", "Code checks", "Freeze", "hash"):
        assert want in joined


def test_point_is_one_line(tmp_path):
    """§14.5: point is 1 bold paragraph."""
    _, prs = h.build(tmp_path, h.deck_spec(h.content(body={"kind": "point", "text": "Gates never move"})))
    tf = h.text_shape(prs.slides[0], "Gates never move").text_frame
    assert len([p for p in tf.paragraphs if p.text.strip()]) == 1
    runs = [r for p in tf.paragraphs for r in p.runs if r.text.strip()]
    assert runs and all(r.font.bold is True or (r.font.bold is None and p.font.bold is True)
                        for p in tf.paragraphs for r in p.runs if r.text.strip())


def test_block_size_must_be_number(tmp_path):
    """§14.5 with §13.4: a non-numeric size is a listed problem."""
    msg = build_err(tmp_path, h.deck_spec(h.content(body=[{"kind": "point", "text": "x", "size": "big"}])))
    assert "slides[0].body" in msg


def test_wide_figure_under_block_not_enlarged(tmp_path):
    """§14.5: a wide figure that already fits keeps the §3.2 width limit."""
    png = h.make_png(tmp_path / "wide.png", 2400, 400)
    body = [{"kind": "point", "text": "Wide chart"}, {"kind": "figure", "path": str(png)}]
    _, prs = h.build(tmp_path, h.deck_spec(h.content(body=body)))
    p = h.box(pictures(prs.slides[0])[0])
    assert (p[2] - p[0]) <= 11.8 + TOL
    assert (p[2] - p[0]) / (p[3] - p[1]) == pytest.approx(6.0, rel=0.02)
