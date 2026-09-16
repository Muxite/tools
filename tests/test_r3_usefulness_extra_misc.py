"""MANIFEST §14 usefulness fixes: diagram/chart, papers, bundle_lint and render_office (§14.6-§14.9).

Companion to the other tests/test_r3_usefulness_*.py files (helpers in helpers_usefulness.py).
"""
from __future__ import annotations

import hashlib  # noqa: F401
import json  # noqa: F401
import re  # noqa: F401
import sys
from pathlib import Path

import pytest  # noqa: F401

sys.path.insert(0, str(Path(__file__).resolve().parent))
import helpers_usefulness as h  # noqa: E402
from helpers_usefulness import gitenv  # noqa: E402,F401  (fixture)


# ===================================================================== from the diagram_chart group

def render(spec, **kw):
    return h.call("diagram_render", spec=spec, **kw)


def nodes(*ids, **fixed):
    out = []
    for i in ids:
        d = {"id": i, "label": i.replace("_", " ")}
        if i in fixed:
            d["rank"] = fixed[i]
        out.append(d)
    return out


def edges(*pairs):
    return [{"from": a, "to": b} for a, b in pairs]


def ranks(res):
    return {k: v["rank"] for k, v in res["nodes"].items()}


CHAIN = ("goal", "intent", "requirements", "plan", "freeze", "dispatch", "gates")


def chain_spec(**extra):
    spec = {"nodes": nodes(*CHAIN), "edges": edges(*zip(CHAIN, CHAIN[1:]))}
    spec.update(extra)
    return spec


def chart(**spec):
    res = h.call("chart_bar", spec=spec)
    return res, h.svg(res["svg"])


def test_rank_zero_explicit_matches_default():
    """§14.6: rank 0 on a source changes nothing."""
    base = {"nodes": nodes("a", "b"), "edges": edges(("a", "b"))}
    fixed = {"nodes": nodes("a", "b", a=0), "edges": edges(("a", "b"))}
    assert ranks(render(base)) == ranks(render(fixed)) == {"a": 0, "b": 1}


def test_wrap_larger_than_ranks_is_one_row():
    """§14.6: with wrap >= the number of ranks everything stays in 1 row (x increases with rank)."""
    res = render(chain_spec(wrap=20))
    b = h.boxes(res)
    for left, right in zip(CHAIN, CHAIN[1:]):
        assert b[left][2] <= b[right][0]


def test_wrap_deterministic():
    """§4.2/§14.6: wrapped layouts are deterministic."""
    assert render(chain_spec(wrap=2))["svg"] == render(chain_spec(wrap=2))["svg"]


def test_png_fallback_mermaid(monkeypatch, tmp_path):
    """§14.6: the fallback also serves Mermaid input."""
    pytest.importorskip("pymupdf")
    h.hide_converters(monkeypatch, tmp_path)
    render_args = {"mermaid": "flowchart LR\n  a(Planner) --> b[Gate]\n", "png": str(tmp_path / "m.png")}
    h.call("diagram_render", **render_args)
    assert (tmp_path / "m.png").stat().st_size > 100


def test_axis_false_explicit():
    """§14.6: axis false draws no axis group."""
    _, root = chart(categories=["a"], values=[1], axis=False)
    assert h.with_class(root, "g", "axis") == []


def test_axis_bars_still_proportional():
    """§14.6 with §5: the axis does not change bar lengths."""
    res, _ = chart(categories=["a", "b"], values=[10, 40], axis=True)
    lengths = [b["length"] for b in res["bars"]]
    assert lengths[1] == pytest.approx(4 * lengths[0], abs=0.5)


def test_highlight_color_without_highlight():
    """§14.6: highlight_color alone highlights nothing."""
    _, root = chart(categories=["a", "b"], values=[1, 2], highlight_color="failed")
    assert {h.fill(b) for b in h.with_class(root, "rect", "bar")} == {"#9a9a9a"}


def test_newline_label_bar_still_reported():
    """§14.6/§5: the bars list still carries every bar; the label keeps its text."""
    res, root = chart(categories=["kept\ntools", "failed"], values=[222, 215])
    assert [b["value"] for b in res["bars"]] == [222, 215]
    assert len(h.with_class(root, "rect", "bar")) == 2


# ===================================================================== from the papers group

def body(pid, d, **kw):
    return h.call("papers_body", id=pid, dir=str(d), **kw)


def grep(pid, d, pattern, **kw):
    return h.call("papers_peek", id=pid, mode="grep", dir=str(d), pattern=pattern, **kw)


def test_appendix_false_is_default(tmp_path):
    """§14.7: appendix false keeps the §9.2 default."""
    h.ff_txt(tmp_path / "2310.03714.txt", ["A", "B", "C", "D\nReferences\n[1] y", "Appendix\nE"])
    assert body("2310.03714", tmp_path, appendix=False)["end"] == 4


def test_width_larger_than_text_keeps_text(tmp_path):
    """§14.7: a hit shorter than width is unchanged."""
    h.marker_txt(tmp_path / "2604.00392.txt", ["short line with gate\n"])
    full = grep("2604.00392", tmp_path, "gate", context=0)["hits"][0]["text"]
    cut = grep("2604.00392", tmp_path, "gate", context=0, width=500)["hits"][0]["text"]
    assert cut == full == "short line with gate"


def test_width_keeps_page_and_line(tmp_path):
    """§14.7: width changes only `text`."""
    h.marker_txt(tmp_path / "2604.00392.txt", ["p1\n", "filler\n" + "y" * 200 + " arbor " + "y" * 200 + "\n"])
    a = grep("2604.00392", tmp_path, "arbor")["hits"][0]
    b = grep("2604.00392", tmp_path, "arbor", width=40)["hits"][0]
    assert (a["page"], a["line"]) == (b["page"], b["line"]) == (2, 9)


def test_four_spaces_are_not_layout(tmp_path):
    """§14.7: the run must be 5 or more spaces."""
    h.marker_txt(tmp_path / "2401.00002.txt", ["A    B\nC    D\nE    F\n"])
    assert body("2401.00002", tmp_path)["layout_text"] is False


def test_reextract_when_text_missing(tmp_path):
    """§14.7: reextract with no .txt simply extracts it."""
    pytest.importorskip("pymupdf")
    h.make_pdf(tmp_path / "2308.00352.pdf", ["metagpt page one", "metagpt page two"])
    r = h.call("papers_fetch", ids=["2308.00352"], dir=str(tmp_path), base_url="http://127.0.0.1:9/", delay=0,
               reextract=True)
    assert r["results"][0]["pages"] == 2
    t = (tmp_path / "2308.00352.txt").read_text(encoding="utf-8")
    assert "\n\n===== page 2 =====\n" in t and "metagpt page two" in t


def test_list_cli_missing_summary(tmp_path, capsys):
    """§14.7 through `tundlekit papers list`."""
    h.marker_txt(tmp_path / "2402.00001.txt", ["x\n"])
    code, data, _, _ = h.run_cli(capsys, ["papers", "list", "--dir", tmp_path, "--json"])
    assert code == 0
    assert data["missing_summary"] == ["2402.00001"]
    assert data["papers"][0]["layout_text"] is False


# ===================================================================== from the bundle_render group

def lint(root, **kw):
    r = h.call("bundle_lint", root=str(root), **kw)
    h.check_shape(r)
    return r


def severities(r, rule):
    return {f["path"].rstrip("/"): f["severity"] for f in h.of(r, rule)}


def _table(*names):
    return "| Item | Note |\n|---|---|\n" + "".join(f"| `{n}` | x |\n" for n in names)


def b014(root):
    f = h.of(lint(root), "B014")
    assert len(f) <= 1
    return f


def _deck(path, slides):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for frames in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, (text, size) in enumerate(frames):
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3 + 0.9 * i), Inches(8), Inches(0.6))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            if size is not None:
                r.font.size = Pt(size)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def slide_text(out, n):
    return (out / f"slide-{n:02d}.txt").read_text(encoding="utf-8").rstrip("\n")


def test_not_a_git_repo_stays_warning(tmp_path, gitenv):
    """§14.8: outside a git work tree nothing is ignored, so junk stays a warning."""
    root = h.tundle_tree(tmp_path / "t", init_git=False, files={
        ".gitignore": "__pycache__/\n", "papers/README.md": "x", "papers/__pycache__/a.pyc": "x"})
    assert severities(lint(root), "B005") == {"papers/__pycache__": "warning"}


def test_ignore_option_still_drops_b005(tmp_path, gitenv):
    """§13.1 with §14.8: `ignore` drops B005 findings whatever their severity."""
    root = h.tundle_tree(tmp_path / "t", files={".gitignore": "*.pyc\n", "tools/x.pyc": "x", "tools/y.tmp": "y"})
    assert h.of(lint(root, ignore=["B005"]), "B005") == []


def test_b014_only_readmes(tmp_path):
    """§14.8: README.md files are not installers; a count of 0 gives no finding."""
    root = h.tundle_tree(tmp_path / "t", files={"setup/README.md": "# Setup\n",
                                                 "setup/windows/README.md": _table()})
    assert b014(root) == []


def test_b014_ignore(tmp_path):
    """§13.1 with §14.8: `ignore` drops B014."""
    root = h.tundle_tree(tmp_path / "t", files={"setup/README.md": "x", "setup/w/README.md": _table("a.exe"),
                                                 "setup/w/a.exe": "a"})
    assert len(h.of(lint(root), "B014")) == 1
    assert h.of(lint(root, ignore=["B014"]), "B014") == []


def test_title_multi_slide_mixed(tmp_path):
    """§14.9: each slide picks its own largest frame."""
    src = _deck(tmp_path / "in" / "d.pptx", [
        [("THE EVIDENCE", 11), ("Arbor merges only through a held-out gate", 30)],
        [("Capability capsules", 40), ("From goal text to gates", 24), ("Muk", 14)],
    ])
    out = tmp_path / "out"
    h.call("render_office", src=str(src), out_dir=str(out), backend="text")
    assert slide_text(out, 1).split("\n")[0] == "TITLE: Arbor merges only through a held-out gate"
    assert slide_text(out, 2).split("\n")[0] == "TITLE: Capability capsules"
    assert "THE EVIDENCE" in slide_text(out, 1)
