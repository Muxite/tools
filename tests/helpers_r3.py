"""Shared helpers for the review-round-3 tests (MANIFEST §15.4, §15.6-§15.10).

Not a conftest: other suites write into tests/ in parallel. Test modules import it with
`import helpers_r3 as r3` (pytest puts tests/ on sys.path). Nothing here imports tundlekit at module
level, so collection succeeds before the tools exist. Each tool call imports only the module that owns
the tool (§15.10), then goes through `registry.call` (§1).
"""
from __future__ import annotations

import hashlib
import importlib
import json
import os
import re
import shutil
import subprocess
import time
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
SKILLS_DIR = REPO / "skills"
TUNDLE = Path(r"C:\Users\m50066326\Downloads\tundle")   # real material, read only

# §15.10: owning module of each new tool
OWNER = {
    "deck_diff": "deck",
    "docx_diff": "textlint",
    "translate_terms": "translate",
    "bundle_source": "bundle",
    "bundle_setup_table": "bundle",
    "papers_summary": "papers",
    "papers_index_check": "papers",
    # existing tools used to build or check fixtures
    "deck_build": "deck",
    "deck_inspect": "deck",
    "bundle_lint": "bundle",
    "bundle_verify": "bundle",
}


# ------------------------------------------------------------------------------------ tool access
def registry():
    return importlib.import_module("tundlekit.registry")


def tool_error():
    return registry().ToolError


def call(name: str, **args):
    """Import the owning module (§15.10), then registry.call (§1). Path values become strings."""
    importlib.import_module("tundlekit." + OWNER[name])
    clean = {}
    for k, v in args.items():
        if isinstance(v, Path):
            v = str(v)
        elif isinstance(v, list):
            v = [str(x) if isinstance(x, Path) else x for x in v]
        clean[k] = v
    return registry().call(name, clean)


def get_tool(name: str):
    importlib.import_module("tundlekit." + OWNER[name])
    return registry().TOOLS[name]


def run_cli(capsys, argv):
    """tundlekit.cli.main(argv) in-process (§0.4). Returns (code, parsed JSON or None, stdout, stderr)."""
    from tundlekit import cli

    capsys.readouterr()
    try:
        code = cli.main([str(a) for a in argv])
    except SystemExit as exc:  # argparse usage errors (§2.1a)
        code = exc.code if isinstance(exc.code, int) else 2
    out, err = capsys.readouterr()
    data = None
    if out.strip():
        try:
            data = json.loads(out)
        except ValueError:
            data = None
    return code, data, out, err


def check_shape(result):
    """§0.3 checker result shape."""
    assert isinstance(result["ok"], bool)
    assert isinstance(result["findings"], list)
    assert set(result["counts"]) >= {"error", "warning", "info"}
    for sev in ("error", "warning", "info"):
        assert result["counts"][sev] == sum(1 for f in result["findings"] if f["severity"] == sev)
    assert result["ok"] == (result["counts"]["error"] == 0)
    for f in result["findings"]:
        assert set(f) >= {"rule", "severity", "path", "line", "message", "excerpt"}
        assert f["severity"] in ("error", "warning", "info")
        assert "\\" not in f["path"]
        assert f["line"] is None or (isinstance(f["line"], int) and f["line"] >= 1)
    keys = [(f["path"], f["line"] or 0, f["rule"]) for f in result["findings"]]
    assert keys == sorted(keys)


def by_rule(result, rule):
    return [f for f in result["findings"] if f["rule"] == rule]


def write(path: Path, text: str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    return path


def slash(p) -> str:
    return str(p).replace("\\", "/")


def hint_matches(hint: str, path: Path, line: int) -> bool:
    """§15.4: a hint is `path:line`. The path form is not pinned, so compare by file name and line."""
    h = slash(hint)
    return h.endswith(f"/{path.name}:{line}") or h == f"{path.name}:{line}"


def need_pptx():
    pytest.importorskip("pptx")
    pytest.importorskip("PIL")


def need_git():
    if shutil.which("git") is None:
        pytest.skip("git not installed")


# ------------------------------------------------------------------------------------ decks (§3, §15.4)
SAY = " ".join(f"word{i}" for i in range(25))


def content(title, time_="0:30", say=SAY, **kw):
    s = {"type": "content", "eyebrow": "Research", "title": title,
         "body": {"kind": "bullets", "items": ["first point", "second point"]},
         "notes": {"time": time_, "say": say}}
    s.update(kw)
    return s


def title_slide(title):
    return {"type": "title", "title": title, "subtitle": "Review round 3", "byline": "A. Speaker",
            "notes": {"time": "0:15", "say": "Hello and welcome."}}


def base_spec():
    """Title slide, then content slides numbered 2..5 (§3.2 numbering)."""
    return {"meta": {"id": "capsule"}, "slides": [
        title_slide("Capability capsules for AI research"),
        content("Kept tools fail held-out tests", source="Beyond Task Completion, Table 4",
                say="Most tools that the agent kept did not pass the held-out tests that we wrote "
                    "before the run started, so keeping a tool is not evidence that it works."),
        content("Execution Broker isolates side effects",
                body={"kind": "lines", "items": ["one broker per run",
                                                 "the broker leases each capability for one run only"]}),
        content("Library growth needs retirement"),
        content("Gates decide what enters the library"),
    ]}


def build_deck(path: Path, spec=None, **kw) -> Path:
    need_pptx()
    call("deck_build", spec=spec or base_spec(), out=str(path), **kw)
    return path


def open_deck(path):
    from pptx import Presentation

    return Presentation(str(path))


def _frames(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            elif sh.has_text_frame:
                yield sh.text_frame
    return list(walk(slide.shapes))


def set_frame_text(tf, new):
    """Replace a frame's text while keeping the first run's formatting."""
    paras = tf.paragraphs
    first = paras[0]
    runs = first.runs
    runs[0].text = new
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def slide_by_number(prs, number):
    for s in prs.slides:
        if any(tf.text.strip() == number for tf in _frames(s)):
            return s
    raise AssertionError(f"no slide numbered {number}")


def edit_frame(src: Path, dst: Path, number: str, old: str, new: str) -> Path:
    """Copy src to dst with the frame whose whole text is `old` (on slide `number`) set to `new`."""
    prs = open_deck(src)
    s = slide_by_number(prs, number) if number else prs.slides[0]
    for tf in _frames(s):
        if tf.text == old:
            set_frame_text(tf, new)
            break
    else:
        raise AssertionError(f"no frame {old!r}")
    prs.save(str(dst))
    return dst


def edit_notes(src: Path, dst: Path, number: str, new_notes: str) -> Path:
    prs = open_deck(src)
    slide_by_number(prs, number).notes_slide.notes_text_frame.text = new_notes
    prs.save(str(dst))
    return dst


def set_hidden(src: Path, dst: Path, number: str, hidden: bool) -> Path:
    prs = open_deck(src)
    el = slide_by_number(prs, number)._element
    if hidden:
        el.set("show", "0")
    elif "show" in el.attrib:
        del el.attrib["show"]
    prs.save(str(dst))
    return dst


def remove_slide(src: Path, dst: Path, number: str) -> Path:
    prs = open_deck(src)
    target = slide_by_number(prs, number)
    lst = prs.slides._sldIdLst
    for sld_id in list(lst):
        if prs.part.related_part(sld_id.rId) is target.part:
            prs.part.drop_rel(sld_id.rId)
            lst.remove(sld_id)
            break
    prs.save(str(dst))
    return dst


def move_slide(src: Path, dst: Path, old_index: int, new_index: int) -> Path:
    """Move the slide at 0-based old_index to new_index."""
    prs = open_deck(src)
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[old_index]
    lst.remove(el)
    lst.insert(new_index, el)
    prs.save(str(dst))
    return dst


def add_slide(src: Path, dst: Path, texts, notes=None) -> Path:
    from pptx.util import Inches, Pt

    prs = open_deck(src)
    layout = prs.slides[-1].slide_layout
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    y = 0.4
    for text, size in texts:
        tb = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(10), Inches(0.8))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(size)
        y += 1.0
    if notes is not None:
        s.notes_slide.notes_text_frame.text = notes
    prs.save(str(dst))
    return dst


def changes_of(result, field):
    return [c for c in result["changes"] if c["field"] == field]


# ------------------------------------------------------------------------------------ docx (§15.4)
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def make_docx(path: Path, paragraphs) -> Path:
    """A minimal .docx written by hand. Each paragraph is a str, or a list of run texts, or
    (style_id, [runs])."""
    body = []
    for p in paragraphs:
        style = None
        if isinstance(p, tuple):
            style, runs = p
        elif isinstance(p, str):
            runs = [p]
        else:
            runs = p
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        rs = "".join(f'<w:r><w:t xml:space="preserve">{_esc(r)}</w:t></w:r>' for r in runs if r != "")
        body.append(f"<w:p>{ppr}{rs}</w:p>")
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f'<w:document xmlns:w="{_W}"><w:body>{"".join(body)}<w:sectPr/></w:body></w:document>')
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>")
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
    return path


# ------------------------------------------------------------------------------------ bundle (§15.7)
def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def local_noon(y, m, d) -> float:
    return time.mktime((y, m, d, 12, 0, 0, 0, 0, -1))


def installer(path: Path, data: bytes = b"MZ fake installer\n", date=(2026, 3, 14)) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    ts = local_noon(*date)
    os.utime(path, (ts, ts))
    return path


def source_text(program, version, name, date, digest, url=None, install=None) -> str:
    """The §15.7 SOURCE.md body, without assuming anything about a trailing newline."""
    head = f"# {program} {version}" if version else f"# {program}"
    return "\n".join([
        head,
        "",
        f"- File:       {name}",
        f"- Source:     {url if url is not None else '<official download URL>'}",
        f"- Downloaded: {date}",
        f"- SHA-256:    {digest}",
        f"- Install:    {install if install is not None else '<steps>'}",
    ])


def norm_text(t: str) -> str:
    return t.replace("\r\n", "\n").rstrip("\n")


CHANGELOG = "# Changelog\n\n<!-- entries -->\n\n## 2026.09.16.1  (2026-09-16 10:00)\n\nx\n\n_1 added, 0 modified, 0 removed, 0 renamed_\n"


def make_tundle(root: Path) -> Path:
    """A lint-clean tundle root (§2.7 B010 satisfied), under git (§14.8 uses git check-ignore)."""
    need_git()
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    subprocess.run(["git", "init", "-q", "-b", "main", str(root)], check=True, capture_output=True, timeout=60)
    write(root / "VERSION", "2026.09.16.1\n")
    write(root / "CHANGELOG.md", CHANGELOG)
    write(root / "setup" / "README.md", "# Setup\n")
    return root


def b006_paths(root: Path):
    r = call("bundle_lint", root=str(root))
    return sorted(f["path"].rstrip("/") for f in r["findings"] if f["rule"] == "B006")


def row_cells(row: str):
    return [c.strip() for c in row.strip().strip("|").split("|")]


# ------------------------------------------------------------------------------------ papers (§9, §15.8)
def marker_txt(path: Path, pages) -> Path:
    """§9 fetch text format."""
    return write(path, "".join(f"\n\n===== page {n} =====\n{body}" for n, body in enumerate(pages, 1)))


def ff_txt(path: Path, pages) -> Path:
    """pdftotext text format (form feeds)."""
    return write(path, "\f".join(pages) + "\f")


def paper_pages(page1, n_pages, refs_page=None):
    pages = [page1]
    for i in range(2, n_pages + 1):
        body = f"Section text on page {i}.\nMore words here.\n"
        if refs_page == i:
            body = f"Conclusion text.\nReferences\n[1] Someone. A paper. 2024.\n"
        pages.append(body)
    return pages


SUMMARY_BODY = ("## Summary\nText.\n\n## How it works (p2-4)\n- a\n\n## Results (p5)\n- b\n\n"
                "## Limitations\n- c\n\n## Relevance\nd\n")


def summary_file(papers_dir: Path, pid: str, short: str, body: str = SUMMARY_BODY) -> Path:
    return write(Path(papers_dir) / "summaries" / f"{pid} - {short}.md",
                 f"# {pid} · {short}\n\nA. Author · arXiv {pid} · 10 pp (8 body)\nRead: p1-8\n\n{body}")


def skeleton(pid, title, authors, pages, body) -> str:
    return "\n".join([
        f"# {pid} · {title}",
        "",
        f"{authors} · arXiv {pid} · {pages} pp ({body} body)",
        "Read: <pages read>",
        "",
        "## Summary",
        "",
        "## How it works",
        "",
        "## Results",
        "",
        "## Limitations",
        "",
        "## Relevance",
    ])


# ------------------------------------------------------------------------------------ markdown (§11)
def parse_frontmatter(text: str):
    text = text.lstrip("\ufeff").replace("\r\n", "\n")
    if not text.startswith("---\n"):
        return None, text
    end = re.search(r"^---[ \t]*$", text[4:], re.M)
    if not end:
        return None, text
    head, body = text[4:4 + end.start()], text[4 + end.end():]
    fields = {}
    lines = head.split("\n")
    i = 0
    while i < len(lines):
        m = re.match(r"^([A-Za-z0-9_-]+):[ \t]*(.*)$", lines[i])
        i += 1
        if not m:
            continue
        key, val = m.group(1), m.group(2).strip()
        if val in ("|", ">", "|-", ">-", "|+", ">+"):
            block = []
            while i < len(lines) and (lines[i].startswith((" ", "\t")) or not lines[i].strip()):
                block.append(lines[i].strip())
                i += 1
            val = (" " if val.startswith(">") else "\n").join(block).strip()
        else:
            while i < len(lines) and lines[i].startswith((" ", "\t")) and lines[i].strip():
                val += " " + lines[i].strip()
                i += 1
            if len(val) >= 2 and val[0] == val[-1] and val[0] in "\"'":
                val = val[1:-1]
        fields[key] = val
    return fields, body


LINK_RE = re.compile(r"!?\[[^\]]*\]\(\s*<?([^)\s>]+)>?(?:\s+\"[^\"]*\")?\s*\)")


def relative_links(md: str):
    out = []
    for t in LINK_RE.findall(md):
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", t) or t.startswith(("#", "/")):
            continue
        out.append(t)
    return out


def table_headers(md: str):
    """Header rows of Markdown tables, as lists of cells with emphasis and code marks removed."""
    lines = md.replace("\r\n", "\n").split("\n")
    out = []
    for i, line in enumerate(lines[:-1]):
        nxt = lines[i + 1].strip()
        if line.strip().startswith("|") and re.match(r"^\|?\s*:?-{3,}", nxt):
            cells = [re.sub(r"[*_`]", "", c).strip() for c in line.strip().strip("|").split("|")]
            out.append(cells)
    return out


def cli_command_exists(group: str, command: str) -> bool:
    import contextlib
    import io

    from tundlekit import cli

    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
            code = cli.main([group, command, "--help"])
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else (0 if exc.code is None else 2)
    return code == 0
