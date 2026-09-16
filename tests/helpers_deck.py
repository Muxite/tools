"""Shared helpers for the deck / chart / palette tests (MANIFEST §3, §5, §6).

Not a conftest: other suite authors write into tests/ in parallel. Import with
`from helpers_deck import ...` (pytest puts this directory on sys.path).
tundlekit is imported lazily so that collection succeeds before the package exists.
"""
from __future__ import annotations

import copy
import importlib
import json
import pathlib
import re
import xml.etree.ElementTree as ET

import pytest

EMU_PER_INCH = 914400
SVG_NS = "http://www.w3.org/2000/svg"

# §6 values, restated here so the tests do not depend on the module under test
STAGE = {
    "intent": "#0072B2", "binding": "#2A8FA8", "freeze": "#7B3FA0", "dispatch": "#D98200",
    "gates": "#008A63", "claims": "#B8527F", "build": "#3D4FB0", "library": "#8C5A2B",
    "external": "#6b6b6b",
}

# 25 plain words: enough for D002 (>= 20) and within D001 for 0:30 at 2.3 w/s (69 words)
SAY25 = " ".join(f"word{i}" for i in range(25))

_MODULE_FOR_PREFIX = {"deck": "deck", "chart": "chart", "palette": "palette"}


# ------------------------------------------------------------------ tool access (§0.2, §1)
def registry():
    return importlib.import_module("tundlekit.registry")


def ToolError():  # noqa: N802 - returns the class
    return registry().ToolError


def call(name: str, **args):
    """Run a registered tool through registry.call (§1). Imports only the module that owns it (§0.1)."""
    importlib.import_module("tundlekit." + _MODULE_FOR_PREFIX[name.split("_")[0]])
    return registry().call(name, args)


def get_tool(name: str):
    importlib.import_module("tundlekit." + _MODULE_FOR_PREFIX[name.split("_")[0]])
    return registry().TOOLS[name]


def run_cli(argv, capsys):
    """Run tundlekit.cli.main (§0.4); returns (exit code, parsed JSON stdout or None, stderr)."""
    cli = importlib.import_module("tundlekit.cli")
    capsys.readouterr()
    code = cli.main([str(a) for a in argv])
    out = capsys.readouterr()
    data = None
    if out.out.strip():
        try:
            data = json.loads(out.out)
        except ValueError:
            data = None
    return code, data, out.err


def need_office():
    pytest.importorskip("pptx")
    pytest.importorskip("PIL")


# ------------------------------------------------------------------ spec builders (§3.1)
def content(title="A plain content title", time="0:30", say=SAY25, eyebrow="Research", body=None, **kw):
    s = {"type": "content", "eyebrow": eyebrow, "title": title,
         "body": body if body is not None else {"kind": "bullets", "items": ["first point", "second point"]},
         "notes": {"time": time, "say": say}}
    s.update(kw)
    return s


def title_slide(title="Deck title", time="0:15", say="Hello.", **kw):
    s = {"type": "title", "title": title, "subtitle": "The subtitle", "byline": "A. Speaker",
         "notes": {"time": time, "say": say}}
    s.update(kw)
    return s


def divider(title="Part two", time="0:05", say="", **kw):
    s = {"type": "divider", "title": title, "subtitle": "Divider subtitle", "notes": {"time": time, "say": say}}
    s.update(kw)
    return s


def spec(*slides, **meta):
    out = {"slides": [copy.deepcopy(s) for s in slides]}
    if meta:
        out["meta"] = meta
    return out


def write_json(path: pathlib.Path, obj) -> pathlib.Path:
    path.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    return path


def make_png(path: pathlib.Path, w: int, h: int, fmt="PNG") -> pathlib.Path:
    from PIL import Image

    Image.new("RGB", (w, h), (200, 30, 30)).save(str(path), fmt)
    return path


# ------------------------------------------------------------------ deck build + read back
def build(tmp_path, the_spec, name="out.pptx", **kw):
    """Build via deck_build (§3.2); returns (result, python-pptx Presentation)."""
    need_office()
    from pptx import Presentation

    out = tmp_path / name
    res = call("deck_build", spec=the_spec, out=str(out), **kw)
    return res, Presentation(str(out))


def iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def frames(slide):
    return [sh.text_frame for sh in iter_shapes(slide.shapes) if sh.has_text_frame]


def frame_texts(slide):
    return [f.text for f in frames(slide)]


def table_cells(slide):
    out = []
    for sh in iter_shapes(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.extend(c.text for c in row.cells)
    return out


def all_text(slide):
    return "\n".join(frame_texts(slide) + table_cells(slide))


NUMBER_RE = re.compile(r"^\d+[a-z]?$")


def numbers_on(slide):
    """Texts of frames whose whole text is a slide number (§3.3 pptx rule)."""
    return [t for t in frame_texts(slide) if NUMBER_RE.match(t.strip())]


def notes_text(slide):
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text


def is_hidden(slide):
    return slide._element.get("show") == "0"


def pictures(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return [sh for sh in iter_shapes(slide.shapes) if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def charts(slide):
    return [sh.chart for sh in iter_shapes(slide.shapes) if getattr(sh, "has_chart", False) and sh.has_chart]


def tables(slide):
    return [sh.table for sh in iter_shapes(slide.shapes) if getattr(sh, "has_table", False) and sh.has_table]


def frame_is_bold(tf):
    runs = [r for p in tf.paragraphs for r in p.runs if r.text.strip()]
    if not runs:
        return False
    for p in tf.paragraphs:
        for r in p.runs:
            if r.text.strip() and not (r.font.bold is True or (r.font.bold is None and p.font.bold is True)):
                return False
    return True


MONO_HINTS = ("consolas", "courier", "mono", "menlo", "monaco", "lucida console", "inconsolata",
              "source code", "fira code", "cascadia")


def is_mono(name):
    return bool(name) and any(h in name.lower() for h in MONO_HINTS)


def point_fill(series, idx):
    """srgb fill (upper-case hex) of point idx: its dPt fill, else the series fill, else None."""
    ser = series._element
    vals = ser.xpath(f'./c:dPt[c:idx/@val="{idx}"]/c:spPr/a:solidFill/a:srgbClr/@val')
    if vals:
        return str(vals[0]).upper()
    vals = ser.xpath("./c:spPr/a:solidFill/a:srgbClr/@val")
    return str(vals[0]).upper() if vals else None


def data_labels_on(chart):
    vals = chart._chartSpace.xpath(".//c:dLbls/c:showVal/@val")
    return any(str(v) in ("1", "true") for v in vals)


# ------------------------------------------------------------------ hand-made pptx (for lint / inspect)
def make_pptx(path, slides):
    """slides: list of dicts with keys
    texts: [(text, size_pt)] text boxes in shape order; table: [[str]] (added first if table_first);
    notes: str or None; hidden: bool."""
    need_office()
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for d in slides:
        s = prs.slides.add_slide(blank)

        def add_table():
            rows = d["table"]
            t = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(4), Inches(6), Inches(1)).table
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell

        if d.get("table") and d.get("table_first"):
            add_table()
        y = 0.3
        for text, size in d.get("texts", []):
            tb = s.shapes.add_textbox(Inches(0.5), Inches(y), Inches(10), Inches(0.6))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(size)
            y += 0.7
        if d.get("table") and not d.get("table_first"):
            add_table()
        if d.get("notes") is not None:
            s.notes_slide.notes_text_frame.text = d["notes"]
        if d.get("hidden"):
            s._element.set("show", "0")
    prs.save(str(path))
    return path


# ------------------------------------------------------------------ findings (§0.3)
def rules(result, rule=None):
    fs = result["findings"]
    return [f for f in fs if rule is None or f["rule"] == rule]


def check_shape(result):
    assert isinstance(result["ok"], bool)
    assert set(result["counts"]) >= {"error", "warning", "info"}
    for f in result["findings"]:
        assert set(f) >= {"rule", "severity", "path", "line", "message", "excerpt"}
        assert f["severity"] in ("error", "warning", "info")
        assert f["line"] is None or (isinstance(f["line"], int) and f["line"] >= 1)
    for sev in ("error", "warning", "info"):
        assert result["counts"][sev] == sum(1 for f in result["findings"] if f["severity"] == sev)
    assert result["ok"] == (result["counts"]["error"] == 0)
    keys = [(f["path"], f["line"] or 0, f["rule"]) for f in result["findings"]]
    assert keys == sorted(keys)


# ------------------------------------------------------------------ SVG (§5)
def svg_root(text):
    root = ET.fromstring(text)
    assert root.tag == f"{{{SVG_NS}}}svg"
    return root


def classes(el):
    return (el.get("class") or "").split()


def find_class(root, tag, cls):
    return [e for e in root.iter(f"{{{SVG_NS}}}{tag}") if cls in classes(e)]


def text_of(el):
    return "".join(el.itertext())


def all_svg_text(root):
    return [text_of(e) for e in root.iter(f"{{{SVG_NS}}}text")]


def fill_of(el):
    f = el.get("fill")
    if f is None:
        m = re.search(r"fill\s*:\s*([^;]+)", el.get("style") or "")
        f = m.group(1).strip() if m else None
    return f.lower() if f else None


def chart_svg(**args):
    res = call("chart_bar", **args)
    return res, svg_root(res["svg"])
