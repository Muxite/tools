"""Shared helpers for the render / papers / translate / skills / docs tests (MANIFEST §8-§12).

Not a conftest: other suites write into tests/ in parallel. Test modules add this directory to
sys.path and import it. Nothing here imports tundlekit at module level, so collection succeeds
before the package is implemented.
"""
from __future__ import annotations

import contextlib
import io
import json
import os
import re
import shutil
import sys
import threading
import zipfile
from functools import lru_cache
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SKILLS_DIR = REPO / "skills"

# §11 required skills
REQUIRED_SKILLS = [
    "tundle-bundle",
    "deck-builder",
    "diagram-maker",
    "report-writing",
    "deliverable-review",
    "paper-reading",
    "zh-en-translation",
    "held-out-build-gate",
]

# Groups listed in the manifest sections (§2-§10).
GROUPS = {"bundle", "deck", "diagram", "chart", "palette", "text", "render", "papers", "translate"}


# --------------------------------------------------------------------------------------------- tools / CLI

def registry():
    from tundlekit import registry as reg

    reg.load_all()
    return reg


def call_tool(name, args=None):
    return registry().call(name, args or {})


def tool_error():
    from tundlekit.registry import ToolError

    return ToolError


def run_cli(argv, capsys):
    """Run tundlekit.cli.main(argv) in-process. Returns (code, parsed_json_or_None, stdout, stderr).

    A SystemExit (argparse) is turned into its code.
    """
    from tundlekit import cli

    capsys.readouterr()
    try:
        code = cli.main([str(a) for a in argv])
    except SystemExit as exc:  # argparse usage errors / --help
        code = exc.code if isinstance(exc.code, int) else (0 if exc.code is None else 2)
    out, err = capsys.readouterr()
    data = None
    if out.strip():
        try:
            data = json.loads(out)
        except ValueError:
            data = None
    return code, data, out, err


@lru_cache(maxsize=None)
def cli_command_exists(group: str, command: str) -> bool:
    """§11/§12: a `tundlekit <group> <command>` is real when `<group> <command> --help` exits 0."""
    from tundlekit import cli

    buf_out, buf_err = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(buf_out), contextlib.redirect_stderr(buf_err):
            code = cli.main([group, command, "--help"])
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else (0 if exc.code is None else 2)
    except Exception:
        return False
    return code == 0


# --------------------------------------------------------------------------------------------- fixtures: files

def make_pdf(path: Path, texts, width=200, height=100):
    """A PDF with 1 page per entry in `texts` (pymupdf). Returns path."""
    import pymupdf

    doc = pymupdf.open()
    for t in texts:
        page = doc.new_page(width=width, height=height)
        page.insert_text((10, 30), t, fontsize=11)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return path


def pdf_bytes(texts, width=200, height=100) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for t in texts:
        page = doc.new_page(width=width, height=height)
        page.insert_text((10, 30), t, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_png(path: Path, w: int, h: int, color=(200, 30, 30)):
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (w, h), color).save(str(path))
    return path


def png_size(path) -> tuple[int, int]:
    from PIL import Image

    with Image.open(str(path)) as im:
        return im.size


def make_pptx(path: Path, slides):
    """slides: list of (texts: list[str], notes: str | None). Each text is its own text box, in order."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for texts, notes in slides:
        slide = prs.slides.add_slide(blank)
        for i, t in enumerate(texts):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5 + i), Inches(6), Inches(0.8))
            box.text_frame.text = t
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _xml_escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def make_docx(path: Path, paragraphs):
    """A minimal .docx written by hand. paragraphs: list of (style_id or None, [run texts])."""
    body = []
    for style, runs in paragraphs:
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        rs = "".join(f'<w:r><w:t xml:space="preserve">{_xml_escape(r)}</w:t></w:r>' for r in runs)
        body.append(f"<w:p>{ppr}{rs}</w:p>")
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}"><w:body>{"".join(body)}<w:sectPr/></w:body></w:document>'
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
    return path


def write_marker_txt(path: Path, pages, preamble: str = ""):
    """§9 fetch text format: `preamble`, then for each page N '\\n\\n===== page N =====\\n' + text."""
    text = preamble + "".join(f"\n\n===== page {n} =====\n{body}" for n, body in enumerate(pages, 1))
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    return text


def write_ff_txt(path: Path, pages, trailing_ff=True):
    """pdftotext-style text: pages separated by form feeds."""
    text = "\f".join(pages) + ("\f" if trailing_ff else "")
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    return text


def write_raw(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    return text


def snapshot(path: Path):
    st = path.stat()
    return path.read_bytes(), st.st_mtime_ns


# --------------------------------------------------------------------------------------------- arXiv stand-in

class ArxivServer:
    """A localhost HTTP server. `routes` maps a URL path ('/pdf/2401.00001') to (status, body bytes)."""

    def __init__(self):
        self.routes: dict[str, tuple[int, bytes]] = {}
        self.requests: list[dict] = []  # {"path", "headers", "time"}
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                import time as _t

                outer.requests.append({"path": self.path, "headers": dict(self.headers.items()),
                                       "time": _t.monotonic()})
                status, body = outer.routes.get(self.path, (404, b"not found"))
                self.send_response(status)
                self.send_header("Content-Type", "application/pdf" if body.startswith(b"%PDF-") else "text/html")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *a):
                pass

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.port = self.httpd.server_address[1]
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)

    def url(self, prefix="/pdf/"):
        return f"http://127.0.0.1:{self.port}{prefix}"

    def paths(self):
        return [r["path"] for r in self.requests]

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, *exc):
        self.httpd.shutdown()
        self.httpd.server_close()


def no_proxy(monkeypatch):
    for k in list(os.environ):
        if k.lower() in ("http_proxy", "https_proxy", "all_proxy"):
            monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    monkeypatch.setenv("no_proxy", "127.0.0.1,localhost")


# --------------------------------------------------------------------------------------------- safety

def forbid_writes(monkeypatch):
    """Make every common delete/copy entry point fail, for refusal tests that must never touch disk."""
    import pathlib

    def boom(*a, **k):
        raise AssertionError(f"filesystem modification attempted: {a!r}")

    for mod, name in [(shutil, "rmtree"), (shutil, "copy"), (shutil, "copy2"), (shutil, "copyfile"),
                      (shutil, "move"), (os, "remove"), (os, "unlink"), (os, "rmdir")]:
        monkeypatch.setattr(mod, name, boom)
    for name in ("unlink", "rmdir", "write_bytes", "write_text"):
        monkeypatch.setattr(pathlib.Path, name, boom)


# --------------------------------------------------------------------------------------------- markdown

FENCE_RE = re.compile(r"^[ \t]*(`{3,}|~{3,})")
INLINE_RE = re.compile(r"(`+)(.+?)\1")


def code_texts(md: str) -> list[str]:
    """Contents of fenced code blocks (whole blocks) and inline code spans outside them."""
    out, block, fence = [], [], None
    for line in md.splitlines():
        m = FENCE_RE.match(line)
        if fence is None:
            if m:
                fence = m.group(1)[0] * len(m.group(1))
                block = []
                continue
            out.extend(s.strip() for _, s in INLINE_RE.findall(line))
        else:
            if m and m.group(1).startswith(fence):
                out.append("\n".join(block))
                fence = None
                continue
            block.append(line)
    if fence is not None:
        out.append("\n".join(block))
    return out


def fenced_blocks(md: str) -> list[tuple[str, str]]:
    """(info string, content) for each fenced block."""
    out, block, fence, info = [], [], None, ""
    for line in md.splitlines():
        m = re.match(r"^[ \t]*(`{3,}|~{3,})(.*)$", line)
        if fence is None:
            if m:
                fence, info, block = m.group(1), m.group(2).strip(), []
            continue
        if m and m.group(1).startswith(fence) and not m.group(2).strip():
            out.append((info, "\n".join(block)))
            fence = None
            continue
        block.append(line)
    return out


def prose_without_code(md: str) -> str:
    lines, fence = [], None
    for line in md.splitlines():
        m = FENCE_RE.match(line)
        if fence is None:
            if m:
                fence = m.group(1)
                continue
            lines.append(INLINE_RE.sub(" ", line))
        elif m and m.group(1).startswith(fence):
            fence = None
    return "\n".join(lines)


CMD_RE = re.compile(r"(?<![\w./-])tundlekit[ \t]+([a-z][a-z0-9-]*)(?:[ \t]+([a-z][A-Za-z0-9_.-]*))?")


def cli_mentions(md: str) -> list[tuple[str, str | None]]:
    """(group, command) pairs of `tundlekit <group> <command>` shown in code in a Markdown text."""
    found = []
    for code in code_texts(md):
        for m in CMD_RE.finditer(code):
            found.append((m.group(1), m.group(2)))
    return found


def bad_cli_mentions(md: str) -> list[str]:
    """Mentions that are not real CLI commands (§11, §12)."""
    bad = []
    tools = registry().TOOLS
    for group, command in cli_mentions(md):
        if group == "tools":
            continue
        if group == "call":
            if command is not None and command not in tools:
                bad.append(f"tundlekit call {command}")
            continue
        if command is None:
            if group not in GROUPS:
                bad.append(f"tundlekit {group}")
            continue
        if not cli_command_exists(group, command):
            bad.append(f"tundlekit {group} {command}")
    return sorted(set(bad))


TOOL_SPAN_RE = re.compile(r"`([a-z][a-z0-9]*(?:_[a-z0-9]+)+)`")


def bad_tool_mentions(md: str) -> list[str]:
    tools = registry().TOOLS
    prefixes = {n.split("_")[0] for n in tools}
    bad = []
    for name in TOOL_SPAN_RE.findall(md):
        if name.split("_")[0] in prefixes and name not in tools:
            bad.append(name)
    return sorted(set(bad))


LINK_RE = re.compile(r"!?\[[^\]]*\]\(\s*<?([^)\s>]+)>?(?:\s+\"[^\"]*\")?\s*\)")


def relative_links(md: str) -> list[str]:
    links = []
    for target in LINK_RE.findall(prose_without_code(md)):
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target) or target.startswith(("#", "/")):
            continue
        links.append(target)
    return links


def parse_frontmatter(text: str):
    """Return (fields dict, body str) or (None, text). Supports plain, quoted and block (| >) scalars."""
    text = text.lstrip("\ufeff").replace("\r\n", "\n")
    if not text.startswith("---\n"):
        return None, text
    end = re.search(r"^---[ \t]*$", text[4:], re.M)
    if not end:
        return None, text
    head = text[4:4 + end.start()]
    body = text[4 + end.end():]
    if body.startswith("\n"):
        body = body[1:]
    fields, lines, i = {}, head.split("\n"), 0
    while i < len(lines):
        line = lines[i]
        m = re.match(r"^([A-Za-z0-9_-]+):[ \t]*(.*)$", line)
        i += 1
        if not m:
            continue
        key, val = m.group(1), m.group(2).strip()
        if val in ("|", ">", "|-", ">-", "|+", ">+"):
            block = []
            while i < len(lines) and (lines[i].startswith((" ", "\t")) or not lines[i].strip()):
                block.append(lines[i].strip())
                i += 1
            sep = "\n" if val.startswith("|") else " "
            val = sep.join(b for b in block).strip()
        else:
            # plain multi-line continuation
            while i < len(lines) and lines[i].startswith((" ", "\t")) and lines[i].strip():
                val += " " + lines[i].strip()
                i += 1
            if len(val) >= 2 and val[0] == val[-1] and val[0] in "\"'":
                val = val[1:-1]
        fields[key] = val
    return fields, body
