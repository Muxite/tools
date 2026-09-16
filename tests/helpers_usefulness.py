"""Helpers for the §14 usefulness tests in tests/test_r3_usefulness_extra_*.py.

Nothing here imports tundlekit (or an optional package) at module level.
"""
from __future__ import annotations

import importlib
import json
import os
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

SVG_NS = "http://www.w3.org/2000/svg"
EMU = 914400
SAY25 = " ".join(f"token{i}" for i in range(25))
MARKER = "<!-- entries -->"
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


# ---------------------------------------------------------------- tools / CLI
_OWNER = {"text": "textlint", "deck": "deck", "diagram": "diagram", "chart": "chart", "papers": "papers",
          "bundle": "bundle", "render": "render"}


def call(name, **args):
    """registry.call after importing only the owning module (§0.1); Path values become str."""
    importlib.import_module("tundlekit." + _OWNER[name.split("_")[0]])
    from tundlekit import registry

    return registry.call(name, {k: (str(v) if isinstance(v, Path) else v) for k, v in args.items()})


def tool_error():
    from tundlekit.registry import ToolError

    return ToolError


def run_cli(capsys, argv):
    """tundlekit.cli.main in-process -> (code, parsed JSON or None, stdout, stderr)."""
    from tundlekit import cli

    capsys.readouterr()
    try:
        code = cli.main([str(a) for a in argv])
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else 2
    out, err = capsys.readouterr()
    data = None
    if out.strip():
        try:
            data = json.loads(out)
        except ValueError:
            data = None
    return code, data, out, err


# ---------------------------------------------------------------- findings (§0.3)
def check_shape(result):
    assert isinstance(result["ok"], bool)
    assert set(result["counts"]) >= {"error", "warning", "info"}
    for sev in ("error", "warning", "info"):
        assert result["counts"][sev] == sum(1 for f in result["findings"] if f["severity"] == sev)
    assert result["ok"] == (result["counts"]["error"] == 0)
    keys = [(f["path"], f["line"] or 0, f["rule"]) for f in result["findings"]]
    assert keys == sorted(keys)
    for f in result["findings"]:
        assert set(f) >= {"rule", "severity", "path", "line", "message", "excerpt"}
        assert f["severity"] in ("error", "warning", "info")
        assert f["line"] is None or (isinstance(f["line"], int) and f["line"] >= 1)
        assert "\\" not in f["path"]


def of(result, rule):
    return [f for f in result["findings"] if f["rule"] == rule]


def lines(result, rule):
    return [f["line"] for f in of(result, rule)]


def pairs(result):
    return [(f["rule"], f["line"]) for f in result["findings"]]


def write(path, data=""):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, bytes):
        path.write_bytes(data)
    else:
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(data)
    return path


def lint_text(tmp_path, monkeypatch, body, name="notes.md", **kw):
    monkeypatch.chdir(tmp_path)
    write(tmp_path / name, body)
    res = call("text_lint", paths=[name], **kw)
    check_shape(res)
    return res


# ---------------------------------------------------------------- git / bundle
GIT_LOCATION_VARS = ("GIT_DIR", "GIT_WORK_TREE", "GIT_INDEX_FILE", "GIT_OBJECT_DIRECTORY",
                     "GIT_ALTERNATE_OBJECT_DIRECTORIES", "GIT_COMMON_DIR", "GIT_NAMESPACE",
                     "GIT_CEILING_DIRECTORIES")


@pytest.fixture
def gitenv(tmp_path, monkeypatch):
    """Isolated git config and a fixed identity."""
    home = tmp_path / "_home"
    home.mkdir()
    gcfg = home / ".gitconfig"
    gcfg.write_text("", encoding="utf-8")
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.setenv("USERPROFILE", str(home))
    monkeypatch.setenv("XDG_CONFIG_HOME", str(home / ".config"))
    monkeypatch.setenv("GIT_CONFIG_GLOBAL", str(gcfg))
    monkeypatch.setenv("GIT_CONFIG_NOSYSTEM", "1")
    for k, v in {"GIT_AUTHOR_NAME": "Held Author", "GIT_AUTHOR_EMAIL": "held@example.com",
                 "GIT_COMMITTER_NAME": "Held Committer", "GIT_COMMITTER_EMAIL": "heldc@example.com"}.items():
        monkeypatch.setenv(k, v)
    for k in GIT_LOCATION_VARS:
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("TUNDLEKIT_NOW", "2026-09-16T10:00")
    return gcfg


def git(root, *args):
    env = {k: v for k, v in os.environ.items() if k not in GIT_LOCATION_VARS}
    p = subprocess.run(["git", "-C", str(root), *args], capture_output=True, env=env, timeout=120)
    if p.returncode != 0:
        raise AssertionError(f"git {args} failed: {p.stderr.decode(errors='replace')}")
    return p.stdout.decode("utf-8")


def changelog_text(version):
    return (f"# Changelog\n\nNewest first.\n\n{MARKER}\n\n## {version}  (2026-09-16 09:00)\n\nstart\n\n"
            "_1 added, 0 modified, 0 removed, 0 renamed_\n")


def tundle_tree(root, files=None, version="2026.09.16.4", init_git=True):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    if init_git:
        git(root, "init", "-q", "-b", "main")
    write(root / "VERSION", version + "\n")
    write(root / "CHANGELOG.md", changelog_text(version))
    for rel, data in (files or {}).items():
        write(root / rel, data)
    return root


# ---------------------------------------------------------------- deck specs and pptx
def content(title="Plain slide title", time="0:30", say=SAY25, source="Alita-G, Table 3", body=None, **kw):
    s = {"type": "content", "eyebrow": "Research", "title": title,
         "body": body if body is not None else {"kind": "bullets", "items": ["one point", "another point"]},
         "notes": {"time": time, "say": say}}
    if source is not None:
        s["source"] = source
    s.update(kw)
    return s


def title_slide(title="Capability capsules", subtitle="From goal text to gates", byline="Muk", time="0:15",
                say="Welcome."):
    return {"type": "title", "title": title, "subtitle": subtitle, "byline": byline,
            "notes": {"time": time, "say": say}}


def divider(title="The evidence", subtitle="What the papers show", time="0:05"):
    return {"type": "divider", "title": title, "subtitle": subtitle, "notes": {"time": time, "say": ""}}


def deck_spec(*slides, **meta):
    out = {"slides": list(slides)}
    if meta:
        out["meta"] = meta
    return out


def need_office():
    pytest.importorskip("pptx")
    pytest.importorskip("PIL")


def make_deck(path, slides):
    """slides: list of {"frames": [(text, size_pt, top_in)], "notes": str | None}."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for d in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for text, size, top in d.get("frames", []):
            tb = s.shapes.add_textbox(Inches(0.7), Inches(top), Inches(9), Inches(0.3))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(size)
        if d.get("notes") is not None:
            s.notes_slide.notes_text_frame.text = d["notes"]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return Path(path)


def slide(title, number, footer=None, footer_top=7.1, notes=None, time="0:40", say=SAY25, extra=()):
    frames = [("2 DESIGN", 11, 0.3), (title, 28, 0.6)] + list(extra)
    if footer is not None:
        frames.append((footer, 10, footer_top))
    if number is not None:
        frames.append((str(number), 10, 7.1))
    return {"frames": frames, "notes": notes if notes is not None else f"TIME {time}\nSAY: {say}"}


def mmss(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"


def iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def box(sh):
    return sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU


def text_shape(sl, startswith):
    hits = [sh for sh in iter_shapes(sl.shapes) if sh.has_text_frame and sh.text_frame.text.startswith(startswith)]
    assert hits, f"no text shape starting with {startswith!r}"
    return hits[0]


def build(tmp_path, spec, name="out.pptx", **kw):
    need_office()
    from pptx import Presentation

    out = Path(tmp_path) / name
    res = call("deck_build", spec=spec, out=str(out), **kw)
    return res, Presentation(str(out))


def make_png(path, w, h):
    from PIL import Image

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (w, h), (0, 114, 178)).save(str(path), "PNG")
    return Path(path)


def point_fill(series, idx):
    ser = series._element
    vals = ser.xpath(f'./c:dPt[c:idx/@val="{idx}"]/c:spPr/a:solidFill/a:srgbClr/@val')
    if vals:
        return str(vals[0]).upper()
    vals = ser.xpath("./c:spPr/a:solidFill/a:srgbClr/@val")
    return str(vals[0]).upper() if vals else None


def run_sizes(tf):
    out = set()
    for p in tf.paragraphs:
        for r in p.runs:
            if r.text.strip():
                size = r.font.size or p.font.size
                out.add(size.pt if size is not None else None)
    return out


# ---------------------------------------------------------------- SVG
def tag(el):
    return el.tag.split("}", 1)[-1]


def svg(text):
    root = ET.fromstring(text)
    assert tag(root) == "svg"
    return root


def with_class(root, name, cls):
    return [e for e in root.iter() if tag(e) == name and cls in (e.get("class") or "").split()]


def fill(el):
    f = el.get("fill")
    if f is None:
        for part in (el.get("style") or "").split(";"):
            if ":" in part and part.split(":", 1)[0].strip() == "fill":
                f = part.split(":", 1)[1].strip()
    return f.lower() if f else None


def boxes(res):
    return {k: (v["x"], v["y"], v["x"] + v["w"], v["y"] + v["h"]) for k, v in res["nodes"].items()}


def overlap(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def assert_layout_ok(res):
    b = boxes(res)
    for k, (x0, y0, x1, y1) in b.items():
        assert x0 >= -0.01 and y0 >= -0.01 and x1 <= res["width"] + 0.01 and y1 <= res["height"] + 0.01, k
    keys = list(b)
    for i, m in enumerate(keys):
        for n in keys[i + 1:]:
            assert not overlap(b[m], b[n]), (m, n)


def hide_converters(monkeypatch, tmp_path):
    """No cairosvg, nothing on PATH; pymupdf left alone (skip if it is missing)."""
    for p in ("C:/Program Files/Inkscape", "C:/Program Files (x86)/Inkscape", "/Applications/Inkscape.app",
              "/usr/bin/inkscape", "/usr/bin/rsvg-convert", "/opt/homebrew/bin/rsvg-convert",
              "/usr/local/bin/rsvg-convert"):
        if Path(p).exists():
            pytest.skip(f"a converter is installed at a fixed location: {p}")
    import sys

    monkeypatch.setitem(sys.modules, "cairosvg", None)
    monkeypatch.setenv("PATH", "")
    monkeypatch.chdir(tmp_path)


# ---------------------------------------------------------------- papers
def marker_txt(path, pages):
    return write(path, "".join(f"\n\n===== page {n} =====\n{body}" for n, body in enumerate(pages, 1)))


def ff_txt(path, pages):
    return write(path, "\f".join(pages) + "\f")


def make_pdf(path, texts):
    import pymupdf

    doc = pymupdf.open()
    for t in texts:
        page = doc.new_page(width=300, height=200)
        page.insert_text((20, 40), t, fontsize=11)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return Path(path)


def pdf_bytes(texts):
    import pymupdf

    doc = pymupdf.open()
    for t in texts:
        page = doc.new_page(width=300, height=200)
        page.insert_text((20, 40), t, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data
