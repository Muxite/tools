"""MANIFEST §14.4 deck_lint: several decks, times_file (D013 and combined D008), D014, the D005 additions and
D009 for .pptx. Decks are modelled on the real capsule/general presentations (footers at 7.02 in, shared
deck-times.json ledger)."""
from __future__ import annotations

import json

import pytest

from helpers_deck import SAY25, call, check_shape, content, rules, run_cli, spec, title_slide, write_json

GENERAL_TIMES = {"general": 1005, "capsule": 1085, "general_inserts": 245, "capsule_inserts": 90}


def mmss(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"


def make_deck(path, slides):
    """slides: list of {"frames": [(text, size_pt, top_in)], "notes": str | None}."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for d in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for text, size, top in d.get("frames", []):
            tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(10), Inches(0.3))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(size)
        if d.get("notes") is not None:
            s.notes_slide.notes_text_frame.text = d["notes"]
    prs.save(str(path))
    return path


def content_slide(title, number, footer=None, footer_top=7.02, time="0:30", say=SAY25, extra=()):
    frames = [("1 BACKGROUND", 11, 0.3), (title, 30, 0.58)] + list(extra)
    if footer is not None:
        frames.append((footer, 10, footer_top))
    frames.append((str(number), 10, 7.02))
    return {"frames": frames, "notes": f"TIME {time}\nSAY: {say}"}


def timed_deck(path, seconds, footer="MCP specification, revision 2026-07-28"):
    return make_deck(path, [content_slide("MCP is a protocol, not a tool", 1, footer, time=mmss(seconds))])


def lines(res, rule, path=None):
    return [f["line"] for f in rules(res, rule) if path is None or f["path"] == path]


# ====================================================================== D009 for pptx
def test_d009_pptx_footer_present(tmp_path):
    """§14.4: a text frame at or below 6.9 in holding text other than the number is a footer."""
    p = make_deck(tmp_path / "capsule.pptx", [
        content_slide("DeepSeek Harness (DSH): every layer is a plugin", 3,
                      footer="DSH 0.1.5-rc.2 at c291e7961a: docs/architecture.md:13, :121"),
    ])
    res = call("deck_lint", pptx_path=str(p))
    check_shape(res)
    assert rules(res, "D009") == []


def test_d009_pptx_only_slide_number_at_bottom(tmp_path):
    """§14.4: the slide number alone is not a footer."""
    p = make_deck(tmp_path / "capsule.pptx", [
        content_slide("Sourced slide", 1, footer="Beyond Task Completion, Table 4"),
        content_slide("Unsourced slide", 2),
    ])
    res = call("deck_lint", pptx_path=str(p))
    f = rules(res, "D009")
    assert [x["line"] for x in f] == [2]
    assert f[0]["severity"] == "warning"
    assert f[0]["path"] == str(p).replace("\\", "/")


def test_d009_pptx_text_above_threshold_is_not_footer(tmp_path):
    """§14.4: text whose top is above 6.9 in is body text, not a footer."""
    p = make_deck(tmp_path / "d.pptx", [content_slide("Body only", 1, footer="Not a footer", footer_top=6.5)])
    assert lines(call("deck_lint", pptx_path=str(p)), "D009") == [1]


def test_d009_pptx_threshold_is_inclusive(tmp_path):
    """§14.4: a frame whose top is exactly 6.9 in counts ("at or below")."""
    p = make_deck(tmp_path / "d.pptx", [content_slide("Edge", 1, footer="Paper, Table 2", footer_top=6.9)])
    assert lines(call("deck_lint", pptx_path=str(p)), "D009") == []


def test_d009_pptx_only_content_slides(tmp_path):
    """§14.4/§3.3: slides without a number frame are not content slides and get no D009."""
    p = make_deck(tmp_path / "d.pptx", [
        {"frames": [("AI4Research Capability Capsules", 40, 2.8)], "notes": "TIME 0:15\nSAY: Hello."},
        content_slide("Needs a source", 1),
    ])
    assert lines(call("deck_lint", pptx_path=str(p)), "D009") == [2]


def test_d009_built_deck(tmp_path):
    """§14.4 with §3.2: a built deck's source footer satisfies D009 on the .pptx."""
    pytest.importorskip("pptx")
    pytest.importorskip("PIL")
    s = spec(title_slide(), content(source="MCP specification, revision 2026-07-28"), content())
    out = tmp_path / "built.pptx"
    call("deck_build", spec=s, out=str(out))
    assert lines(call("deck_lint", pptx_path=str(out)), "D009") == [3]


# ====================================================================== several decks
def test_several_pptx_paths(tmp_path):
    """§14.4: pptx_path may be a list; each deck is linted and findings carry their deck's path."""
    a = make_deck(tmp_path / "general.pptx", [content_slide("Bad — title", 1, footer="src")])
    b = make_deck(tmp_path / "capsule.pptx", [content_slide("Fine title", 1, footer="src"),
                                               content_slide("Another — bad one", 2, footer="src")])
    pa, pb = a.as_posix(), b.as_posix()
    res = call("deck_lint", pptx_path=[pa, pb])
    check_shape(res)
    assert sorted((f["path"], f["line"]) for f in rules(res, "D003")) == sorted([(pa, 1), (pb, 2)])
    assert res["ok"] is False


def test_several_pptx_paths_cli(tmp_path, capsys):
    """§14.4: CLI takes several positional paths."""
    a = make_deck(tmp_path / "general.pptx", [content_slide("Fine", 1, footer="src")])
    b = make_deck(tmp_path / "capsule.pptx", [content_slide("Short say", 1, footer="src", say="too short")])
    code, data, _ = run_cli(["deck", "lint", a.as_posix(), b.as_posix(), "--json"], capsys)
    assert code == 0
    assert [(f["rule"], f["path"]) for f in data["findings"]] == [("D002", b.as_posix())]


# ====================================================================== times_file: D013 and D008
def test_times_file_match_no_d013(tmp_path):
    """§14.4: the deck id is the file stem; a matching entry gives no D013."""
    t = write_json(tmp_path / "deck-times.json", GENERAL_TIMES)
    p = timed_deck(tmp_path / "capsule.pptx", 1085)
    res = call("deck_lint", pptx_path=str(p), times_file=str(t))
    check_shape(res)
    assert rules(res, "D013") == []
    assert rules(res, "D008") == []


def test_times_file_mismatch_d013(tmp_path):
    """§14.4 D013: error, line null, when measured core seconds differ from the ledger entry."""
    t = write_json(tmp_path / "deck-times.json", GENERAL_TIMES)
    p = timed_deck(tmp_path / "capsule.pptx", 1100)
    res = call("deck_lint", pptx_path=p.as_posix(), times_file=str(t))
    check_shape(res)
    f = rules(res, "D013")
    assert len(f) == 1
    assert f[0]["severity"] == "error" and f[0]["line"] is None
    assert f[0]["path"] == p.as_posix()
    assert res["ok"] is False


def test_d008_uses_combined_time(tmp_path):
    """§14.4: with times_file, D008 uses the combined core time of every non-_inserts id."""
    t = write_json(tmp_path / "times.json", {"general": 1500, "capsule": 1000})
    p = timed_deck(tmp_path / "capsule.pptx", 1000)
    assert rules(call("deck_lint", pptx_path=str(p)), "D008") == []
    res = call("deck_lint", pptx_path=str(p), times_file=str(t))
    f = rules(res, "D008")
    assert len(f) == 1
    assert f[0]["severity"] == "warning" and f[0]["line"] is None
    assert rules(res, "D013") == []


def test_d008_combined_uses_measured_values(tmp_path):
    """§14.4: the entries for the decks given take their measured values (stale 100 -> measured 1000)."""
    t = write_json(tmp_path / "times.json", {"general": 1750, "capsule": 100})
    p = timed_deck(tmp_path / "capsule.pptx", 1000)
    res = call("deck_lint", pptx_path=str(p), times_file=str(t))
    assert len(rules(res, "D013")) == 1
    f = rules(res, "D008")
    assert len(f) == 1 and f[0]["severity"] == "error"


def test_d008_combined_ignores_inserts_entries(tmp_path):
    """§14.4: `_inserts` entries are left out of the combined time (2400 s equals the target: ok)."""
    t = write_json(tmp_path / "times.json",
                   {"general": 1500, "capsule": 900, "general_inserts": 245, "capsule_inserts": 600})
    p = timed_deck(tmp_path / "capsule.pptx", 900)
    res = call("deck_lint", pptx_path=str(p), times_file=str(t))
    assert rules(res, "D008") == []
    assert rules(res, "D013") == []


def test_ids_override_file_stem(tmp_path):
    """§14.4: `ids` gives each deck's id when the file stem is not the id."""
    t = write_json(tmp_path / "deck-times.json", GENERAL_TIMES)
    p = timed_deck(tmp_path / "AI4Research Capability Capsule Presentation - Muk.pptx", 1085)
    res = call("deck_lint", pptx_path=[str(p)], times_file=str(t), ids=["capsule"])
    check_shape(res)
    assert rules(res, "D013") == []
    res = call("deck_lint", pptx_path=[str(p)], times_file=str(t), ids=["general"])
    assert len(rules(res, "D013")) == 1


def test_two_decks_times_file(tmp_path):
    """§14.4: each deck is compared with its own entry; D008 combines the measured values."""
    t = write_json(tmp_path / "deck-times.json", GENERAL_TIMES)
    g = timed_deck(tmp_path / "general.pptx", 1400)
    c = timed_deck(tmp_path / "capsule.pptx", 1085)
    res = call("deck_lint", pptx_path=[g.as_posix(), c.as_posix()], times_file=str(t))
    check_shape(res)
    assert [f["path"] for f in rules(res, "D013")] == [g.as_posix()]
    f = rules(res, "D008")
    assert len(f) == 1 and f[0]["severity"] == "warning"   # 1400 + 1085 = 41:25


def test_times_file_cli(tmp_path, capsys):
    """§14.4: CLI `--times-file`."""
    t = write_json(tmp_path / "deck-times.json", GENERAL_TIMES)
    g = timed_deck(tmp_path / "general.pptx", 1005)
    c = timed_deck(tmp_path / "capsule.pptx", 1080)
    code, data, _ = run_cli(["deck", "lint", g, c, "--times-file", t, "--json"], capsys)
    assert code == 1
    assert [(f["rule"], f["path"]) for f in data["findings"] if f["severity"] == "error"] == \
        [("D013", c.as_posix())]
    assert json.loads(t.read_text(encoding="utf-8")) == GENERAL_TIMES   # lint never writes the ledger


# ====================================================================== D014
def spec_lint(the_spec):
    res = call("deck_lint", spec=the_spec)
    check_shape(res)
    return res


def src(**kw):
    kw.setdefault("source", "Beyond Task Completion, Table 4")
    return content(**kw)


def test_d014_em_dash_in_face_text():
    """§14.4 D014: an em dash in face text (a bullet) is a warning."""
    body = {"kind": "bullets", "items": ["Kept tools fail — held-out tests catch it"]}
    res = spec_lint(spec(src(), src(body=body)))
    f = rules(res, "D014")
    assert [x["line"] for x in f] == [2]
    assert f[0]["severity"] == "warning"


def test_d014_em_dash_in_say():
    """§14.4 D014: an em dash in SAY."""
    res = spec_lint(spec(src(say=SAY25 + " — and that is the point.")))
    assert lines(res, "D014") == [1]


def test_d014_does_not_repeat_d003_titles():
    """§14.4: titles are covered by D003, and D014 does not repeat them."""
    res = spec_lint(spec(src(title="MCP — a protocol, not a tool")))
    assert lines(res, "D003") == [1]
    assert lines(res, "D014") == []


@pytest.mark.parametrize("say", [
    SAY25 + " The gate passed.. Then it closed.",
    SAY25 + " The gate passed.  Then it closed.",
    SAY25 + " the tools  The library grows.",
    SAY25 + " first,  Then second.",
], ids=["double-period", "two-spaces-after-period", "two-spaces-after-lower", "two-spaces-after-comma"])
def test_d014_typography_in_say(say):
    """§14.4 D014: a double period, or 2 spaces between a lower-case letter or .,;:) and a capital."""
    assert lines(spec_lint(spec(src(say=say))), "D014") == [1]


@pytest.mark.parametrize("say", [
    SAY25 + " and so on... Then it closed.",
    SAY25 + " The gate passed. Then it closed.",
    SAY25 + " the tools  and the library.",
    SAY25 + " version 3  Then.",
], ids=["ellipsis", "single-space", "two-spaces-before-lower", "two-spaces-after-digit"])
def test_d014_clean_typography(say):
    """§14.4 D014: `...`, single spaces, and 2 spaces before a lower-case word or after a digit are fine."""
    assert lines(spec_lint(spec(src(say=say))), "D014") == []


def test_d014_pptx_say_and_face(tmp_path):
    """§14.4 D014 on a .pptx: SAY from the notes and face text from the frames."""
    p = make_deck(tmp_path / "d.pptx", [
        content_slide("Clean title", 1, footer="src", say=SAY25 + " The gate passed.  Then it closed."),
        content_slide("Clean title", 2, footer="src", extra=[("Models propose — code checks", 18, 2.0)]),
        content_slide("Clean — title", 3, footer="src"),
    ])
    res = call("deck_lint", pptx_path=str(p))
    assert lines(res, "D014") == [1, 2]
    assert lines(res, "D003") == [3]


# ====================================================================== D005 additions
@pytest.mark.parametrize("phrase", [
    "I'll go through the pipeline",
    "I’m going to show the gate",
    "I will go through the ledger",
    "Instead I'll describe the freeze",
    "instead i’ll describe the freeze",
])
def test_d005_new_phrases(phrase):
    """§14.4: D005 also matches I'll go through, I'm going to show, I will go through, Instead I'll."""
    res = spec_lint(spec(src(say=SAY25 + " " + phrase + ".")))
    f = rules(res, "D005")
    assert [x["line"] for x in f] == [1]
    assert f[0]["severity"] == "warning"
